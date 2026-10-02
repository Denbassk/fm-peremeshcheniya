# -*- coding: utf-8 -*-
"""
Расчёт заказа по магазинам из оборотной ведомости Торгсофт + валидация.

Формула:
    =IF(L<=0; (H+I)*1,2; IF(H>0; MAX((H+I)*1,2 - L; 0); 0))
    H = [-] реализация, I = [-] перемещение, L = на конец
Округление: аналог ОКРУГЛВВЕРХ (штучный товар -> целые, весовой -> шаг 0.1).
Для фасовки/шоубоксов действует порог FIX_MIN_FRAC и минимум min_q из справочника.
Нулевые позиции в заказ не попадают.
"""

import os
import re
import math
import sys
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ========================= НАСТРОЙКИ =========================

BASE_DIR = r"C:\Users\denba\Desktop\ПЕРЕМЕЩЕНИЯ"
IN_DIR   = os.path.join(BASE_DIR, "ВХОД")
OUT_DIR  = os.path.join(BASE_DIR, "ЗАКАЗЫ")
LOG_DIR  = os.path.join(BASE_DIR, "ЛОГИ")
ARC_DIR  = os.path.join(BASE_DIR, "АРХИВ")

KOEF = 1.2                 # коэффициент запаса из формулы

# ROUND_MODE:
#   "up"      - аналог ОКРУГЛВВЕРХ: любой остаток -> вверх (рекомендуется)
#   "half_up" - математическое: от 0.5 вверх, ниже - вниз
#   "smart"   - вверх, но остаток <= SMART_FRAC отбрасывается (гасит хвосты от *1,2)
ROUND_MODE  = "up"
SMART_FRAC  = 0.10
ROUND_EPS   = 1e-6         # гашение погрешности float (1.2000000000000002)

WEIGHT_STEP    = 0.1       # шаг округления весового товара, кг
REF_FILE     = os.path.join(BASE_DIR, "Справочник", "transfer_pack.csv")
FIX_MIN_FRAC = 0.34        # меньше 0.34 лотка/шоубокса - не заказывать
REF = {}
WEIGHT_AS_INT  = False     # True -> весовой тоже округлять до целых
TXT_WEIGHT_SEP = "."       # разделитель дробной части в txt для ТСД
DROP_BELOW     = 0.0       # расчёт ниже этого значения считать шумом (0 = выключено)

HEADER_ROW_HINT = 3
TXT_ENCODING = "cp1251"
TXT_NEWLINE  = "\r\n"

BALANCE_TOL     = 0.011    # допуск сверки баланса ведомости
ANOMALY_FACTOR  = 5        # заказ > factor * реализации -> в отчёт аномалий
ANOMALY_ABS     = 300      # заказ больше этого числа единиц -> в отчёт
MAX_REPORT_ROWS = 300      # сколько проблемных строк выводить в отчёт
NEG_END_TOL     = 0.5      # весовой минус мельче этого - шум взвешивания, не аномалия

BC_STD_LEN = (8, 12, 13, 14)   # стандартные длины EAN/UPC
BC_MANUAL = {                  # ручные соответствия, если контрольная не сходится
    # "54881005906": "054881005906",
}

FORMULA_TEXT = "IF(L<=0; (H+I)*1,2; IF(H>0; MAX((H+I)*1,2-L; 0); 0))"
PY_EXPR = "((H + I) * K) if (L <= 0) else (max((H + I) * K - L, 0.0) if (H > 0) else 0.0)"

# ======================= ОТЧЁТ ПРОВЕРОК ======================

class Report(object):
    def __init__(self):
        self.checks = []
        self.problems = []
        self.anomalies = []
        self.log = []

    def say(self, msg):
        print(msg)
        self.log.append(str(msg))

    def check(self, status, name, detail=""):
        self.checks.append((status, name, str(detail)))
        mark = {"OK": "[ OK ]", "WARN": "[ВНИМ]", "ERROR": "[ОШИБ]", "INFO": "[ИНФО]"}[status]
        self.say("%s %s%s" % (mark, name, ("  -  " + str(detail)) if detail else ""))

    def problem(self, kind, shop, bc, name, detail):
        if len(self.problems) < MAX_REPORT_ROWS:
            self.problems.append((kind, shop, bc, name, str(detail)))

    def anomaly(self, shop, bc, name, sale, out, end, raw, qty, why):
        self.anomalies.append((shop, bc, name, sale, out, end, round(raw, 3), qty, why))

    @property
    def errors(self):
        return [c for c in self.checks if c[0] == "ERROR"]

    @property
    def warns(self):
        return [c for c in self.checks if c[0] == "WARN"]

R = Report()

# ========================= УТИЛИТЫ ===========================

def ensure_dirs():
    for d in (BASE_DIR, IN_DIR, OUT_DIR, LOG_DIR, ARC_DIR):
        os.makedirs(d, exist_ok=True)


def norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def safe_name(s):
    s = str(s).strip().replace("/", "-").replace("\\", "-")
    s = re.sub(r'[:*?"<>|]', "", s)
    s = re.sub(r"\s+", " ", s).strip(" .")
    return s or "БЕЗ_АДРЕСА"


def fmt_barcode(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    if isinstance(v, str):
        s = v.strip().replace(" ", "").replace("\u00a0", "")
        return s[:-2] if s.endswith(".0") else s
    try:
        return str(int(round(float(v))))
    except Exception:
        return str(v).strip()


def ean_valid(bc):
    """Контрольная цифра EAN-8/12/13/14. Для внутренних кодов может не сходиться."""
    if not bc.isdigit() or len(bc) not in (8, 12, 13, 14):
        return False
    digits = [int(ch) for ch in bc]
    body, control = digits[:-1], digits[-1]
    body = body[::-1]
    total = 0
    for idx, d in enumerate(body):
        total += d * (3 if idx % 2 == 0 else 1)
    return (10 - total % 10) % 10 == control

def restore_barcode(bc):
    """Восстанавливает ведущие нули, потерянные Excel. -> (код, пояснение)"""
    if not bc or not bc.isdigit():
        return bc, ""
    if bc in BC_MANUAL:
        return BC_MANUAL[bc], "%s -> %s (вручную)" % (bc, BC_MANUAL[bc])
    if len(bc) in BC_STD_LEN and ean_valid(bc):
        return bc, ""
    for L in BC_STD_LEN:
        if L <= len(bc):
            continue
        cand = bc.zfill(L)
        if ean_valid(cand):
            return cand, "%s -> %s" % (bc, cand)
    return bc, ""


def frac_threshold():
    if ROUND_MODE == "up":
        return ROUND_EPS
    if ROUND_MODE == "half_up":
        return 0.5 - 1e-9
    return SMART_FRAC


def apply_rounding(x, step):
    """Округление вверх по шагу с учётом выбранного режима."""
    if x is None or x <= 0:
        return 0.0
    u = round(float(x) / step, 9)
    fl = math.floor(u)
    n = fl if (u - fl) <= frac_threshold() else fl + 1
    return round(n * step, 6)


def calc_order(h, i, l):
    base = (h + i) * KOEF
    if l <= 0:
        return base
    return max(base - l, 0.0) if h > 0 else 0.0


def calc_order_eval(h, i, l):
    """Независимая реализация: вычисление текста формулы."""
    return float(eval(PY_EXPR, {"max": max, "__builtins__": {}},
                      {"H": float(h), "I": float(i), "L": float(l), "K": KOEF}))


def qty_to_text(q, is_weight):
    if not is_weight or WEIGHT_AS_INT:
        return str(int(round(q)))
    dec = max(0, len(str(WEIGHT_STEP).split(".")[-1]))
    s = ("%." + str(dec) + "f") % q
    if float(s) == int(float(s)):
        s = str(int(float(s)))
    return s.replace(".", TXT_WEIGHT_SEP)

# ===================== САМОТЕСТИРОВАНИЕ ======================

def self_test():
    cases = [((1, 0, 6), 0.0), ((1, 0, 0), 1.2), ((3, 0, 7), 0.0), ((0, 0, 5), 0.0),
             ((0, 0, -2), 0.0), ((5, 0, 2), 4.0), ((2, 3, 1), 5.0), ((0, 4, 0), 4.8),
             ((1, 0, -1), 1.2), ((10, 2, 3), 11.4)]
    bad = []
    for (h, i, l), exp in cases:
        got = calc_order(h, i, l)
        got2 = calc_order_eval(h, i, l)
        if abs(got - exp) > 1e-9 or abs(got2 - exp) > 1e-9:
            bad.append("H=%s I=%s L=%s -> %s (ожидалось %s)" % (h, i, l, got, exp))
    if bad:
        R.check("ERROR", "Тест формулы", "; ".join(bad))
        return False
    R.check("OK", "Тест формулы", "%d контрольных случаев совпали" % len(cases))

    rbad = []
    if ROUND_MODE == "up":
        for x, exp in [(1.2, 2), (2.0, 2), (0.2, 1), (4.8, 5), (5.999, 6), (0.0, 0)]:
            if apply_rounding(x, 1) != exp:
                rbad.append("%s -> %s (ожидалось %s)" % (x, apply_rounding(x, 1), exp))
        for x, exp in [(1.046, 1.1), (0.846, 0.9), (3.0, 3.0)]:
            if abs(apply_rounding(x, 0.1) - exp) > 1e-9:
                rbad.append("вес %s -> %s (ожидалось %s)" % (x, apply_rounding(x, 0.1), exp))
    if apply_rounding(3.0000000001, 1) != 3:
        rbad.append("шум float 3.0000000001 -> лишняя единица")

    # порог фасовки: 0.30 лотка -> 0, 0.40 лотка -> 1 лоток
    if apply_rounding_thr(0.30 * 4.5, 4.5, FIX_MIN_FRAC) != 0.0:
        rbad.append("фасовка: 0.30 лотка не отброшена")
    if abs(apply_rounding_thr(0.40 * 4.5, 4.5, FIX_MIN_FRAC) - 4.5) > 1e-9:
        rbad.append("фасовка: 0.40 лотка не поднята до 1 упаковки")
    if rbad:
        R.check("ERROR", "Тест округления", "; ".join(rbad))
        return False
    R.check("OK", "Тест округления",
            "режим '%s', шаг веса %s, порог фасовки %s"
            % (ROUND_MODE, WEIGHT_STEP, FIX_MIN_FRAC))
    return True

# ==================== ЧТЕНИЕ ИСТОЧНИКА =======================

ALIASES = {
    "name":   ["название товара", "назва товару", "наименование"],
    "art":    ["артикул"],
    "bc":     ["штрих-код", "штрих код", "штрихкод"],
    "start":  ["на начало", "на початок"],
    "inc":    ["приход", "прихід"],
    "move_in": ["[+] перемещение", "[+] переміщення"],
    "sale":   ["[-] реализация", "[-] реалізація", "реализация", "реалізація"],
    "out":    ["[-] перемещение", "[-] переміщення"],
    "woff":   ["списание", "списання"],
    "ret":    ["возврат поставщикам", "повернення постачальникам"],
    "end":    ["на конец", "на кінець"],
    "shop":   ["центр учета", "центр обліку", "магазин"],
}


def find_source_file():
    """Самый актуальный файл: приоритет - дата в названии, затем дата изменения."""
    found = []
    for folder in (IN_DIR, BASE_DIR):
        if not os.path.isdir(folder):
            continue
        for nm in os.listdir(folder):
            low = nm.lower()
            if nm.startswith("~$") or not low.endswith((".xlsx", ".xlsm")):
                continue
            if not any(k in low for k in ("перемещ", "переміщ", "оборот", "ведом")):
                continue
            full = os.path.join(folder, nm)
            d = extract_date(full)
            found.append((d, os.path.getmtime(full), full))
    if not found:
        return None

    found.sort(key=lambda t: (t[0], t[1]), reverse=True)
    if len(found) > 1:
        R.check("WARN", "В папке несколько ведомостей",
                "найдено %d, взят файл за %s: %s | остальные: %s"
                % (len(found), found[0][0].strftime("%d.%m.%Y"),
                   os.path.basename(found[0][2]),
                   ", ".join(os.path.basename(f[2]) for f in found[1:4])))
    return found[0][2]


def extract_date(path):
    m = re.search(r"(\d{1,2})[.\-_ ](\d{1,2})[.\-_ ](\d{2,4})", os.path.basename(path))
    if m:
        d, mo, y = (int(g) for g in m.groups())
        y = y + 2000 if y < 100 else y
        try:
            return date(y, mo, d)
        except ValueError:
            pass
    return date.today()


def detect_header_row(path):
    probe = pd.read_excel(path, sheet_name=0, header=None, nrows=15)
    for idx in range(len(probe)):
        vals = [norm(v) for v in probe.iloc[idx].tolist() if pd.notna(v)]
        if any("штрих" in v for v in vals) and any("центр" in v or "магазин" in v for v in vals):
            return idx + 1
    return HEADER_ROW_HINT


def resolve_columns(df):
    cols = {norm(c): c for c in df.columns}
    res = {}
    for key, variants in ALIASES.items():
        hit = None
        for v in variants:
            if v in cols:
                hit = cols[v]
                break
        if hit is None:
            for v in variants:
                for cn, orig in cols.items():
                    if v in cn:
                        hit = orig
                        break
                if hit:
                    break
        res[key] = hit
    return res

# ======================= ПОДГОТОВКА ==========================

def prepare(df, c):
    num = lambda col: pd.to_numeric(df[col], errors="coerce")

    raw_sale, raw_out, raw_end = num(c["sale"]), num(c["out"]), num(c["end"])
    bad_cells = int(raw_sale.isna().sum() + raw_out.isna().sum() + raw_end.isna().sum())
    if bad_cells:
        R.check("WARN", "Нечисловые значения в расчётных столбцах",
                "%d ячеек трактованы как 0" % bad_cells)
    else:
        R.check("OK", "Числовые данные", "нечисловых ячеек в H/I/L нет")

    w = pd.DataFrame({
        "shop": df[c["shop"]].astype(str).str.strip(),
        "name": (df[c["name"]].astype(str).str.strip() if c.get("name") else ""),
        "art":  (df[c["art"]].astype(str).str.strip() if c.get("art") else ""),
        "bc":   df[c["bc"]].map(fmt_barcode),
        "sale": raw_sale.fillna(0.0),
        "out":  raw_out.fillna(0.0),
        "end":  raw_end.fillna(0.0),
        "start": (num(c["start"]).fillna(0.0) if c.get("start") else 0.0),
    })
    w["name"] = w["name"].replace({"nan": "", "None": ""})
    w["art"] = w["art"].replace({"nan": "", "None": ""})
    _fixed, _pairs = [], []
    for b in w["bc"]:
        nb, note = restore_barcode(b)
        _fixed.append(nb)
        if note:
            _pairs.append(note)
    w["bc"] = _fixed
    if _pairs:
        uniq_pairs = sorted(set(_pairs))
        R.check("WARN", "Восстановлены ведущие нули штрих-кодов",
                "%d кодов: %s" % (len(uniq_pairs), "; ".join(uniq_pairs[:6])))
        for p in uniq_pairs:
            R.problem("Ведущий ноль восстановлен (Excel обрезал)", "", p, "",
                      "проверь выгрузку Торгсофта")
    else:
        R.check("OK", "Ведущие нули штрих-кодов", "потерь не обнаружено")

    w["src_row"] = df.index + 1

    dropped = w[(w["bc"] == "") | (w["shop"] == "") | (w["shop"].str.lower() == "nan")]
    for _, r0 in dropped.iterrows():
        R.problem("Пропущена строка (нет штрих-кода или адреса)", r0["shop"], r0["bc"],
                  r0["name"], "строка источника ~%s" % r0["src_row"])
    if len(dropped):
        R.check("WARN", "Строки без штрих-кода/адреса", "исключено %d" % len(dropped))
    else:
        R.check("OK", "Полнота ключевых полей", "штрих-код и адрес есть во всех строках")

    w = w.drop(dropped.index)
    return w


def balance_check(df, c):
    need = ("start", "inc", "move_in", "sale", "out", "woff", "ret", "end")
    if any(not c.get(k) for k in need):
        R.check("WARN", "Сверка баланса ведомости", "нет части столбцов, проверка пропущена")
        return
    g = lambda k: pd.to_numeric(df[c[k]], errors="coerce").fillna(0.0)
    calc = g("start") + g("inc") + g("move_in") - g("sale") - g("out") - g("woff") - g("ret")
    diff = (calc - g("end")).abs()
    bad = int((diff > BALANCE_TOL).sum())
    if bad:
        R.check("WARN", "Сверка баланса (начало+приход+перемещ-расход = конец)",
                "расхождений: %d строк, макс. %.3f" % (bad, diff.max()))
    else:
        R.check("OK", "Сверка баланса ведомости", "все %d строк сходятся" % len(df))


def detect_weight(w):
    frac = lambda s: (s - s.round()).abs() > 0.001
    flag = frac(w["sale"]) | frac(w["out"]) | frac(w["end"]) | frac(w["start"])
    by_bc = flag.groupby(w["bc"]).any()
    art_hint = w.assign(h=w["art"].str.contains(r"[.,]", na=False)).groupby("bc")["h"].any()
    weight = (by_bc | art_hint).to_dict()
    R.check("INFO", "Определение весового товара",
            "весовых артикулов: %d из %d" % (sum(1 for v in weight.values() if v), len(weight)))
    return weight


def barcode_check(w):
    uniq = sorted(set(w["bc"]))
    nondigit = [b for b in uniq if not b.isdigit()]
    badlen = [b for b in uniq if b.isdigit() and len(b) not in (8, 12, 13, 14)]
    badsum = [b for b in uniq if b.isdigit() and len(b) in (8, 12, 13, 14) and not ean_valid(b)]
    if nondigit:
        R.check("ERROR", "Штрих-коды: посторонние символы",
                "%d кодов, напр.: %s" % (len(nondigit), ", ".join(nondigit[:5])))
    else:
        R.check("OK", "Штрих-коды: только цифры", "проверено %d кодов" % len(uniq))
    if badlen:
        R.check("WARN", "Штрих-коды: нетипичная длина",
                "%d кодов, напр.: %s" % (len(badlen), ", ".join(badlen[:5])))
    if badsum:
        R.check("WARN", "Штрих-коды: не сходится контрольная цифра",
                "%d кодов (норма для внутренних), напр.: %s" % (len(badsum), ", ".join(badsum[:5])))

    mix = w.groupby("bc")["name"].nunique()
    mix = mix[mix > 1]
    if len(mix):
        R.check("WARN", "Один штрих-код у разных названий", "%d кодов" % len(mix))
        for b in list(mix.index)[:20]:
            names = sorted(set(w.loc[w["bc"] == b, "name"]))
            R.problem("Один код - разные товары", "", b, " | ".join(names[:3]), "")
    return not bool(nondigit)

# ========================== РАСЧЁТ ===========================

def load_reference():
    """Справочник упаковок из BigQuery: barcode, pname, unit_type, order_step, min_q, n, confidence."""
    if not os.path.exists(REF_FILE):
        R.check("WARN", "Справочник упаковок", "нет файла %s, работа по эвристике" % REF_FILE)
        return {}

    try:
        d = pd.read_csv(REF_FILE, sep=None, engine="python",
                        encoding="utf-8-sig", dtype=str, keep_default_na=False)
    except Exception as e:
        R.check("ERROR", "Справочник упаковок",
                "не удалось прочитать %s: %s" % (REF_FILE, e))
        return {}

    d.columns = [str(cn).strip().lstrip("\ufeff").lower() for cn in d.columns]
    lack = [k for k in ("barcode", "unit_type", "order_step") if k not in d.columns]
    if lack:
        R.check("ERROR", "Справочник упаковок",
                "нет столбцов %s; в файле: %s" % (lack, list(d.columns)))
        return {}

    def _num(v):
        s = str(v).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
        if s == "" or s.lower() in ("nan", "none", "null"):
            return None
        try:
            x = float(s)
        except ValueError:
            return None
        return x if x > 0 else None

    out, broken, dups = {}, [], 0
    for _, r0 in d.iterrows():
        bc = fmt_barcode(r0["barcode"])
        if not bc:
            continue
        bc, _ = restore_barcode(bc)
        st = _num(r0["order_step"])
        if st is None:
            broken.append("%s '%s' -> '%s'" % (bc, str(r0.get("pname", ""))[:26],
                                               r0["order_step"]))
            continue
        if bc in out:
            dups += 1
        out[bc] = {
            "unit_type": str(r0["unit_type"]).strip(),
            "step": st,
            "min_q": _num(r0["min_q"]) if "min_q" in d.columns else None,
            "confidence": str(r0.get("confidence", "")).strip(),
        }

    if broken:
        R.check("WARN", "Справочник: нечисловой шаг",
                "%d строк пропущено (похоже на порчу Excel - даты вместо чисел), напр.: %s"
                % (len(broken), "; ".join(broken[:3])))
    if dups:
        R.check("WARN", "Справочник: дубли штрих-кодов",
                "%d повторов, взята последняя строка" % dups)

    cnt = lambda t: sum(1 for v in out.values() if v["unit_type"] == t)
    mq = sum(1 for v in out.values() if v.get("min_q"))
    R.check("OK", "Справочник упаковок",
            "%d позиций: недел. %d, фасовка %d, вес %d; с min_q: %d"
            % (len(out), cnt("sht_nedelimyy"), cnt("ves_fix"), cnt("ves_plav"), mq))
    return out


def resolve_step(bc, heur_weight):
    """-> (шаг, порог дробной части, минимальный заказ, пояснение)"""
    info = REF.get(bc)
    if info is None or not info.get("step"):
        s = WEIGHT_STEP if (heur_weight and not WEIGHT_AS_INT) else 1.0
        return s, frac_threshold(), 0.0, "эвристика: " + ("вес" if heur_weight else "шт")

    s = float(info["step"])
    ut = str(info.get("unit_type", ""))
    mq = float(info.get("min_q") or 0.0)
    tail = (", min %g" % mq) if mq > 0 else ""

    if ut == "sht_nedelimyy":
        return s, FIX_MIN_FRAC, mq, "шоубокс %g шт%s" % (s, tail)
    if ut == "ves_fix":
        return s, FIX_MIN_FRAC, mq, "фасовка %g кг%s" % (s, tail)
    if ut == "ves_plav":
        return (s if s > 0 else WEIGHT_STEP), frac_threshold(), mq, "вес, шаг %g кг" % s
    return s, frac_threshold(), mq, (ut or "не указано")


def apply_rounding_thr(x, step, thr):
    if x is None or x <= 0:
        return 0.0
    u = round(float(x) / float(step), 9)
    fl = math.floor(u)
    n = fl if (u - fl) <= thr else fl + 1
    return round(n * float(step), 6)


def build_orders(w, weight_map):
    import numpy as np

    h, i, l = w["sale"].to_numpy(float), w["out"].to_numpy(float), w["end"].to_numpy(float)
    vec = np.where(l <= 0, (h + i) * KOEF,
                   np.where(h > 0, np.maximum((h + i) * KOEF - l, 0.0), 0.0))
    loop = np.array([calc_order(a, b, cc) for a, b, cc in zip(h, i, l)])
    ev = np.array([calc_order_eval(a, b, cc) for a, b, cc in zip(h, i, l)])

    d1 = float(np.max(np.abs(vec - loop))) if len(vec) else 0.0
    d2 = float(np.max(np.abs(vec - ev))) if len(vec) else 0.0
    if max(d1, d2) > 1e-9:
        R.check("ERROR", "Сверка трёх реализаций формулы",
                "макс. расхождение %.12f" % max(d1, d2))
    else:
        R.check("OK", "Сверка трёх реализаций формулы",
                "векторная = построчная = разбор текста, %d строк" % len(vec))

    w = w.copy()
    w["raw"] = vec

    grp = (w.groupby(["shop", "bc"], as_index=False)
            .agg(name=("name", "first"), art=("art", "first"),
                 sale=("sale", "sum"), out=("out", "sum"),
                 end=("end", "sum"), raw=("raw", "sum"), dup=("bc", "size")))

    dups = int((grp["dup"] > 1).sum())
    R.check("INFO", "Дубли товара внутри магазина",
            "сведено позиций: %d" % dups if dups else "дублей нет")

    grp["is_weight"] = grp["bc"].map(lambda b: bool(weight_map.get(b, False)))
    if DROP_BELOW > 0:
        grp.loc[grp["raw"] < DROP_BELOW, "raw"] = 0.0

    _s, _t, _m, _nt = [], [], [], []
    for _b, _hw in zip(grp["bc"], grp["is_weight"]):
        a, b2, m2, c2 = resolve_step(_b, _hw)
        _s.append(a); _t.append(b2); _m.append(m2); _nt.append(c2)
    grp["step"], grp["thr"], grp["min_q"], grp["unit_note"] = _s, _t, _m, _nt

    in_ref = int(grp["bc"].map(lambda b: b in REF).sum())
    R.check("INFO", "Покрытие справочником",
            "%d из %d позиций найдены в transfer_pack" % (in_ref, len(grp)))

    _ex = grp["bc"].map(lambda b: REF.get(b, {}).get("unit_type") == "exclude")
    if int(_ex.sum()):
        R.check("INFO", "Исключено справочником", "%d позиций" % int(_ex.sum()))
        grp = grp[~_ex].copy()

    # 1) базовое округление по шагу и порогу
    grp["base"] = [apply_rounding_thr(x, s, t)
                   for x, s, t in zip(grp["raw"], grp["step"], grp["thr"])]

    # 2) нижняя граница: страховка при нулевом остатке + min_q из справочника
    floors, n_ins, n_mq = [], 0, 0
    for base, raw, end, step, thr, mq in zip(grp["base"], grp["raw"], grp["end"],
                                             grp["step"], grp["thr"], grp["min_q"]):
        fl = 0.0
        if base <= 0 and raw > 0 and end <= 0 and thr > 0.01:
            fl = float(step)
            n_ins += 1
        if mq and mq > 0 and max(base, fl) > 0:
            mq_c = apply_rounding_thr(mq, step, ROUND_EPS)
            if max(base, fl) + 1e-9 < mq_c:
                fl = mq_c
                n_mq += 1
        floors.append(fl)
    grp["floor"] = floors
    grp["qty"] = [max(b, f) for b, f in zip(grp["base"], grp["floor"])]

    if n_ins:
        R.check("INFO", "Страховка при нулевом остатке",
                "поднято до 1 упаковки: %d позиций (порог %s срезал бы в ноль)"
                % (n_ins, FIX_MIN_FRAC))
    if n_mq:
        R.check("INFO", "Поднято до min_q справочника", "%d позиций" % n_mq)

    _chk = grp[grp["step"] > 1.0]
    _bad = _chk[((_chk["qty"] / _chk["step"]) -
                 (_chk["qty"] / _chk["step"]).round()).abs() > 1e-6]
    if len(_bad):
        R.check("ERROR", "Кратность нарушена", "%d позиций" % len(_bad))
    else:
        R.check("OK", "Кратность соблюдена",
                "все заказы кратны шагу (%d позиций с шагом > 1)" % len(_chk))

    # контроль направления округления
    # для фасовки/шоубоксов недобор допустим, но не больше порога FIX_MIN_FRAC от упаковки
    plain = grp[(grp["qty"] > 0) & (grp["thr"] <= 0.01)]
    wrong = plain[plain["qty"] + 1e-9 < plain["raw"]]

    packs = grp[(grp["qty"] > 0) & (grp["thr"] > 0.01)]
    over = packs[packs["raw"] - packs["qty"] > packs["thr"] * packs["step"] + 1e-9]
    cut = int((packs["raw"] - packs["qty"] > 1e-9).sum())

    if len(wrong) or len(over):
        R.check("ERROR", "Округление не вверх",
                "штучных/весовых занижено %d, фасовок ниже допустимого %d"
                % (len(wrong), len(over)))
        for _, b0 in list(wrong.iterrows())[:10] + list(over.iterrows())[:10]:
            R.problem("Занижено против расчёта", b0["shop"], b0["bc"], b0["name"],
                      "расчёт %.3f -> заказ %.3f, шаг %g" % (b0["raw"], b0["qty"], b0["step"]))
    else:
        R.check("OK", "Округление вверх",
                "занижений нет; неполная упаковка отброшена по порогу %s: %d позиций"
                % (FIX_MIN_FRAC, cut))

    res = grp[grp["qty"] > 0].copy()
    zero = len(grp) - len(res)
    R.check("INFO", "Отброшено нулевых позиций", zero)

    for _, r0 in res.iterrows():
        why = []
        if r0["sale"] <= 0 and r0["qty"] > 0:
            why.append("заказ без реализации (остаток <= 0)")
        if r0["end"] < 0:
            _piece = float(r0["step"]).is_integer() and float(r0["step"]) >= 1
            if _piece or r0["end"] <= -NEG_END_TOL:
                why.append("отрицательный остаток %.3f%s"
                           % (r0["end"], " (штучный товар!)" if _piece else ""))
        if r0["qty"] > ANOMALY_ABS:
            why.append("очень крупный заказ")
        if r0["sale"] > 0 and r0["qty"] > ANOMALY_FACTOR * r0["sale"] and r0["qty"] >= 10:
            why.append("заказ > %dx реализации" % ANOMALY_FACTOR)
        if why:
            R.anomaly(r0["shop"], r0["bc"], r0["name"], r0["sale"], r0["out"],
                      r0["end"], r0["raw"], r0["qty"], "; ".join(why))
    if R.anomalies:
        R.check("WARN", "Аномальные позиции", "%d шт. - см. лист 'Аномалии'" % len(R.anomalies))
    else:
        R.check("OK", "Аномальные позиции", "не обнаружены")

    return res.sort_values(["shop", "name"], kind="stable")

# ========================== ЗАПИСЬ ===========================

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
ERR_FILL = PatternFill("solid", fgColor="FCE4E4")

HDRS = ["№", "Название товара", "Артикул", "Штрих-код", "[-] реализация",
        "[-] перемещение", "на конец", "Расчёт (скрипт)", "Расчёт (формула Excel)",
        "ЗАКАЗ", "Контроль", "Упаковка"]
WIDTHS = [5, 44, 11, 15, 13, 14, 10, 15, 17, 10, 14, 24]
LAST_COL = len(HDRS)          # 12 -> L


def excel_qty_formula(row_i, step, thr, floor=0.0):
    """Excel-формула округления, идентичная логике скрипта (с порогом строки и минимумом)."""
    s = repr(float(step))
    t = repr(float(thr))
    u = "I{r}/{s}".format(r=row_i, s=s)
    f = "IF({u}-INT({u})<={t},INT({u}),INT({u})+1)*{s}".format(u=u, t=t, s=s)
    if floor and float(floor) > 0:
        f = "MAX({f},{fl})".format(f=f, fl=repr(float(floor)))
    return f


def write_shop_xlsx(path, shop, day, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказ"
    ws["A1"] = "Заказ: %s" % shop
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Дата ведомости: %s     Позиций: %d     Единиц: %s     Формула: %s     "
                "Округление: %s, порог фасовки %s"
                % (day.strftime("%d.%m.%Y"), len(rows),
                   round(float(rows["qty"].sum()), 3),
                   FORMULA_TEXT, ROUND_MODE, FIX_MIN_FRAC))
    ws["A2"].font = Font(italic=True, size=9)

    hr = 4
    for j, h in enumerate(HDRS, start=1):
        c0 = ws.cell(row=hr, column=j, value=h)
        c0.font = Font(bold=True)
        c0.fill = HEAD_FILL
        c0.border = BORDER
        c0.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    r = hr
    for n, (_, row) in enumerate(rows.iterrows(), start=1):
        r += 1
        step = float(row["step"])
        thr = float(row["thr"])
        floor = float(row.get("floor", 0.0) or 0.0)

        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value=row["name"])
        ws.cell(row=r, column=3, value=row["art"])
        bc = ws.cell(row=r, column=4, value=row["bc"])
        bc.number_format = "@"
        ws.cell(row=r, column=5, value=float(row["sale"]))
        ws.cell(row=r, column=6, value=float(row["out"]))
        ws.cell(row=r, column=7, value=float(row["end"]))
        ws.cell(row=r, column=8, value=round(float(row["raw"]), 4))
        ws.cell(row=r, column=9, value=(
            "=IF(G{r}<=0,(E{r}+F{r})*{k},IF(E{r}>0,MAX((E{r}+F{r})*{k}-G{r},0),0))"
        ).format(r=r, k=repr(KOEF)))
        q = ws.cell(row=r, column=10, value=round(float(row["qty"]), 3))
        q.font = Font(bold=True)
        ws.cell(row=r, column=11, value=(
            '=IF(AND(ABS(H{r}-I{r})<0.0005,ABS({f}-J{r})<0.0005),"OK","РАСХОЖДЕНИЕ")'
        ).format(r=r, f=excel_qty_formula(r, step, thr, floor)))
        ws.cell(row=r, column=12, value=row.get("unit_note", ""))

        for j in range(1, LAST_COL + 1):
            cell = ws.cell(row=r, column=j)
            cell.border = BORDER
            if 5 <= j <= 10:
                cell.number_format = "0.###"
                cell.alignment = Alignment(horizontal="center")

    last = r
    r += 1
    ws.cell(row=r, column=9, value="ИТОГО:").font = Font(bold=True)
    tot = ws.cell(row=r, column=10, value=round(float(rows["qty"].sum()), 3))
    tot.font = Font(bold=True)
    tot.fill = WARN_FILL

    ws["A3"] = ('=IF(COUNTIF(K5:K{n},"РАСХОЖДЕНИЕ")=0,"Excel-контроль: OK, расчёт скрипта '
                'совпал с формулой во всех строках","Excel-контроль: РАСХОЖДЕНИЙ "'
                '&COUNTIF(K5:K{n},"РАСХОЖДЕНИЕ"))').format(n=last)
    ws["A3"].font = Font(bold=True, color="1F6F3F")

    for j, wd in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = wd
    ws.freeze_panes = "A%d" % (hr + 1)
    ws.auto_filter.ref = "A%d:%s%d" % (hr, get_column_letter(LAST_COL), last)
    wb.save(path)


def write_shop_txt(path, rows):
    lines = []
    for _, row in rows.iterrows():
        _int = float(row["step"]).is_integer() and float(row["step"]) >= 1
        lines.append("%s;%s" % (row["bc"], qty_to_text(float(row["qty"]), not _int)))
    with open(path, "w", encoding=TXT_ENCODING, errors="replace", newline="") as f:
        f.write(TXT_NEWLINE.join(lines) + TXT_NEWLINE)
    return lines

# =================== ПРОВЕРКА ЗАПИСАННОГО =====================

TXT_LINE_RE = re.compile(r"^\d{6,14};\d+(?:[.,]\d+)?$")


def verify_written(shop_dir, fname, rows):
    """Обратное чтение xlsx и txt со сверкой с расчётом."""
    problems = []
    xlsx_path = os.path.join(shop_dir, fname + ".xlsx")
    txt_path = os.path.join(shop_dir, fname + ".txt")

    expected = {}
    for _, row in rows.iterrows():
        expected[row["bc"]] = round(float(row["qty"]), 3)

    wb = load_workbook(xlsx_path, data_only=False)
    ws = wb.active
    got = {}
    r = 5
    while ws.cell(row=r, column=4).value not in (None, ""):
        b = fmt_barcode(ws.cell(row=r, column=4).value)
        v = ws.cell(row=r, column=10).value
        got[b] = round(float(v), 3) if isinstance(v, (int, float)) else None
        r += 1
    wb.close()
    if got != expected:
        problems.append("xlsx: позиций %d вместо %d или количества не совпали"
                        % (len(got), len(expected)))

    with open(txt_path, "r", encoding=TXT_ENCODING) as f:
        raw_lines = [ln.strip() for ln in f.read().splitlines() if ln.strip()]
    if len(raw_lines) != len(expected):
        problems.append("txt: строк %d вместо %d" % (len(raw_lines), len(expected)))
    tgot = {}
    for ln in raw_lines:
        if not TXT_LINE_RE.match(ln):
            problems.append("txt: неверный формат строки '%s'" % ln)
            continue
        b, q = ln.split(";")
        tgot[b] = round(float(q.replace(",", ".")), 3)
    if tgot != expected:
        problems.append("txt: содержимое не совпало с расчётом")
    if os.path.getsize(txt_path) == 0:
        problems.append("txt: файл пустой")
    return problems

# ======================= ОТЧЁТЫ ==============================

def write_summary(path, day, summary, src):
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = "Сводка заказов на %s" % day.strftime("%d.%m.%Y")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Источник: %s" % src
    ws["A2"].font = Font(italic=True, size=9)
    for j, h in enumerate(["№", "Адрес (Центр учета)", "Позиций", "Единиц",
                           "Проверка файлов", "Папка"], start=1):
        c0 = ws.cell(row=4, column=j, value=h)
        c0.font = Font(bold=True)
        c0.fill = HEAD_FILL
    for n, it in enumerate(summary, start=1):
        ws.cell(row=4 + n, column=1, value=n)
        ws.cell(row=4 + n, column=2, value=it["shop"])
        ws.cell(row=4 + n, column=3, value=it["positions"])
        ws.cell(row=4 + n, column=4, value=it["units"])
        st = ws.cell(row=4 + n, column=5, value=it["verify"])
        st.fill = OK_FILL if it["verify"] == "OK" else ERR_FILL
        ws.cell(row=4 + n, column=6, value=it["folder"])
    for j, wd in enumerate([5, 32, 10, 10, 34, 62], start=1):
        ws.column_dimensions[get_column_letter(j)].width = wd
    ws.freeze_panes = "A5"
    wb.save(path)


def write_validation(path, day, src, shops_all, shops_with_order):
    wb = Workbook()
    ws = wb.active
    ws.title = "Проверки"
    ws["A1"] = "Отчёт валидации заказа на %s" % day.strftime("%d.%m.%Y")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Источник: %s   |   запуск %s" % (src, datetime.now().strftime("%d.%m.%Y %H:%M"))
    ws["A2"].font = Font(italic=True, size=9)
    verdict = "ЕСТЬ ОШИБКИ" if R.errors else ("ЕСТЬ ПРЕДУПРЕЖДЕНИЯ" if R.warns else "ВСЁ ЧИСТО")
    v = ws["A3"]
    v.value = "ИТОГ: %s (ошибок %d, предупреждений %d)" % (verdict, len(R.errors), len(R.warns))
    v.font = Font(bold=True, size=12)
    v.fill = ERR_FILL if R.errors else (WARN_FILL if R.warns else OK_FILL)

    for j, h in enumerate(["Статус", "Проверка", "Детали"], start=1):
        c0 = ws.cell(row=5, column=j, value=h)
        c0.font = Font(bold=True)
        c0.fill = HEAD_FILL
    for n, (st, nm, det) in enumerate(R.checks, start=1):
        ws.cell(row=5 + n, column=1, value=st).fill = {
            "OK": OK_FILL, "WARN": WARN_FILL, "ERROR": ERR_FILL,
            "INFO": PatternFill("solid", fgColor="EDEDED")}[st]
        ws.cell(row=5 + n, column=2, value=nm)
        ws.cell(row=5 + n, column=3, value=det)
    for j, wd in enumerate([10, 52, 90], start=1):
        ws.column_dimensions[get_column_letter(j)].width = wd

    ws2 = wb.create_sheet("Проблемы данных")
    for j, h in enumerate(["Тип", "Адрес", "Штрих-код", "Товар", "Детали"], start=1):
        c0 = ws2.cell(row=1, column=j, value=h)
        c0.font = Font(bold=True)
        c0.fill = HEAD_FILL
    for n, row in enumerate(R.problems, start=1):
        for j, v2 in enumerate(row, start=1):
            ws2.cell(row=1 + n, column=j, value=v2)
    for j, wd in enumerate([38, 26, 16, 42, 40], start=1):
        ws2.column_dimensions[get_column_letter(j)].width = wd

    ws3 = wb.create_sheet("Аномалии")
    for j, h in enumerate(["Адрес", "Штрих-код", "Товар", "реализация", "перемещение",
                           "на конец", "Расчёт", "ЗАКАЗ", "Причина"], start=1):
        c0 = ws3.cell(row=1, column=j, value=h)
        c0.font = Font(bold=True)
        c0.fill = HEAD_FILL
    for n, row in enumerate(R.anomalies, start=1):
        for j, v2 in enumerate(row, start=1):
            ws3.cell(row=1 + n, column=j, value=v2)
    for j, wd in enumerate([26, 16, 42, 12, 13, 11, 11, 10, 44], start=1):
        ws3.column_dimensions[get_column_letter(j)].width = wd

    ws4 = wb.create_sheet("Адреса")
    ws4.cell(row=1, column=1, value="Адрес (Центр учета)").font = Font(bold=True)
    ws4.cell(row=1, column=2, value="Есть заказ").font = Font(bold=True)
    for n, s in enumerate(shops_all, start=1):
        ws4.cell(row=1 + n, column=1, value=s)
        ws4.cell(row=1 + n, column=2, value="да" if s in shops_with_order else "нет")
    ws4.column_dimensions["A"].width = 34
    wb.save(path)

# =========================== MAIN ============================

def main():
    ensure_dirs()
    R.say("=" * 70)
    R.say("РАСЧЁТ ЗАКАЗА С ВАЛИДАЦИЕЙ  |  %s" % datetime.now().strftime("%d.%m.%Y %H:%M:%S"))
    R.say("=" * 70)

    if not self_test():
        R.say("Самотестирование не пройдено, работа прервана.")
        return 2

    src = find_source_file()
    if not src:
        R.check("ERROR", "Файл-источник", "не найден, положите ведомость в %s" % IN_DIR)
        return 2
    day = extract_date(src)
    R.check("OK", "Файл-источник", src)
    R.check("INFO", "Дата заказа", day.strftime("%d.%m.%Y"))

    hr = detect_header_row(src)
    R.check("OK" if hr == HEADER_ROW_HINT else "WARN", "Строка заголовков", hr)

    df = pd.read_excel(src, sheet_name=0, header=hr - 1).dropna(how="all")
    R.check("INFO", "Прочитано строк", len(df))

    c = resolve_columns(df)
    missing = [k for k in ("bc", "sale", "out", "end", "shop") if not c.get(k)]
    if missing:
        R.check("ERROR", "Обязательные столбцы",
                "не найдены: %s; заголовки файла: %s" % (missing, list(df.columns)))
        return 2
    R.check("OK", "Сопоставление столбцов",
            "H='%s', I='%s', L='%s', адрес='%s'" % (c["sale"], c["out"], c["end"], c["shop"]))

    balance_check(df, c)
    w = prepare(df, c)
    if not barcode_check(w):
        return 2
    weight_map = detect_weight(w)

    shops_all = sorted(set(w["shop"]))
    R.check("INFO", "Уникальных адресов", "%d: %s" % (len(shops_all), "; ".join(shops_all)))

    names, collide = {}, []
    for s in shops_all:
        names.setdefault(safe_name(s).lower(), []).append(s)
    for k, v in names.items():
        if len(v) > 1:
            collide.append(" / ".join(v))
    if collide:
        R.check("WARN", "Совпадение имён папок после очистки",
                "; ".join(collide) + " - будут суффиксы (2), (3)")
    else:
        R.check("OK", "Имена файлов/папок", "коллизий нет")

    REF.clear()
    REF.update(load_reference())

    orders = build_orders(w, weight_map)
    if R.errors:
        R.say("Обнаружены критические ошибки, файлы не записаны.")
        write_validation(os.path.join(LOG_DIR, "ВАЛИДАЦИЯ_СБОЙ_%s.xlsx"
                                      % datetime.now().strftime("%Y-%m-%d_%H-%M-%S")),
                         day, src, shops_all, set())
        return 2
    if orders.empty:
        R.check("WARN", "Результат", "по формуле ни одна позиция не дала количество > 0")
        return 0

    day_dir = os.path.join(OUT_DIR, day.strftime("%Y-%m-%d"))
    os.makedirs(day_dir, exist_ok=True)

    R.say("-" * 70)
    summary, used, verify_fail = [], {}, 0
    for shop, rows in orders.groupby("shop", sort=True):
        base_nm = safe_name(shop)
        nm, k = base_nm, 2
        while nm.lower() in used:
            nm = "%s (%d)" % (base_nm, k)
            k += 1
        used[nm.lower()] = shop

        shop_dir = os.path.join(day_dir, nm)
        os.makedirs(shop_dir, exist_ok=True)
        write_shop_xlsx(os.path.join(shop_dir, nm + ".xlsx"), shop, day, rows)
        write_shop_txt(os.path.join(shop_dir, nm + ".txt"), rows)

        probs = verify_written(shop_dir, nm, rows)
        status = "OK" if not probs else "СБОЙ: " + "; ".join(probs)
        if probs:
            verify_fail += 1
            for p in probs:
                R.problem("Проверка записанных файлов", shop, "", "", p)
        summary.append({"shop": shop, "positions": len(rows),
                        "units": round(float(rows["qty"].sum()), 3),
                        "folder": shop_dir, "verify": status})
        R.say("%-24s позиций: %4d   единиц: %9s   %s"
              % (shop, len(rows), round(float(rows["qty"].sum()), 3),
                 "OK" if not probs else "СБОЙ"))

    if verify_fail:
        R.check("ERROR", "Обратная сверка записанных файлов",
                "расхождения в %d магазинах" % verify_fail)
    else:
        R.check("OK", "Обратная сверка записанных файлов",
                "xlsx и txt перечитаны, совпали с расчётом (%d магазинов)" % len(summary))

    tot_units = round(sum(s["units"] for s in summary), 3)
    if abs(tot_units - round(float(orders["qty"].sum()), 3)) > 0.001:
        R.check("ERROR", "Сверка итогов", "сумма по файлам не равна сумме расчёта")
    else:
        R.check("OK", "Сверка итогов", "всего единиц %s" % tot_units)

    write_summary(os.path.join(day_dir, "_СВОДКА.xlsx"), day, summary, src)
    write_validation(os.path.join(day_dir, "_ВАЛИДАЦИЯ.xlsx"), day, src,
                     shops_all, set(s["shop"] for s in summary))

    R.say("-" * 70)
    R.say("Магазинов с заказом : %d из %d" % (len(summary), len(shops_all)))
    R.say("Всего позиций       : %d" % sum(s["positions"] for s in summary))
    R.say("Всего единиц        : %s" % tot_units)
    R.say("Ошибок / предупр.   : %d / %d" % (len(R.errors), len(R.warns)))
    R.say("Результат           : %s" % day_dir)
    R.say("Отчёт валидации     : %s" % os.path.join(day_dir, "_ВАЛИДАЦИЯ.xlsx"))

    lp = os.path.join(LOG_DIR, "log_%s.txt" % datetime.now().strftime("%Y-%m-%d_%H-%M-%S"))
    with open(lp, "w", encoding="utf-8") as f:
        f.write("\n".join(R.log))
    print("Лог: %s" % lp)
    return 2 if R.errors else 0


if __name__ == "__main__":
    try:
        code = main()
    except Exception as e:
        print("СБОЙ: %s: %s" % (type(e).__name__, e))
        import traceback
        traceback.print_exc()
        code = 2
    try:
        input("\nГотово. Нажмите Enter для выхода...")
    except EOFError:
        pass
    sys.exit(code)
