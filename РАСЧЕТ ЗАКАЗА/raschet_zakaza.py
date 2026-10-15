# -*- coding: utf-8 -*-
"""
Расчёт заказа по магазинам из оборотной ведомости Торгсофт + валидация.

Формула:
    =IF(L<=0; (H+I)*1,2 - L; IF(H>0; MAX((H+I)*1,2 - L; 0); 0))
    Минус на конец (L<0) = потребность, добавляется к заказу (NEG_AS_NEED).
    H = [-] реализация, I = [-] перемещение, L = на конец
Округление: аналог ОКРУГЛВВЕРХ (штучный товар -> целые, весовой -> шаг 0.1).
Для фасовки/шоубоксов действует порог FIX_MIN_FRAC и минимум min_q из справочника.
Нулевые позиции в заказ не попадают.
"""

import os
import re
import io
import math
import sys
import glob
import time
from datetime import date, datetime, timedelta, timezone

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ========================= НАСТРОЙКИ =========================

def _app_dir():
    """Папка, в которой лежит exe (или сам .py). Абсолютных путей в коде нет:
    программа работает из любой папки и на любом компьютере.
    Можно переопределить переменной окружения FM_BASE_DIR."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    try:
        return os.path.dirname(os.path.abspath(__file__))
    except NameError:
        return os.path.dirname(os.path.abspath(sys.argv[0]))


# В окне (console=False) стандартных потоков нет: без этой заглушки
# первый же print уронил бы программу с AttributeError.
if sys.stdout is None:
    sys.stdout = open(os.devnull, "w")
if sys.stderr is None:
    sys.stderr = sys.stdout

BASE_DIR = os.environ.get("FM_BASE_DIR", "").strip() or _app_dir()
IN_DIR   = os.path.join(BASE_DIR, "ВХОД")   # прежняя общая папка: читается, не создаётся
IN_FM_DIR   = os.path.join(BASE_DIR, u"Ведомости Family")
IN_UA_DIR   = os.path.join(BASE_DIR, u"Ведомости ЮА Маркет")
DIST_IN_DIR = os.path.join(BASE_DIR, u"Списки раздачи")
OUT_DIR  = os.path.join(BASE_DIR, "ЗАКАЗЫ")
LOG_DIR  = os.path.join(BASE_DIR, "ЛОГИ")
ARC_DIR  = os.path.join(BASE_DIR, "АРХИВ")

KOEF = 1.2                 # коэффициент запаса из формулы

# PATCH-NEG-END v1: минусовой остаток на конец = продали то, чего ещё нет на приходе.
# True  -> минус добавляется к заказу: (H+I)*1,2 + |L|
# False -> прежнее поведение: минус трактуется как ноль
NEG_AS_NEED = True

# ROUND_MODE:
#   "up"      - аналог ОКРУГЛВВЕРХ: любой остаток -> вверх (рекомендуется)
#   "half_up" - математическое: от 0.5 вверх, ниже - вниз
#   "smart"   - вверх, но остаток <= SMART_FRAC отбрасывается (гасит хвосты от *1,2)
ROUND_MODE  = "up"
SMART_FRAC  = 0.10
ROUND_EPS   = 1e-6         # гашение погрешности float (1.2000000000000002)

WEIGHT_STEP    = 0.1       # шаг округления весового товара, кг

# ---- справочник упаковок: BigQuery + локальный кэш ----
BQ_PROJECT  = "family-market-analytics"
BQ_TABLE    = "family-market-analytics.family_market.transfer_pack_v5"
BQ_MANUAL   = "family-market-analytics.family_market.transfer_pack_manual"
REF_FILE    = os.path.join(BASE_DIR, "Справочник", "transfer_pack.csv")  # кэш и офлайн-фолбэк
REF_TTL_H   = 12          # свежесть кэша, часов; старше - идём в BQ
REF_OFFLINE = False       # True -> в BQ не ходить вообще, только кэш

_HERE = _app_dir()

CRED_ENV  = "GOOGLE_APPLICATION_CREDENTIALS"
CRED_MASK = "family-market-analytics-*.json"
CRED_DIRS = [
    os.path.join(_HERE, "credentials"),
    os.path.join(BASE_DIR, "credentials"),
    os.path.join(os.path.expanduser("~"), ".gcp"),
]

FIX_MIN_FRAC = 0.34        # меньше 0.34 лотка/шоубокса - не заказывать
MIN_Q_MIN_N  = 5           # min_q применяется только если наблюдений >= 5

# ---- сводный заказ поставщику ----
# Перемещение РЦ -> магазин: коробку можно вскрыть (шаг 0.1 / 1 шт).
# Заказ поставщику: только целыми коробками.
MAKE_SUPPLIER_ORDER = True   # писать сводный файл _ЗАКАЗ_ПОСТАВЩИКУ.xlsx

# ---- раздача конкретного списка с Полевой (режим выбирается в окне) ----
DIST_FILE        = ""                   # путь к списку; пусто = обычный расчёт
# источник раздачи - склад своей базы, получателем быть не может (решение 14.10)
DIST_SOURCE_SHOP    = u"Полевая-Склад"      # Family
DIST_SOURCE_SHOP_UA = u"Полевая-Склад ЮА"   # ЮА Маркет: склад и опт
DIST_LIMIT_STOCK = True                 # резать потребность под наличие
DIST_DIR_SUFFIX  = u"_РАЗДАЧА"          # своя папка результата
# Позиции из списка, которых на точках никогда не было: истории нет,
# формула их не видит. Раздаём разнарядкой.
DIST_NEW_MODE = "oborot"   # oborot = пропорционально обороту точки
                           # ravno  = поровну всем точкам
                           # net    = не раздавать, оставить на Полевой
DIST_NEW_TOP  = 0          # 0 = всем; N = только N крупнейшим точкам
DIST_NEW_MIN  = 3        # меньше этого на точку не везём (шт или кг)


# ---- что не перемещаем никогда ----
# Розлив, весовая развеска, кеги и напитки из кофеаппарата - это
# комплектация или производство на точке, между магазинами не ездят.
DIST_SKIP_BC = (
    "4820002713199",   # Оболонь Пиво Сеньор Картель 0,33л з/б - не перемещаем
    "2978950020707",   # Вода Імператорська 1л
)
DIST_SKIP_NAME = (
    u"імператорськ",      # вода на розлив
    u"овочі/фрукти",      # весовая развеска
    u"розливн",           # Оболонь Київське Розливне - наливают на точке
    u" кег",              # БІР  Кег ... - комплектация
    u"кегове",            # Кегове Бір Хаус ...
    u"зерн.кава",         # готовый напиток из кофеаппарата
    u"прем.кава",         # он же, вторая линейка
    u"кришка д/стак",     # тех.товар: крышки к стаканам кофеаппарата
)
# Группы из файла вывоза, которые не раздаём (колонка "Группа").
# Сравнение по подстроке, регистр не важен.
DIST_SKIP_GROUP = (
    u"тех.товар",
)
DIST_SKIP_ONLY_IN_DIST = False

# ---- как делить остаток с Полевой (PATCH-DIST-SHARE v1) ----
# Доля точки = DIST_BC_WEIGHT * продажи ЭТОГО ШК + остальное * общий оборот точки.
# Только оборот (прежнее поведение) = DIST_SHARE_BY_BC = False.
DIST_SHARE_BY_BC = True
DIST_BC_WEIGHT   = 0.7    # 1.0 = только продажи ШК; 0.0 = только оборот точки
DIST_SKIP_DEAD   = True   # не везти туда, где ШК лежит (остаток > 0) и не продаётся


def dist_on():
    return bool(DIST_FILE) and os.path.exists(DIST_FILE)


def dist_suffix():
    return DIST_DIR_SUFFIX if dist_on() else ""


# ---- два юрлица: две базы Торгсофта (PATCH-ENTITY v1) ----
# У магазинов базы ЮА Маркет в «Центре учета» в конце всегда «ЮА».
# Ведомость каждой базы считается отдельно: своя папка результата и свои txt
# (в ТСД заливаются из своей базы). Имена магазинов не переписываются.
UA_SHOP_SUFFIX    = u"ЮА"
ENTITY_LABEL      = {"FM": u"Family", "UA": u"ЮА Маркет"}
ENTITY_DIR_SUFFIX = {"FM": u"", "UA": u"_ЮА"}   # ЗАКАЗЫ\<дата>_ЮА
ENTITY            = "FM"                         # юрлицо текущего прогона


def is_ua_shop(s):
    return bool(re.search(r"\s%s$" % UA_SHOP_SUFFIX.lower(), str(s).strip().lower()))


def entity_of_shops(shops):
    """FM / UA по магазинам ведомости; MIX - вперемешку; None - магазинов нет."""
    flags = {is_ua_shop(x) for x in shops
             if str(x).strip() and str(x).strip().lower() != "nan"}
    if not flags:
        return None
    if len(flags) > 1:
        return "MIX"
    return "UA" if flags.pop() else "FM"


def entity_suffix():
    return ENTITY_DIR_SUFFIX.get(ENTITY, "")


def dist_source():
    """Склад, с которого раздаём: у каждой базы свой."""
    return DIST_SOURCE_SHOP_UA if ENTITY == "UA" else DIST_SOURCE_SHOP


# ---- склады: заказ на них не считается никогда (решение пользователя 14.10) ----
# Полевая-Склад (Family) - только склад, продаж нет.
# Полевая-Склад ЮА - склад и опт: продажи есть, но перемещения на него не делаем.
NO_ORDER_SHOPS = (u"Полевая-Склад", u"Полевая-Склад ЮА")


def is_no_order_shop(s):
    return route_key(s) in {route_key(x) for x in NO_ORDER_SHOPS}


def drop_no_order_shops(w):
    """Строки складов убираются до расчёта: ни заказа, ни доли в раздаче."""
    m = w["shop"].map(is_no_order_shop)
    if m.any():
        R.check("INFO", u"Склад не получатель",
                u"%s: исключено %d строк, заказ на склад не считается"
                % ("; ".join(sorted(set(w.loc[m, "shop"]))), int(m.sum())))
    return w[~m]
SUPPLIER_MIN_FRAC   = ROUND_EPS   # ROUND_EPS = всегда вверх; 0.34 = экономный вариант
INDIVISIBLE = ("ves_fix", "sht_nedelimyy")   # не дробится даже при перемещении
SUP_MAX_PCS_AUTO = 24    # коробка штучного товара крупнее - только через ручной список
# --- Приходная коробка != шаг перемещения ------------------------------
# Для этих групп число в скобках = транспортная коробка поставщика.
# В магазин перемещаем блоками, а не коробками.
BOX_ONLY_PATTERNS = ("пакет", "пляшка", "бутыл", "запальнич", "зажигалк")
# Слово должно стоять В НАЧАЛЕ названия, иначе ловится "Майонез 170 гр пакет",
# "Оливки пакет 250 г", "Водафон Стартовий Пакет" - это не упаковочный товар.
TRANSFER_BLOCK_RULES = (
    ("пакет", 50),
    ("пляшка", 5),
    ("бутыл", 5),
    ("запальнич", 5),
    ("зажигалк", 5),
)
# Пакеты: всегда кратно 50 и всегда вверх (49 -> 50, 55 -> 100), порог не применяется.
BAG_BLOCK = 50
BAG_ALWAYS_UP = True
# Пакеты, у которых слово стоит не первым (бренд впереди) - ведём списком.
BAG_EXTRA_BC = ()   # пакеты, у которых слово стоит не первым словом названия
# Кузя Сервіс - бытовые мусорные пакеты в рулонах, возятся поштучно: под правило 50 не попадают.

# ---- стаканы: в магазин везём упаковками, а не поштучно ----
# поштучный ШК -> (упаковочный ШК, сколько штук в упаковке)
CUP_PACKS = {
    "2978950015550": ("2978950003595", 50),    # Паперовий Малюнок 175 мл -> Упаковка 50 шт
    "2978950015567": ("2978950003601", 50),    # Паперовий Малюнок 250 мл -> Упаковка 50 шт
    "2978950017660": ("2978950003694", 100),   # Пластик 180 мл -> Упаковка 100 шт
    "2978950015574": ("2978950003717", 50),    # Пластик 300 мл -> Упаковка 50 шт
    "2978950015598": ("2978950003700", 50),    # Пластик 500 мл -> Упаковка 50 шт
    "2978950038856": ("2978950003687", 100),   # Пластик 100 мл -> вместимость уточнить
}
# Стакан Паперовий Малюнок 340 мл (2978950003618) упаковки не имеет - остаётся поштучно.

# ---- пакеты, которые возят только упаковками ----
# Поштучная карточка заменяется упаковочной: закупщица заказывает упаковки,
# а не тысячи штук. Слева - карточка «1 шт», справа - «Упаковка N шт».
BAG_PACKS = {
    "2978950015505": ("2978950004523", 100),   # Пакет Майка 100 шт -> Упаковка 100 шт
    "2978950015482": ("2978950004493", 250),   # Пакет Смайл (250шт) -> Упаковка 250 шт
}

# Общая таблица замен «поштучно -> упаковка»
PACK_SWAP = dict(CUP_PACKS)
PACK_SWAP.update(BAG_PACKS)

# ---- комплектность: к пивным бутылкам едут крышки и ручки ----
BC_KRISHKA = "2978950002598"        # Кришка біла - по одной на каждую бутылку
BC_RUCHKA  = "2978950002604"        # Ручка 1-2 л ПЕТ - только к бутылкам 1; 1,5; 2 л
BC_PLYASHKI_ALL = ("2978950002550", "2978950002567", "2978950017301",
                   "2978950002574", "2978950002581")          # 0,5 / 1 / 1,5 / 2 / 3 л
BC_PLYASHKI_RUCHKA = ("2978950002567", "2978950017301", "2978950002574")  # 1 / 1,5 / 2 л
# Доля блока, ниже которой потребность не тянем на целый блок.
# 0.0 = всегда округлять вверх до блока.
BLOCK_MIN_FILL = 0.5

FIX_MIN_FRAC = 0.34        # меньше 0.34 лотка/шоубокса - не заказывать
MIN_Q_MIN_N  = 5           # min_q применяется только если наблюдений >= 5
MIN_Q_MAX    = 100         # min_q крупнее - мусор из единичного прихода
VES_MIN_Q_N  = 10          # столько перемещений нужно, чтобы верить порции весового
VES_MIN_FILL = 0.34        # мельче этой доли порции весовой товар не возим

# ---- весовой товар в ящиках, которые не вскрывают ----
# Ящик мармелада, зефира, печенья на РЦ не подфасовывают: либо целый ящик,
# либо ничего. Дробное количество по такой позиции закупщица исполнить не может.
VES_BOX_RULE      = True   # включить правило «ящик или ноль»
VES_BOX_MIN       = 1.5    # ящиком считаем вес от этого значения, кг.
                           # Ниже - это порция (конфеты Рошен, ХБФ идут по 1 кг),
                           # её вскрывать можно, правило «ящик или ноль» не для неё.
VES_BOX_STOCK_MAX = 1.0    # остаток не больше этого - везём ящик
VES_BOX_QTY       = 1      # сколько ящиков везём за одну волну
# Реализация 0, а на остатке хвост: по одной выгрузке не отличить "продалось
# вчистую" от "лежит мёртвым". Решаем по величине хвоста: крошка = товар
# кончился, везём ящик; заметный остаток = ждём следующей волны, иначе
# доложенный сверху ящик уйдёт в списание.
VES_BOX_NEED_DEMAND = True   # False -> прежнее поведение (ящик без спроса)
VES_BOX_CRUMB       = 0.3    # хвост не больше этого, кг - считаем, что кончился

# Поставщики, у которых есть подфасовка: ящик вскрывают, возим порциями.
# Ключ - начало названия карточки, значение - порция в кг.
# До Бочкового: скобка в названии - это вес ящика прихода, а не шаг перемещения.
# Проверено по приходам на РЦ (incoming_transactions, delivery_type='РЦ'):
# у этих поставщиков приходы НЕ кратны скобке из названия - значит ящика,
# который нельзя вскрыть, там нет.
VES_PORTION_BY_NAME = {
    u"до бочкового":     1.0,   # есть подфасовка, возим по 1 кг
    u"роганський мк":    0.1,   # весовая колбаса, приходы 1,624 / 2,198 кг
    u"безлюдовський мк": 0.1,   # то же самое
    u"глобино мк":       0.1,   # то же самое
}
REF = {}
WEIGHT_CONFLICTS = []      # справочник говорит «шт», ведомость показывает дробь
PACK_MIN_FRAC = ROUND_EPS  # связки/шоубоксы: всегда вверх (60 при связке 50 -> 100)
MAX_KG_STEP  = 10.0        # фасовки весового товара больше 10 кг не бывает
MAX_BOX_PCS  = 200         # больше - это цена в артикуле, а не вложение ящика
REF = {}
WEIGHT_CONFLICTS = []      # справочник говорит «шт», ведомость показывает дробь
WEIGHT_AS_INT  = False     # True -> весовой тоже округлять до целых
TXT_WEIGHT_SEP = "."       # разделитель дробной части в txt для ТСД
DROP_BELOW     = 0.0       # расчёт ниже этого значения считать шумом (0 = выключено)

HEADER_ROW_HINT = 3
TXT_ENCODING = "cp1251"
TXT_NEWLINE  = "\r\n"

BALANCE_TOL     = 0.011    # допуск сверки баланса ведомости
ANOMALY_FACTOR  = 5        # заказ > factor * реализации -> в отчёт аномалий
ANOMALY_ABS     = 3000     # заказ больше этого числа единиц -> в отчёт.
# Ведомость за месяц: 300 давало сотни ложных срабатываний (яйцо, пакеты, крышки).
MAX_REPORT_ROWS = 300      # сколько проблемных строк выводить в отчёт
NEG_END_TOL     = 0.5      # весовой минус мельче этого - шум взвешивания, не аномалия

BC_STD_LEN = (8, 12, 13, 14)   # стандартные длины EAN/UPC
BC_MANUAL = {                  # ручные соответствия, если контрольная не сходится
    # "54881005906": "054881005906",
}

if NEG_AS_NEED:
    FORMULA_TEXT = "IF(L<=0; (H+I)*1,2-L; IF(H>0; MAX((H+I)*1,2-L; 0); 0))"
    PY_EXPR = "((H + I) * K - L) if (L <= 0) else (max((H + I) * K - L, 0.0) if (H > 0) else 0.0)"
else:
    FORMULA_TEXT = "IF(L<=0; (H+I)*1,2; IF(H>0; MAX((H+I)*1,2-L; 0); 0))"
    PY_EXPR = "((H + I) * K) if (L <= 0) else (max((H + I) * K - L, 0.0) if (H > 0) else 0.0)"

# ======================= ОТЧЁТ ПРОВЕРОК ======================

class Report(object):
    def __init__(self):
        self.checks = []
        self.problems = []
        self.anomalies = []
        self.log = []

    def say(self, msg):
        print(msg)
        self.log.append(str(msg))

    def check(self, status, name, detail=""):
        self.checks.append((status, name, str(detail)))
        mark = {"OK": "[ OK ]", "WARN": "[ВНИМ]", "ERROR": "[ОШИБ]", "INFO": "[ИНФО]"}[status]
        self.say("%s %s%s" % (mark, name, ("  -  " + str(detail)) if detail else ""))

    def problem(self, kind, shop, bc, name, detail):
        if len(self.problems) < MAX_REPORT_ROWS:
            self.problems.append((kind, shop, bc, name, str(detail)))

    def anomaly(self, shop, bc, name, sale, out, end, raw, qty, why):
        self.anomalies.append((shop, bc, name, sale, out, end, round(raw, 3), qty, why))

    @property
    def errors(self):
        return [c for c in self.checks if c[0] == "ERROR"]

    @property
    def warns(self):
        return [c for c in self.checks if c[0] == "WARN"]

R = Report()
RUN_INFO = {}   # итоги последнего прогона для окна

# ========================= УТИЛИТЫ ===========================

def ensure_dirs():
    for d in (BASE_DIR, IN_FM_DIR, IN_UA_DIR, DIST_IN_DIR, OUT_DIR, LOG_DIR, ARC_DIR):
        os.makedirs(d, exist_ok=True)


def norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def safe_name(s):
    s = str(s).strip().replace("/", "-").replace("\\", "-")
    s = re.sub(r'[:*?"<>|]', "", s)
    s = re.sub(r"\s+", " ", s).strip(" .")
    return s or "БЕЗ_АДРЕСА"


def fmt_barcode(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    if isinstance(v, str):
        s = v.strip().replace(" ", "").replace("\u00a0", "")
        return s[:-2] if s.endswith(".0") else s
    try:
        return str(int(round(float(v))))
    except Exception:
        return str(v).strip()


def ean_valid(bc):
    """Контрольная цифра EAN-8/12/13/14. Для внутренних кодов может не сходиться."""
    if not bc.isdigit() or len(bc) not in (8, 12, 13, 14):
        return False
    digits = [int(ch) for ch in bc]
    body, control = digits[:-1], digits[-1]
    body = body[::-1]
    total = 0
    for idx, d in enumerate(body):
        total += d * (3 if idx % 2 == 0 else 1)
    return (10 - total % 10) % 10 == control

def restore_barcode(bc):
    """Восстанавливает ведущие нули, потерянные Excel. -> (код, пояснение)"""
    if not bc or not bc.isdigit():
        return bc, ""
    if bc in BC_MANUAL:
        return BC_MANUAL[bc], "%s -> %s (вручную)" % (bc, BC_MANUAL[bc])
    if len(bc) in BC_STD_LEN and ean_valid(bc):
        return bc, ""
    for L in BC_STD_LEN:
        if L <= len(bc):
            continue
        cand = bc.zfill(L)
        if ean_valid(cand):
            return cand, "%s -> %s" % (bc, cand)
    return bc, ""


def frac_threshold():
    if ROUND_MODE == "up":
        return ROUND_EPS
    if ROUND_MODE == "half_up":
        return 0.5 - 1e-9
    return SMART_FRAC


def apply_rounding(x, step):
    """Округление вверх по шагу с учётом выбранного режима."""
    if x is None or x <= 0:
        return 0.0
    u = round(float(x) / step, 9)
    fl = math.floor(u)
    n = fl if (u - fl) <= frac_threshold() else fl + 1
    return round(n * step, 6)


def calc_order(h, i, l):
    base = (h + i) * KOEF
    if l <= 0:
        return (base - l) if NEG_AS_NEED else base
    return max(base - l, 0.0) if h > 0 else 0.0


def calc_order_eval(h, i, l):
    """Независимая реализация: вычисление текста формулы."""
    return float(eval(PY_EXPR, {"max": max, "__builtins__": {}},
                      {"H": float(h), "I": float(i), "L": float(l), "K": KOEF}))


def qty_to_text(q, is_weight):
    if not is_weight or WEIGHT_AS_INT:
        return str(int(round(q)))
    s = ("%.3f" % float(q)).rstrip("0").rstrip(".")
    return (s or "0").replace(".", TXT_WEIGHT_SEP)

# ===================== САМОТЕСТИРОВАНИЕ ======================

def self_test():
    cases = [((1, 0, 6), 0.0), ((1, 0, 0), 1.2), ((3, 0, 7), 0.0), ((0, 0, 5), 0.0),
             ((5, 0, 2), 4.0), ((2, 3, 1), 5.0), ((0, 4, 0), 4.8), ((10, 2, 3), 11.4)]
    if NEG_AS_NEED:
        cases += [((0, 0, -2), 2.0), ((1, 0, -1), 2.2), ((3, 1, -2.5), 7.3)]
    else:
        cases += [((0, 0, -2), 0.0), ((1, 0, -1), 1.2)]
    bad = []
    for (h, i, l), exp in cases:
        got = calc_order(h, i, l)
        got2 = calc_order_eval(h, i, l)
        if abs(got - exp) > 1e-9 or abs(got2 - exp) > 1e-9:
            bad.append("H=%s I=%s L=%s -> %s (ожидалось %s)" % (h, i, l, got, exp))
    if bad:
        R.check("ERROR", "Тест формулы", "; ".join(bad))
        return False
    R.check("OK", "Тест формулы", "%d контрольных случаев совпали" % len(cases))

    rbad = []
    if ROUND_MODE == "up":
        for x, exp in [(1.2, 2), (2.0, 2), (0.2, 1), (4.8, 5), (5.999, 6), (0.0, 0)]:
            if apply_rounding(x, 1) != exp:
                rbad.append("%s -> %s (ожидалось %s)" % (x, apply_rounding(x, 1), exp))
        for x, exp in [(1.046, 1.1), (0.846, 0.9), (3.0, 3.0)]:
            if abs(apply_rounding(x, 0.1) - exp) > 1e-9:
                rbad.append("вес %s -> %s (ожидалось %s)" % (x, apply_rounding(x, 0.1), exp))
    if apply_rounding(3.0000000001, 1) != 3:
        rbad.append("шум float 3.0000000001 -> лишняя единица")

    # min_q из единичного прихода не должен попадать в заказ
    if _sane_min_q({"min_q": 1000.0, "n": 1}) != 0.0:
        rbad.append("min_q=1000 при n=1 не отброшен")
    if _sane_min_q({"min_q": 6.0, "n": 40}) != 6.0:
        rbad.append("достоверный min_q=6 при n=40 потерян")
    if _sane_min_q({"min_q": 2.7, "n": 1}) != 0.0:
        rbad.append("весовой min_q при n=1 не отброшен")

    # заказ поставщику: 5.1 кг потребности при коробке 2.2 -> 6.6 (3 коробки)
    if abs(apply_rounding_thr(5.1, 2.2, ROUND_EPS) - 6.6) > 1e-9:
        rbad.append("заказ поставщику: 5.1 при коробке 2.2 не поднят до 6.6")
    # перемещение того же товара дробится по 0.1
    if abs(apply_rounding_thr(2.73, 0.1, ROUND_EPS) - 2.8) > 1e-9:
        rbad.append("перемещение: 2.73 не округлено до 2.8")

    # порог фасовки: 0.30 лотка -> 0, 0.40 лотка -> 1 лоток
    if apply_rounding_thr(0.30 * 4.5, 4.5, FIX_MIN_FRAC) != 0.0:
        rbad.append("фасовка: 0.30 лотка не отброшена")
    if abs(apply_rounding_thr(0.40 * 4.5, 4.5, FIX_MIN_FRAC) - 4.5) > 1e-9:
        rbad.append("фасовка: 0.40 лотка не поднята до 1 упаковки")
    if rbad:
        R.check("ERROR", "Тест округления", "; ".join(rbad))
        return False
    R.check("OK", "Тест округления",
            "режим '%s', шаг веса %s, порог фасовки %s"
            % (ROUND_MODE, WEIGHT_STEP, FIX_MIN_FRAC))

    pcases = [
        ("АВК Ваф.Трубочки Бам-Бук Згущене Молоко 1кг (4,4)", "4,4", True, None, 4.4),
        ("АВК Цукерки Труфальє 1кг", "6", True, None, 6.0),
        ("Щедрики Печиво 4,5кг", "", True, 4.5, None),
        ("Печиво Вівсяне 2,4кг (9,6)", "9,6", True, 2.4, 9.6),
        ("ДО Бочкового Розлив 1кг", "1000", True, None, None),
        ("АВК Kresko Банановий Смак 140г", "24", False, None, 24.0),
        ("Якобс Монарх 3в1 13г (20)", "20", False, 20.0, 20.0),
        ("МакКофе Оригінал Стик 18г", "26", False, 26.0, 26.0),
        ("Нескафе Класик 250г банка", "6", False, None, 6.0),
        ("Агропром Соломка Для Коктейля (*100) 1 шт", "", False, 1.0, None),
        ("Агропром Стакан Паперовий Малюнок 175 мл Упаковка 50 шт", "", False, 1.0, None),
        ("Агропром Комплект №1 (42)", "", False, None, 42.0),
    ]
    pbad = []
    for nm, ar, wt, exp_s, exp_b in pcases:
        g = parse_pack_from_name(nm, ar, wt)
        if g["step"] != exp_s or g["box"] != exp_b:
            pbad.append("%s -> шаг %s / ящик %s (ожидалось %s / %s)"
                        % (nm[:34], g["step"], g["box"], exp_s, exp_b))
    if pbad:
        R.check("ERROR", "Тест парсера названий", "; ".join(pbad))
        return False
    R.check("OK", "Тест парсера названий",
            "%d случаев: кг вне скобок = шаг, скобка/артикул = ящик, шоубокс кофе"
            % len(pcases))

    if apply_rounding_thr(60, 50, PACK_MIN_FRAC) != 100:
        R.check("ERROR", "Тест связок", "60 при связке 50 не дало 100")
        return False
    R.check("OK", "Тест связок", "60 -> 100, 101 -> 150 (всегда вверх)")

    ebad = []
    for nm, exp in ((u"Шевченко 341 ЮА", True), (u"Бучми 32 юа", True),
                    (u"Шевченко 341", False), (u"Каюа 5", False), (u"Каюа", False)):
        if is_ua_shop(nm) != exp:
            ebad.append(u"%s -> %s" % (nm, not exp))
    for shops, exp in (([u"Бучми 32 ЮА", u"Шевченко 341 ЮА"], "UA"),
                       ([u"Бучми 32", u"Шевченко 341"], "FM"),
                       ([u"Бучми 32", u"Шевченко 341 ЮА"], "MIX"), ([], None)):
        if entity_of_shops(shops) != exp:
            ebad.append(u"%s -> %s" % ("; ".join(shops), entity_of_shops(shops)))
    if route_key(u"Пр-т Ювілейний 67 ЮА") != route_key(u"Пр-т Ювілейний 67"):
        ebad.append(u"сверка куста не видит магазин ЮА")
    for nm, exp in ((u"Полевая-Склад", True), (u"Полевая-Склад ЮА", True),
                    (u"полевая-склад юа", True), (u"Полевая-Магазин", False),
                    (u"Полевая-Магазин ЮА", False), (u"Бучми 32", False)):
        if is_no_order_shop(nm) != exp:
            ebad.append(u"склад: %s -> %s" % (nm, not exp))
    for pth, exp in ((os.path.join(IN_FM_DIR, u"a.xlsx"), "FM"),
                     (os.path.join(IN_UA_DIR, u"a.xlsx"), "UA"),
                     (os.path.join(IN_DIR, u"a.xlsx"), None),
                     (os.path.join(IN_UA_DIR, u"x", u"a.xlsx"), None)):
        if folder_entity(pth) != exp:
            ebad.append(u"папка: %s -> %s" % (pth, folder_entity(pth)))
    if ebad:
        R.check("ERROR", "Тест юрлиц", "; ".join(ebad))
        return False
    R.check("OK", "Тест юрлиц", u"«... ЮА» = база ЮА Маркет, смешанная ведомость ловится, "
                                  u"склады не получатели, база по папке ведомости")
    return True

# ==================== ЧТЕНИЕ ИСТОЧНИКА =======================

ALIASES = {
    "name":   ["название товара", "назва товару", "наименование"],
    "art":    ["артикул"],
    "bc":     ["штрих-код", "штрих код", "штрихкод"],
    "start":  ["на начало", "на початок"],
    "inc":    ["приход", "прихід"],
    "move_in": ["[+] перемещение", "[+] переміщення"],
    "sale":   ["[-] реализация", "[-] реалізація", "реализация", "реалізація"],
    "out":    ["[-] перемещение", "[-] переміщення"],
    "woff":   ["списание", "списання"],
    "ret":    ["возврат поставщикам", "повернення постачальникам"],
    "end":    ["на конец", "на кінець"],
    "shop":   ["центр учета", "центр обліку", "магазин"],
}


def folder_entity(path):
    """FM / UA, если файл лежит в папке ведомостей своей базы, иначе None."""
    d = os.path.normcase(os.path.abspath(os.path.dirname(path)))
    for ent, folder in (("FM", IN_FM_DIR), ("UA", IN_UA_DIR)):
        if d == os.path.normcase(os.path.abspath(folder)):
            return ent
    return None


def _source_candidates():
    """Ведомости в папках баз, в прежней ВХОД и рядом с программой, свежие первыми.
    В папке базы берётся любой xlsx, в остальных - только с «ведомость» в имени."""
    found = []
    for folder in (IN_FM_DIR, IN_UA_DIR, IN_DIR, BASE_DIR):
        if not os.path.isdir(folder):
            continue
        own = folder in (IN_FM_DIR, IN_UA_DIR)
        for nm in os.listdir(folder):
            low = nm.lower()
            if nm.startswith("~$") or not low.endswith((".xlsx", ".xlsm")):
                continue
            if not own and not any(k in low for k in
                                   ("перемещ", "переміщ", "оборот", "ведом")):
                continue
            full = os.path.join(folder, nm)
            d = extract_date(full)
            found.append((d, os.path.getmtime(full), full))
    found.sort(key=lambda t: (t[0], t[1]), reverse=True)
    return found


def find_source_file(quiet=False):
    """Автопоиск: приоритет - дата в названии, затем дата изменения.
    Используется как запасной путь, когда окно выбора файла недоступно,
    и как подсказка по умолчанию для проводника (quiet=True)."""
    found = _source_candidates()
    if not found:
        return None
    if len(found) > 1 and not quiet:
        R.check("WARN", "В папке несколько ведомостей",
                "найдено %d, взят файл за %s: %s | остальные: %s"
                % (len(found), found[0][0].strftime("%d.%m.%Y"),
                   os.path.basename(found[0][2]),
                   ", ".join(os.path.basename(f[2]) for f in found[1:4])))
    return found[0][2]


def read_shops(path, nrows=None):
    """Магазины (Центр учета) ведомости. nrows - быстрый взгляд на начало файла."""
    try:
        df = pd.read_excel(path, sheet_name=0, header=detect_header_row(path) - 1,
                           nrows=nrows)
        col = resolve_columns(df).get("shop")
        if not col:
            return []
        v = df[col].dropna().astype(str).str.strip()
        return sorted(set(x for x in v if x and x.lower() != "nan"))
    except Exception:
        return []


def file_entity(path):
    """FM / UA / MIX / None по первым строкам ведомости - быстро, для окна."""
    return entity_of_shops(read_shops(path, nrows=400))


def find_sources_by_entity(quiet=False):
    """{"FM": путь, "UA": путь} - свежайшая ведомость каждой базы.
    База файла - по его папке («Ведомости Family» / «Ведомости ЮА Маркет»);
    в ВХОД и рядом с программой - по магазинам. Чужой файл в папке базы
    не перекладывается молча: его поймает сверка юрлица."""
    by = {}
    for _d, _mt, full in _source_candidates():
        ent = folder_entity(full) or ("UA" if file_entity(full) == "UA" else "FM")
        by.setdefault(ent, []).append(full)
    for ent, lst in by.items():
        if len(lst) > 1 and not quiet:
            R.check("WARN", u"В папке несколько ведомостей %s" % ENTITY_LABEL[ent],
                    u"найдено %d, взят %s | остальные: %s"
                    % (len(lst), os.path.basename(lst[0]),
                       ", ".join(os.path.basename(f) for f in lst[1:4])))
    return {ent: lst[0] for ent, lst in by.items()}


def extract_date(path):
    """Дата заказа. Тонкое место: она берётся из имени файла и ни на один
    расчёт не влияет - только на имя папки в ЗАКАЗЫ и подписи в отчётах.
    Ошибиться здесь = записать заказ не в ту папку, а не посчитать не то."""
    return guess_day(path)[0]


def guess_day(path):
    """(дата, откуда взята). Порядок: имя файла -> отметка выгрузки -> сегодня."""
    m = re.search(r"(\d{1,2})[.\-_ ](\d{1,2})[.\-_ ](\d{2,4})", os.path.basename(path))
    if m:
        d, mo, y = (int(g) for g in m.groups())
        y = y + 2000 if y < 100 else y
        try:
            return date(y, mo, d), "из имени файла"
        except ValueError:
            pass
    st = read_export_stamp(path)
    if st:
        return st.date(), "из отметки выгрузки: в имени файла даты нет"
    return date.today(), "сегодняшняя: в имени файла даты нет"


# ==================== МАРШРУТЫ ОТГРУЗКИ ======================
#
# Три куста, каждый отгружается дважды в неделю. Состав снят с фактических
# ведомостей 01-05 и 07.09.2026 (вт, ср, чт, пт, сб, пн) - это ровно одна
# полная неделя, поэтому адреса здесь в том виде, в каком их пишет Торгсофт,
# а не в старом русском написании из графика.
#   A - пн и чт   B - вт и пт   C - ср и сб    Воскресенье отгрузки нет.
# Зачем: если куст в файле не совпал с днём недели - выгрузили не тот период
# или ошиблись датой; если внутри куста магазина нет - он выпал из выгрузки.

ROUTES = {
    "A": [u"Амосова 5А", u"Богдана Хмельницького 8", u"Валентинівська 50А",
          u"Гарібальді 1", u"Зубенко 31В/5", u"Краснодарська 171з",
          u"Переяславська 23", u"Пр-т Героїв Харкова 160", u"Пр-т Ювілейний 67",
          u"Роганська 130/4", u"Роганська 148", u"Салтівське шосе 264В"],
    "B": [u"Іскрінський 19", u"Байрона 138/1", u"Байрона 156", u"Байрона 163",
          u"Грозненська 38", u"Зернова 6/5", u"Качанівська 19",
          u"Небесної Сотні 14/1", u"Ньютона 102", u"Ньютона 111",
          u"Олімпійська 9А", u"Петра Григоренка 37", u"Танкопія 16"],
    "C": [u"Астрономічна 44Г", u"Бучми 32", u"Бучми 52", u"Бучми Джерело",
          u"Валентинівська 24", u"Гвардійців Широнінців 54", u"Зубенко 23",
          u"Михайля Семенка 17", u"Нескорених 33",
          u"Пр-т Тракторобудiвникiв 95", u"Шевченко 341"],
}
ROUTE_BY_DOW = {0: "A", 1: "B", 2: "C", 3: "A", 4: "B", 5: "C"}   # пн..сб

# Точки, отсутствие которых в выгрузке - норма, а не пропажа.
# Убрать из ROUTES совсем, когда закроются окончательно.
ROUTE_PAUSED = {}   # напр. {u"Адрес": u"магазин закрывается"}

# Улицы переименованы, в старых выгрузках и графике встречаются прежние
# названия - это те же самые магазины.
ROUTE_ALIASES = {
    u"Героїв Праці 33":         u"Нескорених 33",
    u"Героев труда 33":         u"Нескорених 33",
    u"Героїв Сталінграда 138/1": u"Байрона 138/1",
    u"Героев Сталинграда 138":   u"Байрона 138/1",
    u"Героїв Сталінграда 156":   u"Байрона 156",
    u"Героев Сталинграда 156":   u"Байрона 156",
    u"Героїв Сталінграда 163":   u"Байрона 163",
    u"Героев Сталинграда 163":   u"Байрона 163",
    u"Болградська 38":           u"Грозненська 38",   # имя в базе ЮА Маркет
}
DOW_RU = [u"понедельник", u"вторник", u"среда", u"четверг",
          u"пятница", u"суббота", u"воскресенье"]

# латиница, набранная вместо кириллицы, встречается в адресах Торгсофта
_LAT2CYR = {"a": u"а", "c": u"с", "e": u"е", "i": u"і", "o": u"о",
            "p": u"р", "x": u"х", "y": u"у"}


def route_key(s):
    """Ключ сравнения адреса: регистр, пробелы, дроби и латинские двойники
    не должны мешать сопоставлению."""
    t = re.sub(r"\s+", " ", str(s).strip().lower())
    t = re.sub(r"\s%s$" % UA_SHOP_SUFFIX.lower(), "", t)   # «Бучми 32 ЮА» = «Бучми 32»
    t = "".join(_LAT2CYR.get(ch, ch) for ch in t)
    return re.sub(r"[/\\.-]", "", t)


def check_route(day, shops_in_file, other=()):
    """Сверка состава магазинов с графиком отгрузки.
    other - магазины ведомости второго юрлица на тот же день: их здесь
    не считаем, но и пропавшими не называем."""
    if dist_on():
        R.check("INFO", u"График отгрузки",
                u"раздача с %s: куст не проверяем, точек в файле %d"
                % (dist_source(), len(set(shops_in_file))))
        return ROUTE_BY_DOW.get(day.weekday()) or "A"
    alias = {route_key(k): route_key(v) for k, v in ROUTE_ALIASES.items()}
    have, own = {}, set()
    for x in list(other) + list(shops_in_file):
        k = route_key(x)
        have[alias.get(k, k)] = x
    for x in shops_in_file:
        k = route_key(x)
        own.add(alias.get(k, k))
    if other:
        R.check("INFO", u"Вторая ведомость",
                u"%d магазинов другого юрлица учтены при сверке куста, "
                u"считаются отдельно" % len(set(other)))
    scores = {k: len(have.keys() & {route_key(a) for a in v})
              for k, v in ROUTES.items()}
    best = max(scores, key=lambda k: scores[k])
    exp = ROUTE_BY_DOW.get(day.weekday())

    if exp is None:
        R.check("WARN", u"График отгрузки",
                u"%s - отгрузки нет, в файле похоже на куст %s"
                % (DOW_RU[day.weekday()], best))
        use = best
    elif best != exp and scores[best] > scores.get(exp, 0):
        R.check("WARN", u"График отгрузки",
                u"дата %s это %s (куст %s), а состав магазинов похож на куст %s"
                u" - проверьте дату или период выгрузки"
                % (day.strftime("%d.%m.%Y"), DOW_RU[day.weekday()], exp, best))
        use = best
    else:
        use = exp

    want = ROUTES[use]
    miss = [a for a in want if route_key(a) not in have]
    extra = [have[k] for k in have
             if k in own and k not in {route_key(a) for a in want}]

    paused = [a for a in miss if a in ROUTE_PAUSED]
    miss = [a for a in miss if a not in ROUTE_PAUSED]
    for a in paused:
        R.check("INFO", u"Точка вне отгрузки", u"%s - %s" % (a, ROUTE_PAUSED[a]))
    if miss:
        R.check("WARN", u"Магазина нет в выгрузке",
                u"куст %s, %s, не хватает %d из %d: %s"
                % (use, DOW_RU[day.weekday()], len(miss), len(want),
                   "; ".join(miss)))
    else:
        R.check("OK", u"График отгрузки",
                u"куст %s (%s), все %d магазинов на месте"
                % (use, DOW_RU[day.weekday()], len(want)))
    if extra:
        R.check("WARN", u"Магазин вне графика этого дня",
                u"%s - в кусте %s его быть не должно" % ("; ".join(extra), use))
    return use


# ============ ВЫБОР ФАЙЛА, ДАТЫ И КОНТРОЛЬ ПОВТОРОВ ============
#
# Задача блока - убрать три тихие ошибки:
#   1. посчитали не тот файл (старую выгрузку, вчерашнюю ведомость);
#   2. дата угадана из имени и заказ лёг не в ту папку;
#   3. повторный прогон подмешал свежие файлы к прошлым в той же папке.
# Ни одна из них не портит числа - все три портят то, что уходит в ТСД.

USE_GUI            = True   # False -> без окон, как раньше (автопоиск)
STALE_EXPORT_HOURS = 18     # выгрузка старше -> предупреждение
DATE_SANITY_DAYS   = 10     # дата дальше от сегодня -> предупреждение
HISTORY_FILE       = "история_запусков.csv"
REPLACED_DIR       = "_замененные"


def read_export_stamp(path):
    """Когда выгрузка сделана в Торгсофте. Берётся из служебных свойств xlsx
    (docProps/core.xml), а не из даты файла: копирование и пересохранение
    дату файла меняют, а эту отметку - нет. Возвращает datetime или None."""
    try:
        import zipfile
        with zipfile.ZipFile(path) as z:
            raw = z.read("docProps/core.xml").decode("utf-8", "ignore")
        m = (re.search(r"<dcterms:created[^>]*>([^<]+)<", raw)
             or re.search(r"<dcterms:modified[^>]*>([^<]+)<", raw))
        if not m:
            return None
        t = m.group(1).strip()
        dt = datetime.strptime(t[:19], "%Y-%m-%dT%H:%M:%S")
        if t.endswith("Z") or "+" in t[10:]:
            dt = dt.replace(tzinfo=timezone.utc).astimezone().replace(tzinfo=None)
        return dt
    except Exception:
        return None


def source_fingerprint(path):
    """Отпечаток содержимого: размер + sha1. Один и тот же файл под другим
    именем или после копирования даёт тот же отпечаток."""
    try:
        import hashlib
        h = hashlib.sha1()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                h.update(chunk)
        return "%d_%s" % (os.path.getsize(path), h.hexdigest()[:16])
    except Exception:
        return ""


def history_read():
    """Список прошлых запусков из ЛОГИ \\ история_запусков.csv."""
    rows, p = [], os.path.join(LOG_DIR, HISTORY_FILE)
    if not os.path.isfile(p):
        return rows
    try:
        with io.open(p, encoding="utf-8") as f:
            for line in f.read().splitlines()[1:]:
                parts = line.split(";")
                if len(parts) >= 5:
                    rows.append(dict(zip(("run", "day", "file", "fp", "shops",
                                          "base"), parts[:6])))
    except Exception:
        pass
    return rows


def history_write(day, src, fp, shops, base=u""):
    p = os.path.join(LOG_DIR, HISTORY_FILE)
    new = not os.path.isfile(p)
    try:
        with io.open(p, "a", encoding="utf-8", newline="") as f:
            if new:
                f.write(u"запуск;дата заказа;файл;отпечаток;магазинов;база\n")
            f.write(u"%s;%s;%s;%s;%d;%s\n"
                    % (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                       day.strftime("%Y-%m-%d"),
                       os.path.basename(src).replace(";", ","), fp, shops, base))
    except Exception:
        pass


def history_same_file(fp):
    """Последний запуск с тем же отпечатком или None."""
    if not fp:
        return None
    same = [h for h in history_read() if h.get("fp") == fp]
    return same[-1] if same else None


def day_dir_state(day, suffix=""):
    """(путь, магазинов внутри, время последней записи) для папки этой даты."""
    d = os.path.join(OUT_DIR, day.strftime("%Y-%m-%d") + suffix)
    if not os.path.isdir(d):
        return d, 0, None
    try:
        shops = [n for n in os.listdir(d)
                 if os.path.isdir(os.path.join(d, n)) and not n.startswith("_")]
        ts = datetime.fromtimestamp(os.path.getmtime(d))
    except Exception:
        shops, ts = [], None
    return d, len(shops), ts


def locked_files(day_dir):
    """Файлы в папке даты, открытые сейчас в Excel. Ловим их до записи,
    а не на середине: иначе половина магазинов запишется, половина нет."""
    busy = []
    if not os.path.isdir(day_dir):
        return busy
    for root_, _dirs, files in os.walk(day_dir):
        for nm in files:
            if nm.startswith("~$"):
                busy.append(os.path.join(root_, nm[2:]))
                continue
            if not nm.lower().endswith((".xlsx", ".txt")):
                continue
            p = os.path.join(root_, nm)
            try:
                with open(p, "r+b"):
                    pass
            except Exception:
                busy.append(p)
    return sorted(set(busy))


def clear_day_dir(day_dir):
    r"""Прошлый результат этой даты уводится в ЗАКАЗЫ\_замененные\<дата>_<время>.
    Не удаляется: вернуть можно всегда. Смысл - чтобы в папке даты не осталось
    txt магазина, который в новом расчёте заказа не получил."""
    import shutil
    if not os.path.isdir(day_dir) or not os.listdir(day_dir):
        return None
    dst = os.path.join(OUT_DIR, REPLACED_DIR,
                       "%s_%s" % (os.path.basename(day_dir),
                                  datetime.now().strftime("%H-%M-%S")))
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    shutil.move(day_dir, dst)
    return dst


def archive_source(src):
    """Ведомость после успешного расчёта уезжает в АРХИВ - чтобы в папке
    ведомостей всегда лежал ровно один файл и выбирать было не из чего."""
    import shutil
    try:
        dst = os.path.join(ARC_DIR, os.path.basename(src))
        k = 2
        while os.path.exists(dst):
            stem, ext = os.path.splitext(os.path.basename(src))
            dst = os.path.join(ARC_DIR, "%s (%d)%s" % (stem, k, ext))
            k += 1
        shutil.move(src, dst)
        return dst
    except Exception as e:
        R.check("WARN", "Перенос ведомости в АРХИВ",
                "не удался (%s), файл остался на месте" % type(e).__name__)
        return None


def _gui_pick(default_src):
    """Проводник + окно подтверждения. Возвращает кортеж настроек,
    None - если отменили, False - если окна недоступны (тогда автопоиск)."""
    import tkinter as tk
    from tkinter import filedialog, ttk

    root = tk.Tk()
    root.withdraw()
    try:
        root.call("tk", "scaling", 1.3)
    except Exception:
        pass
    # без этого диалог открывается ЗА окном консоли и выглядит как зависание
    try:
        root.attributes("-topmost", True)
    except Exception:
        pass

    start_dir = IN_FM_DIR if os.path.isdir(IN_FM_DIR) else BASE_DIR
    src = filedialog.askopenfilename(
        parent=root,
        title="Выберите оборотную ведомость из Торгсофта",
        initialdir=(os.path.dirname(default_src)
                    if default_src else start_dir),
        initialfile=(os.path.basename(default_src) if default_src else ""),
        filetypes=[("Книги Excel", "*.xlsx *.xlsm"), ("Все файлы", "*.*")])
    if not src:
        root.destroy()
        return None
    src = os.path.abspath(src)

    guess, origin = guess_day(src)
    stamp = read_export_stamp(src)
    prev  = history_same_file(source_fingerprint(src))

    state = {"ok": False}

    BG, CARD, FG   = "#f4f6f9", "#ffffff", "#16202c"
    MUTED, BORD    = "#6b7885", "#dfe4ea"
    ACC, ACC_H     = "#2f6fed", "#1f5bd0"
    RED, GRN, AMB  = "#c62828", "#1b7a3d", "#a86400"
    F, FB, FS      = ("Segoe UI", 10), ("Segoe UI Semibold", 10), ("Segoe UI", 9)

    win = tk.Toplevel(root)
    win.title(u"Расчёт заказа перемещений")
    win.configure(bg=BG)
    win.resizable(False, False)

    st = ttk.Style(win)
    try:
        st.theme_use("clam")
    except Exception:
        pass
    st.configure(".", background=CARD, foreground=FG, font=F)
    st.configure("Card.TFrame", background=CARD)
    st.configure("Cap.TLabel", background=CARD, foreground=MUTED, font=FS)
    st.configure("Val.TLabel", background=CARD, foreground=FG,
                 font=("Segoe UI Semibold", 11))
    st.configure("Note.TLabel", background=CARD, foreground=MUTED, font=FS)
    st.configure("Red.TLabel",  background=CARD, foreground=RED, font=FS)
    st.configure("Grn.TLabel",  background=CARD, foreground=GRN, font=FS)
    st.configure("Amb.TLabel",  background=CARD, foreground=AMB, font=FS)
    st.configure("Go.TButton", background=ACC, foreground="#ffffff",
                 font=FB, borderwidth=0, focuscolor=ACC, padding=(20, 9))
    st.map("Go.TButton", background=[("active", ACC_H), ("pressed", ACC_H)])
    st.configure("Ghost.TButton", background=CARD, foreground=MUTED,
                 font=F, borderwidth=0, focuscolor=CARD, padding=(16, 9))
    st.map("Ghost.TButton", background=[("active", "#e9edf3")])
    st.configure("TCheckbutton", background=CARD, foreground=FG, font=FS,
                 focuscolor=CARD)
    st.map("TCheckbutton", background=[("active", CARD)])
    st.configure("D.TEntry", fieldbackground="#ffffff", bordercolor=BORD,
                 lightcolor=BORD, darkcolor=BORD, padding=7)

    # шапка
    head = tk.Frame(win, bg=ACC)
    head.pack(fill="x")
    tk.Label(head, text=u"Расчёт заказа перемещений", bg=ACC, fg="#ffffff",
             font=("Segoe UI Semibold", 14)).pack(anchor="w", padx=22, pady=(16, 2))
    tk.Label(head, text=u"РЦ → магазины · проверьте перед запуском", bg=ACC,
             fg="#cfe0ff", font=FS).pack(anchor="w", padx=22, pady=(0, 16))

    body = tk.Frame(win, bg=CARD)
    body.pack(fill="both", expand=True)
    pad = dict(padx=22, anchor="w")

    def cap(t):
        tk.Label(body, text=t, bg=CARD, fg=MUTED, font=FS).pack(pady=(14, 1), **pad)

    def line(t, style="Note.TLabel"):
        w = ttk.Label(body, text=t, style=style, justify="left")
        w.pack(pady=(2, 0), **pad)
        return w

    def rule():
        tk.Frame(body, bg=BORD, height=1).pack(fill="x", padx=22, pady=(14, 0))

    cap(u"ВЕДОМОСТЬ")
    tk.Label(body, text=os.path.basename(src), bg=CARD, fg=FG,
             font=("Segoe UI Semibold", 11)).pack(**pad)
    line(os.path.dirname(src))

    if stamp:
        age = (datetime.now() - stamp).total_seconds() / 3600.0
        line(u"Выгружено из Торгсофта %s · %.0f ч назад"
             % (stamp.strftime("%d.%m.%Y в %H:%M"), age),
             "Amb.TLabel" if age > STALE_EXPORT_HOURS else "Grn.TLabel")
    else:
        line(u"Отметки времени выгрузки в файле нет")
    if prev:
        line(u"Этот же файл уже считался %s → ЗАКАЗЫ\\%s"
             % (prev["run"], prev["day"]), "Red.TLabel")

    rule()
    cap(u"ДАТА ЗАКАЗА")
    drow = tk.Frame(body, bg=CARD)
    drow.pack(padx=22, anchor="w")
    dv = tk.StringVar(value=guess.strftime("%d.%m.%Y"))
    ttk.Entry(drow, textvariable=dv, width=11, justify="center",
              style="D.TEntry", font=("Segoe UI Semibold", 12)).pack(side="left")
    tk.Label(drow, text="  " + origin, bg=CARD, fg=MUTED,
             font=FS).pack(side="left")
    route_l = line(u"")
    warn_l  = line(u"", "Red.TLabel")

    rule()
    box = tk.Frame(body, bg=CARD)
    box.pack(fill="x", padx=18, pady=(10, 0))
    clean_v, arch_v = tk.IntVar(value=1), tk.IntVar(value=1)
    chk = ttk.Checkbutton(box, variable=clean_v,
                          text=u"прошлый результат этой даты убрать в ЗАКАЗЫ\\%s"
                               % REPLACED_DIR)
    chk.pack(anchor="w", pady=1)
    ttk.Checkbutton(box, variable=arch_v,
                    text=u"после расчёта убрать ведомость в АРХИВ").pack(
                        anchor="w", pady=1)

    def parse_dv():
        t = dv.get().strip().replace("/", ".").replace("-", ".")
        m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})$", t)
        if not m:
            return None
        d, mo, y = (int(g) for g in m.groups())
        y = y + 2000 if y < 100 else y
        try:
            return date(y, mo, d)
        except ValueError:
            return None

    def refresh(*_a):
        d = parse_dv()
        if d is None:
            route_l.config(text=u"")
            warn_l.config(text=u"Дата не разобрана, нужен вид ДД.ММ.ГГГГ")
            chk.pack_forget()
            return
        rk = ROUTE_BY_DOW.get(d.weekday())
        if rk:
            route_l.config(text=u"%s · куст %s · ожидается %d магазинов"
                                % (DOW_RU[d.weekday()].capitalize(), rk,
                                   len(ROUTES[rk])), style="Note.TLabel")
        else:
            route_l.config(text=u"Воскресенье — отгрузки по графику нет",
                           style="Amb.TLabel")
        msgs = []
        _dir, n, ts = day_dir_state(d)
        if n:
            msgs.append(u"Папка ЗАКАЗЫ\\%s уже есть: %d магазинов, записана %s"
                        % (d.strftime("%Y-%m-%d"), n,
                           ts.strftime("%d.%m в %H:%M") if ts else "?"))
            chk.pack(anchor="w", pady=1, before=box.winfo_children()[-1])
        else:
            chk.pack_forget()
        off = (d - date.today()).days
        if abs(off) > DATE_SANITY_DAYS:
            msgs.append(u"Дата отстоит от сегодняшней на %d дн." % off)
        warn_l.config(text="\n".join(msgs))

    dv.trace_add("write", refresh)

    def on_ok(*_a):
        if parse_dv() is None:
            refresh()
            return
        state["ok"] = True
        win.destroy()

    def on_cancel(*_a):
        state["ok"] = False
        win.destroy()

    foot = tk.Frame(win, bg=BG)
    foot.pack(fill="x")
    ttk.Button(foot, text=u"Рассчитать", style="Go.TButton",
               command=on_ok).pack(side="right", padx=(0, 22), pady=16)
    ttk.Button(foot, text=u"Отмена", style="Ghost.TButton",
               command=on_cancel).pack(side="right", padx=8, pady=16)
    for w_ in foot.winfo_children():
        w_.configure(takefocus=1)
    tk.Frame(foot, bg=BG).pack(side="left")

    refresh()
    win.bind("<Return>", on_ok)
    win.bind("<Escape>", on_cancel)
    win.protocol("WM_DELETE_WINDOW", on_cancel)
    win.update_idletasks()
    sw, sh = win.winfo_screenwidth(), win.winfo_screenheight()
    win.geometry("%dx%d+%d+%d"
                 % (max(520, win.winfo_reqwidth()), win.winfo_reqheight(),
                    max(0, (sw - max(520, win.winfo_reqwidth())) // 2),
                    max(0, (sh - win.winfo_reqheight()) // 3)))
    win.grab_set()
    win.lift()
    win.focus_force()
    try:
        win.attributes("-topmost", True)
        win.after(400, lambda: win.attributes("-topmost", False))
    except Exception:
        pass
    root.wait_window(win)

    day = parse_dv() if state["ok"] else None
    clean = bool(clean_v.get())
    arch = bool(arch_v.get())
    root.destroy()
    if not state["ok"] or day is None:
        return None
    return src, day, clean, arch, ("указана вручную"
                                   if day != guess else origin)


GUI_CHOICE  = None  # заполняется окном перед запуском расчёта
GUI_SOURCES = None  # окно: [(юрлицо, путь)] - по ведомости на базу
JOB_CHOICE  = None  # выбор для текущего прогона (ставит main)
OTHER_SHOPS = []    # магазины ведомости второго юрлица - для сверки куста


def pick_source():
    """(файл, дата, чистить_папку, архивировать, откуда_дата) или None.
    Порядок: --auto -> автопоиск без окон; путь аргументом -> он;
    иначе проводник. Любая осечка окон = откат на прежнее поведение."""
    if JOB_CHOICE:
        return JOB_CHOICE
    if GUI_CHOICE:
        return GUI_CHOICE
    args = [a for a in sys.argv[1:] if a.strip()]
    auto = any(a.lower().lstrip("-/") in ("auto", "a") for a in args)
    paths = [a for a in args if not a.startswith(("-", "/"))]
    given = os.path.abspath(paths[0]) if paths and os.path.isfile(paths[0]) else None

    if not auto and USE_GUI:
        res = False
        try:
            res = _gui_pick(given or find_source_file(quiet=True))
        except Exception as e:
            R.check("WARN", "Окно выбора файла",
                    "не открылось (%s: %s), взят автопоиск"
                    % (type(e).__name__, e))
        if res is None:
            return None
        if res:
            return res

    src = given or find_source_file()
    if not src:
        return None
    day, origin = guess_day(src)
    return src, day, True, False, origin



def detect_header_row(path):
    probe = pd.read_excel(path, sheet_name=0, header=None, nrows=15)
    for idx in range(len(probe)):
        vals = [norm(v) for v in probe.iloc[idx].tolist() if pd.notna(v)]
        if any("штрих" in v for v in vals) and any("центр" in v or "магазин" in v for v in vals):
            return idx + 1
    return HEADER_ROW_HINT


def resolve_columns(df):
    cols = {norm(c): c for c in df.columns}
    res = {}
    for key, variants in ALIASES.items():
        hit = None
        for v in variants:
            if v in cols:
                hit = cols[v]
                break
        if hit is None:
            for v in variants:
                for cn, orig in cols.items():
                    if v in cn:
                        hit = orig
                        break
                if hit:
                    break
        res[key] = hit
    return res

# ======================= ПОДГОТОВКА ==========================

def prepare(df, c):
    num = lambda col: pd.to_numeric(df[col], errors="coerce")

    raw_sale, raw_out, raw_end = num(c["sale"]), num(c["out"]), num(c["end"])
    bad_cells = int(raw_sale.isna().sum() + raw_out.isna().sum() + raw_end.isna().sum())
    if bad_cells:
        R.check("WARN", "Нечисловые значения в расчётных столбцах",
                "%d ячеек трактованы как 0" % bad_cells)
    else:
        R.check("OK", "Числовые данные", "нечисловых ячеек в H/I/L нет")

    w = pd.DataFrame({
        "shop": df[c["shop"]].astype(str).str.strip(),
        "name": (df[c["name"]].astype(str).str.strip() if c.get("name") else ""),
        "art":  (df[c["art"]].astype(str).str.strip() if c.get("art") else ""),
        "bc":   df[c["bc"]].map(fmt_barcode),
        "sale": raw_sale.fillna(0.0),
        "out":  raw_out.fillna(0.0),
        "end":  raw_end.fillna(0.0),
        "start": (num(c["start"]).fillna(0.0) if c.get("start") else 0.0),
    })
    w["name"] = w["name"].replace({"nan": "", "None": ""})
    w["art"] = w["art"].replace({"nan": "", "None": ""})
    _fixed, _pairs = [], []
    for b in w["bc"]:
        nb, note = restore_barcode(b)
        _fixed.append(nb)
        if note:
            _pairs.append(note)
    w["bc"] = _fixed
    if _pairs:
        uniq_pairs = sorted(set(_pairs))
        R.check("WARN", "Восстановлены ведущие нули штрих-кодов",
                "%d кодов: %s" % (len(uniq_pairs), "; ".join(uniq_pairs[:6])))
        for p in uniq_pairs:
            R.problem("Ведущий ноль восстановлен (Excel обрезал)", "", p, "",
                      "проверь выгрузку Торгсофта")
    else:
        R.check("OK", "Ведущие нули штрих-кодов", "потерь не обнаружено")

    w["src_row"] = df.index + 1

    dropped = w[(w["bc"] == "") | (w["shop"] == "") | (w["shop"].str.lower() == "nan")]
    for _, r0 in dropped.iterrows():
        R.problem("Пропущена строка (нет штрих-кода или адреса)", r0["shop"], r0["bc"],
                  r0["name"], "строка источника ~%s" % r0["src_row"])
    if len(dropped):
        R.check("WARN", "Строки без штрих-кода/адреса", "исключено %d" % len(dropped))
    else:
        R.check("OK", "Полнота ключевых полей", "штрих-код и адрес есть во всех строках")

    w = w.drop(dropped.index)
    return w


def balance_check(df, c):
    need = ("start", "inc", "move_in", "sale", "out", "woff", "ret", "end")
    if any(not c.get(k) for k in need):
        R.check("WARN", "Сверка баланса ведомости", "нет части столбцов, проверка пропущена")
        return
    g = lambda k: pd.to_numeric(df[c[k]], errors="coerce").fillna(0.0)
    calc = g("start") + g("inc") + g("move_in") - g("sale") - g("out") - g("woff") - g("ret")
    diff = (calc - g("end")).abs()
    bad = int((diff > BALANCE_TOL).sum())
    if bad:
        R.check("WARN", "Сверка баланса (начало+приход+перемещ-расход = конец)",
                "расхождений: %d строк, макс. %.3f" % (bad, diff.max()))
    else:
        R.check("OK", "Сверка баланса ведомости", "все %d строк сходятся" % len(df))


def detect_weight(w):
    frac = lambda s: (s - s.round()).abs() > 0.001
    flag = frac(w["sale"]) | frac(w["out"]) | frac(w["end"]) | frac(w["start"])
    by_bc = flag.groupby(w["bc"]).any()
    kg_hint = (w.assign(h=w["name"].str.contains("кг", case=False, na=False))
                .groupby("bc")["h"].any())
    art_hint = (w.assign(h=w["art"].str.contains(r"[.,]", na=False))
                 .groupby("bc")["h"].any()) & kg_hint
    weight = (by_bc | art_hint).to_dict()
    R.check("INFO", "Определение весового товара",
            "весовых артикулов: %d из %d" % (sum(1 for v in weight.values() if v), len(weight)))
    return weight


def barcode_check(w):
    uniq = sorted(set(w["bc"]))
    nondigit = [b for b in uniq if not b.isdigit()]
    badlen = [b for b in uniq if b.isdigit() and len(b) not in (8, 12, 13, 14)]
    badsum = [b for b in uniq if b.isdigit() and len(b) in (8, 12, 13, 14) and not ean_valid(b)]
    if nondigit:
        R.check("ERROR", "Штрих-коды: посторонние символы",
                "%d кодов, напр.: %s" % (len(nondigit), ", ".join(nondigit[:5])))
    else:
        R.check("OK", "Штрих-коды: только цифры", "проверено %d кодов" % len(uniq))
    if badlen:
        R.check("WARN", "Штрих-коды: нетипичная длина",
                "%d кодов, напр.: %s" % (len(badlen), ", ".join(badlen[:5])))
    if badsum:
        R.check("WARN", "Штрих-коды: не сходится контрольная цифра",
                "%d кодов (норма для внутренних), напр.: %s" % (len(badsum), ", ".join(badsum[:5])))

    mix = w.groupby("bc")["name"].nunique()
    mix = mix[mix > 1]
    if len(mix):
        R.check("WARN", "Один штрих-код у разных названий", "%d кодов" % len(mix))
        for b in list(mix.index)[:20]:
            names = sorted(set(w.loc[w["bc"] == b, "name"]))
            R.problem("Один код - разные товары", "", b, " | ".join(names[:3]), "")
    return not bool(nondigit)

# ========================== РАСЧЁТ ===========================
# ================== КРАТНОСТЬ ИЗ НАЗВАНИЯ ====================

_MASK_PATTERNS = [
    r"\d+\s*(?:в|b|in)\s*1\b",                    # 3в1, 2в1, 3in1
    r"\d+(?:[.,]\d+)?\s*(?:мл|мм|см|%)",          # 175 мл, 250 мл
    r"№\s*\d+",                                   # Комплект №1
    r"\d+\s*сорт",                                # 1 сорт
]
_RE_STAR = re.compile(r"\(\s*[*xх]\s*(\d+)\s*\)", re.I)
_RE_UPAK = re.compile(r"(?:упаковка|уп\.?)\s*(\d+)\s*шт", re.I)
_RE_BOX  = re.compile(r"\(\s*(\d+(?:[.,]\d+)?)\s*\)")
_RE_KG   = re.compile(r"(\d+(?:[.,]\d+)?)\s*кг\b", re.I)
_RE_G    = re.compile(r"(\d+(?:[.,]\d+)?)\s*г\b", re.I)
_RE_STICK = re.compile(r"стик|\d\s*(?:в|in)\s*1", re.I)
_RE_COFFEE = re.compile(
    r"якобс|jacobs|мак\s*ко|maccoffee|mccoffee|петровськ|петровск|"
    r"нескафе|nescafe|карт\s*нуар|carte\s*noire|чорна\s*карта|галка|жокей", re.I)

STEP_FROM_NAME = []


def _fnum(s):
    try:
        return float(str(s).replace("\u00a0", "").replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def _mask_name(s):
    out = s
    for p in _MASK_PATTERNS:
        out = re.sub(p, lambda m: "\x00" * len(m.group(0)), out, flags=re.I)
    return out


def _no_brackets(s):
    return re.sub(r"\([^)]*\)", lambda m: "\x00" * len(m.group(0)), s)


def _small_gram(masked):
    mg = _RE_G.search(masked)
    v = _fnum(mg.group(1)) if mg else None
    return bool(v and v <= 30)


def guess_box(masked, artikul, is_weight):
    """Вложение ящика. Только для справки, в расчёт не идёт."""
    lim = MAX_KG_STEP if is_weight else MAX_BOX_PCS
    vals = [v for v in (_fnum(x) for x in _RE_BOX.findall(masked)) if v]
    for v in reversed(vals):
        if 0 < v <= lim:
            return v, "название"
    v = _fnum(artikul)
    if v and 0 < v <= lim:
        return v, "артикул"
    return None, ""


def parse_pack_from_name(name, artikul="", is_weight=False):
    """-> dict(step, thr, box, kind, note). step=None -> решает эвристика."""
    s = re.sub(r"\s+", " ", str(name or "").replace("\u00a0", " ")).strip()
    m = _mask_name(s)
    box, box_src = guess_box(m, artikul, is_weight)
    res = {"step": None, "thr": None, "box": box, "kind": "", "note": ""}

    if is_weight:
        mk = _RE_KG.search(_no_brackets(m))
        v = _fnum(mk.group(1)) if mk else None
        if v and abs(v - 1.0) < 1e-9:
            res["note"] = "цена за 1 кг, вес плавающий"
            return res
        if v and 0.1 <= v <= MAX_KG_STEP:
            res.update(step=v, thr=FIX_MIN_FRAC, kind="ves_fix",
                       note="фасовка %g кг из названия" % v)
        return res

    # штучный товар: стиковый кофе - только целым шоубоксом
    if _RE_STICK.search(s) or (_RE_COFFEE.search(s) and _small_gram(m)):
        if box and float(box).is_integer() and 2 <= box <= 100:
            res.update(step=float(int(box)), thr=PACK_MIN_FRAC, kind="sht_nedelimyy",
                       note="шоубокс %d шт (%s)" % (int(box), box_src))
            return res
    mu = _RE_UPAK.search(m)
    if mu:
        res.update(step=1.0, thr=frac_threshold(), kind="pack_unit",
                   note="SKU = упаковка %s шт, шаг 1" % mu.group(1))
        return res
    ms = _RE_STAR.search(m)
    if ms:
        res.update(step=1.0, thr=frac_threshold(), kind="piece",
                   note="вложение %s шт, шаг 1" % ms.group(1))
    return res



def find_credentials():
    """Ключ сервис-аккаунта: env -> папки-кандидаты. -> путь или None"""
    p = os.environ.get(CRED_ENV, "").strip().strip('"')
    if p and os.path.exists(p):
        return p
    for d in CRED_DIRS:
        if not d or not os.path.isdir(d):
            continue
        hits = sorted(glob.glob(os.path.join(d, CRED_MASK)))
        if not hits:
            hits = sorted(glob.glob(os.path.join(d, "*.json")))
        if hits:
            os.environ[CRED_ENV] = hits[0]
            return hits[0]
    return None


_MQ_PAREN = re.compile(r"\((\d{1,3})\)")


def _mq_in_name(mq, pname):
    """Подтверждён ли min_q числом в скобках названия: '... (25) 18г' + min_q=25."""
    if not pname:
        return False
    want = int(round(float(mq)))
    return any(int(g) == want for g in _MQ_PAREN.findall(str(pname)))


def _min_q_ok(info):
    """min_q достоверен только при статистике, вменяемой величине и подтверждении."""
    if is_box_only(info.get("pname", "")):
        return False  # (50)/(200)/(120) - приход на склад, не шаг перемещения
    mq = info.get("min_q")
    if not mq or float(mq) <= 0:
        return False
    if int(info.get("n") or 0) < MIN_Q_MIN_N:
        return False
    if float(mq) > MIN_Q_MAX:
        return False
    if str(info.get("confidence", "")) == "manual":
        return True
    ut = str(info.get("unit_type", ""))
    if ut == "ves_plav":
        return int(info.get("n") or 0) >= VES_MIN_Q_N   # порция подтверждена историей
    if ut.startswith("ves"):
        return False          # ves_fix: шаг и есть упаковка, min_q не нужен
    if ut == "sht_delimyy":
        return _mq_in_name(mq, info.get("pname", ""))
    return True


def _sane_min_q(info):
    return float(info["min_q"]) if _min_q_ok(info) else 0.0


def refresh_reference_from_bq():
    """Тянет справочник в REF_FILE. -> (успех, сообщение)"""
    if REF_OFFLINE:
        return False, "включён режим офлайн (REF_OFFLINE=True)"
    cred = find_credentials()
    if not cred:
        return False, ("не найден ключ: задай %s или положи json в %s"
                       % (CRED_ENV, os.path.join(_HERE, "credentials")))
    try:
        from google.cloud import bigquery
    except ImportError:
        return False, "нет пакета: pip install google-cloud-bigquery db-dtypes"
    sql = """
    SELECT p.barcode, p.pname,
           COALESCE(m.unit_type,  p.unit_type)  AS unit_type,
           COALESCE(m.order_step, p.order_step) AS order_step,
           COALESCE(m.min_q,      p.min_q)      AS min_q,
           COALESCE(m.step_supplier, p.step_supplier, p.order_step) AS step_supplier,
           p.box_kg, p.n, p.review,
           IF(m.barcode IS NULL, IFNULL(p.confidence, ''), 'manual') AS confidence
    FROM `%s` p
    LEFT JOIN `%s` m USING (barcode)
    """ % (BQ_TABLE, BQ_MANUAL)
    try:
        client = bigquery.Client(project=BQ_PROJECT)
        d = client.query(sql).to_dataframe()
    except Exception as e:
        return False, "%s: %s" % (type(e).__name__, str(e).replace("\n", " ")[:200])
    if d is None or d.empty:
        return False, "запрос вернул 0 строк"
    d["barcode"] = d["barcode"].astype(str).str.strip()
    os.makedirs(os.path.dirname(REF_FILE), exist_ok=True)
    d.to_csv(REF_FILE, index=False, encoding="utf-8-sig")
    return True, "%d позиций, ключ %s" % (len(d), os.path.basename(cred))


def load_reference():
    """Справочник упаковок: BigQuery с кэшем в REF_FILE."""
    fresh = (os.path.exists(REF_FILE)
             and (time.time() - os.path.getmtime(REF_FILE)) < REF_TTL_H * 3600)
    if fresh:
        R.check("INFO", "Справочник: источник",
                "кэш свежее %d ч, запрос в BigQuery не нужен" % REF_TTL_H)
    else:
        ok, msg = refresh_reference_from_bq()
        if ok:
            R.check("OK", "Справочник: обновлён из BigQuery", msg)
        elif os.path.exists(REF_FILE):
            age = (time.time() - os.path.getmtime(REF_FILE)) / 3600.0
            R.check("WARN", "Справочник: BigQuery недоступен",
                    "%s -> работаю по кэшу возрастом %.1f ч" % (msg, age))
        else:
            R.check("WARN", "Справочник: BigQuery недоступен",
                    "%s; кэша нет -> работа по эвристике" % msg)

    if not os.path.exists(REF_FILE):
        return {}

    try:
        d = pd.read_csv(REF_FILE, sep=None, engine="python",
                        encoding="utf-8-sig", dtype=str, keep_default_na=False)
    except Exception as e:
        R.check("ERROR", "Справочник упаковок",
                "не удалось прочитать %s: %s" % (REF_FILE, e))
        return {}

    d.columns = [str(cn).strip().lstrip("\ufeff").lower() for cn in d.columns]
    lack = [k for k in ("barcode", "unit_type", "order_step") if k not in d.columns]
    if lack:
        R.check("ERROR", "Справочник упаковок",
                "нет столбцов %s; в файле: %s" % (lack, list(d.columns)))
        return {}

    def _num(v):
        s = str(v).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
        if s == "" or s.lower() in ("nan", "none", "null"):
            return None
        try:
            x = float(s)
        except ValueError:
            return None
        return x if x > 0 else None

    out, broken, dups = {}, [], 0
    for _, r0 in d.iterrows():
        bc = fmt_barcode(r0["barcode"])
        if not bc:
            continue
        bc, _ = restore_barcode(bc)
        st = _num(r0["order_step"])
        if st is None:
            broken.append("%s '%s' -> '%s'" % (bc, str(r0.get("pname", ""))[:26],
                                               r0["order_step"]))
            continue
        if bc in out:
            dups += 1
        out[bc] = {
            "unit_type": str(r0["unit_type"]).strip(),
            "step": st,
            "min_q": _num(r0["min_q"]) if "min_q" in d.columns else None,
            "n": int(_num(r0.get("n", 0)) or 0),
            "pname": str(r0.get("pname", "")).strip(),
            "step_sup": _num(r0.get("step_supplier", "")) or st,
            "box_kg": _num(r0.get("box_kg", "")),
            "review": str(r0.get("review", "")).strip(),
            "confidence": str(r0.get("confidence", "")).strip(),
        }

    if broken:
        R.check("WARN", "Справочник: нечисловой шаг",
                "%d строк пропущено (похоже на порчу Excel - даты вместо чисел), напр.: %s"
                % (len(broken), "; ".join(broken[:3])))
    if dups:
        R.check("WARN", "Справочник: дубли штрих-кодов",
                "%d повторов, взята последняя строка" % dups)

    cnt = lambda t: sum(1 for v in out.values() if v["unit_type"] == t)
    mq_use = sum(1 for v in out.values() if _min_q_ok(v))
    mq_bad = [v for v in out.values() if v.get("min_q") and not _min_q_ok(v)]
    R.check("OK", "Справочник упаковок",
            "%d позиций: недел. %d, фасовка %d, вес %d; min_q применяется у %d"
            % (len(out), cnt("sht_nedelimyy"), cnt("ves_fix"), cnt("ves_plav"), mq_use))
    cand = [v for v in out.values()
            if v.get("min_q") and not _min_q_ok(v)
            and int(v.get("n") or 0) >= 20
            and v.get("unit_type") == "sht_delimyy"
            and float(v["min_q"]) >= 6]
    if cand:
        cand.sort(key=lambda v: -float(v["min_q"]))
        R.check("WARN", "Похоже на шоубокс, но не подтверждено названием",
                "%d позиций - проверь и внеси в transfer_pack_manual: %s"
                % (len(cand), "; ".join("%s min_q=%g n=%s"
                                        % (v.get("pname", "")[:30], v["min_q"], v.get("n"))
                                        for v in cand[:5])))
        for v in cand[:40]:
            R.problem("Возможный шоубокс без подтверждения", "", "", v.get("pname", ""),
                      "min_q=%g при n=%s; если это блок - задать order_step вручную"
                      % (v["min_q"], v.get("n")))

    if mq_bad:
        R.check("WARN", "Справочник: min_q отброшен как недостоверный",
                "%d позиций (наблюдений < %d или значение > %g), напр.: %s"
                % (len(mq_bad), MIN_Q_MIN_N, MIN_Q_MAX,
                   "; ".join("%s min_q=%g n=%s" % (v.get("pname", "")[:22],
                                                   v["min_q"], v.get("n"))
                             for v in mq_bad[:3])))
    rv = {}
    for v in out.values():
        r_ = v.get("review", "")
        if r_ and r_ != "OK":
            rv[r_] = rv.get(r_, 0) + 1
    if rv:
        R.check("INFO", "Справочник: позиции на проверку",
                ", ".join("%s: %d" % kv for kv in sorted(rv.items())))
    return out

def _norm_name(s: str) -> str:
    s = (s or "").lower()
    for a, b in (("і", "и"), ("ї", "и"), ("є", "е"), ("ы", "и")):
        s = s.replace(a, b)
    return s


def is_box_only(name: str) -> bool:
    """True, если скобки в названии = коробка прихода, а не шаг перемещения."""
    n = _norm_name(name)
    return any(p in n for p in BOX_ONLY_PATTERNS)


def is_pack_card(name: str) -> bool:
    """Карточка, единица которой сама по себе упаковка: «... Упаковка 100 шт».
    Кратность 50 к такой карточке применять нельзя - это было бы 50 упаковок."""
    return bool(re.search(r"упаковка\s*\d+\s*шт", _norm_name(name)))


def is_bag(name: str, bc: str = "") -> bool:
    """Упаковочный пакет, считаемый штуками. Карточка-упаковка сюда не входит."""
    if is_pack_card(name):
        return False
    if str(bc) in BAG_EXTRA_BC:
        return True
    return bool(re.match(r"\s*пакет\b", _norm_name(name)))


def transfer_block(name: str, bc: str = "") -> int:
    if is_bag(name, bc):
        return BAG_BLOCK
    n = _norm_name(name)
    for pat, blk in TRANSFER_BLOCK_RULES:
        if pat == "пакет":
            continue                      # пакеты разобраны выше, по началу названия
        if re.match(r"\s*" + pat, n):     # слово должно быть в начале названия
            return blk
    return 0


def box_from_name(name: str) -> float:
    """Вес ящика весового товара, если он написан в названии:
    «(4кг)», «(1,5кг в ящ)», «1кг (5)», «1 кг(0,8)». Ноль - не нашли.
    Названию верим больше, чем шагу из истории: история показывает,
    сколько отгружали по факту, а ящик физически не делится."""
    n = _norm_name(name).strip()
    for pat in (r"\((\d+[.,]?\d*)\s*кг[^)]*\)",
                r"кг\s*\((\d+[.,]?\d*)\)",
                r"\((\d+[.,]?\d*)\)\s*$"):
        m = re.search(pat, n)
        if m:
            try:
                v = float(m.group(1).replace(",", "."))
            except ValueError:
                continue
            if 0.2 <= v <= MAX_KG_STEP:
                return v
    return 0.0


def resolve_step(bc, heur_weight, name="", art=""):
    """Приоритет: справочник -> кратность из названия -> эвристика.
    -> (шаг, порог дробной части, минимальный заказ, пояснение)"""
    info = REF.get(bc)

    _nm = name or (info or {}).get("pname", "")
    _blk = transfer_block(_nm, bc)
    if _blk:
        _tag = "пакет" if is_bag(_nm, bc) else "блок"
        return float(_blk), ROUND_EPS, float(_blk), \
               "%s: кратно %d шт (коробка - только поставщику)" % (_tag, _blk)

    if info and info.get("step"):
        s = float(info["step"])
        ut = str(info.get("unit_type", ""))
        mq = _sane_min_q(info)                 # min_q при n < MIN_Q_MIN_N отбрасывается
        tail = (", min %g" % mq) if mq > 0 else ""
        rv = str(info.get("review", ""))
        if rv and rv != "OK":
            tail += " [%s]" % rv

        if ut == "sht_delimyy" and heur_weight and not WEIGHT_AS_INT:
            WEIGHT_CONFLICTS.append((bc, name, s))
            return WEIGHT_STEP, frac_threshold(), 0.0, \
                   "вес по ведомости (в справочнике 'шт')"
        if ut == "sht_nedelimyy":
            return s, FIX_MIN_FRAC, mq, "шоубокс %g шт%s" % (s, tail)
        if ut == "ves_fix":
            return s, FIX_MIN_FRAC, mq, "фасовка %g кг%s" % (s, tail)
        if ut == "ves_plav":
            return (s if s > 0 else WEIGHT_STEP), frac_threshold(), mq, \
                   "вес, шаг %g кг%s" % (s, tail)
        return s, frac_threshold(), mq, ((ut or "не указано") + tail)

    g = parse_pack_from_name(name, art, heur_weight)
    if g and g.get("step"):
        STEP_FROM_NAME.append((bc, name, float(g["step"]), g.get("box"), g.get("note", "")))
        return float(g["step"]), float(g["thr"]), 0.0, g.get("note", "из названия")

    s = WEIGHT_STEP if (heur_weight and not WEIGHT_AS_INT) else 1.0
    return s, frac_threshold(), 0.0, "эвристика: " + ("вес" if heur_weight else "шт")

def resolve_supplier_step(bc, transfer_step, unit_type=""):
    """Шаг для заказа поставщику: только целыми коробками.
    -> (шаг, пояснение)"""
    info = REF.get(bc)
    if not info:
        return float(transfer_step), "нет в справочнике"
    ut = str(info.get("unit_type", "") or unit_type)
    sup = float(info.get("step_sup") or info.get("step") or transfer_step)
    if ut in INDIVISIBLE:
        return sup, "неделимо, %g" % sup
    manual = str(info.get("confidence", "")) == "manual"
    if ut == "sht_delimyy" and sup > SUP_MAX_PCS_AUTO and not manual:
        return float(transfer_step), "кратность %g отклонена (транспортная коробка)" % sup
    if sup > float(transfer_step) + 1e-9:
        return sup, "коробка %g (перемещение по %g)" % (sup, transfer_step)
    return sup, "коробка неизвестна, по %g" % sup


def _dist_from_lines(lines):
    """Строки вида 'ШК' или 'ШК;количество' -> {ШК: сумма}."""
    tot = {}
    for ln in lines:
        p = [x.strip() for x in re.split(r"[;\t,]", ln) if x.strip()]
        if not p:
            continue
        q = _fnum(p[1]) if len(p) > 1 else None
        tot[p[0]] = tot.get(p[0], 0.0) + (q or 0.0)
    return tot


def _dist_read_txt(path):
    for enc in ("cp1251", "utf-8-sig"):
        try:
            with io.open(path, encoding=enc) as f:
                return f.read().splitlines()
        except (UnicodeDecodeError, IOError):
            continue
    return []


def _dist_from_txt_dir(folder):
    """Папка ВЫВОЗ: складываем TXT всех складов - это и приехало на Полевую."""
    tot, files = {}, 0
    for nm in sorted(os.listdir(folder)):
        if not nm.lower().endswith(".txt") or nm.upper().startswith(u"МАТРИЦА"):
            continue
        files += 1
        for bc, q in _dist_from_lines(_dist_read_txt(os.path.join(folder, nm))).items():
            tot[bc] = tot.get(bc, 0.0) + q
    return tot, files


def _dist_skip_row(bc, name, group):
    """Причина не перемещать строку вывоза, либо '' если можно."""
    if bc and bc in set(DIST_SKIP_BC):
        return u"ШК в списке исключений"
    low = (u"%s" % (name or u"")).lower()
    for pat in DIST_SKIP_NAME:
        if pat in low:
            return u"название: %s" % pat
    grp = (u"%s" % (group or u"")).strip().lower()
    for pat in DIST_SKIP_GROUP:
        if pat in grp:
            return u"группа: %s" % pat
    return u""


def _dist_pick_col(cols, *keys):
    """Колонка по подстроке имени: заголовки в файле вывоза плавают."""
    for c in cols:
        n = norm(c)
        for k in keys:
            if k in n:
                return c
    return None


DIST_NAMES = {}   # ШК -> название из файла вывоза (в ведомости его нет)


def _dist_skip_row(bc, name, group):
    """Причина не перемещать строку вывоза, либо '' если можно."""
    if bc and bc in set(DIST_SKIP_BC):
        return u"ШК в списке исключений"
    low = (u"%s" % (name or u"")).lower()
    for pat in DIST_SKIP_NAME:
        if pat in low:
            return u"название: %s" % pat
    grp = (u"%s" % (group or u"")).strip().lower()
    for pat in DIST_SKIP_GROUP:
        if pat in grp:
            return u"группа: %s" % pat
    return u""


def _dist_pick_col(cols, *keys):
    """Колонка по подстроке имени: заголовки в файле вывоза плавают."""
    for c in cols:
        n = norm(c)
        for k in keys:
            if k in n:
                return c
    return None


def _dist_from_xlsx(path):
    """Вывоз_<дата>.xlsx: строки с «Вывоз = ДА» после правок закупщицы.
    Розлив, кеги, кофе из аппарата и тех.товар отсекаются здесь же:
    позже нечем - у позиций без истории на точках названия в ведомости
    нет, и фильтр по названию в drop_excluded их не видит."""
    tot, sheets, cut = {}, 0, {}
    for sh in (u"Вне матрицы", u"Матрица на точках"):
        try:
            d = pd.read_excel(path, sheet_name=sh, dtype={u"Штрих-код": str})
        except (ValueError, IOError, KeyError):
            continue
        if u"Штрих-код" not in d.columns:
            continue
        if u"Вывоз" not in d.columns:
            R.check("WARN", u"Лист вывоза без колонки «Вывоз»",
                    u"«%s»: отметки закупщицы на этом листе не прочитаны" % sh)
            continue
        sheets += 1
        d = d[d[u"Вывоз"].astype(str).str.strip().str.upper() == u"ДА"]
        qs = pd.to_numeric(d[u"Количество"], errors="coerce").fillna(0)
        c_nm = _dist_pick_col(d.columns, u"назва", u"название", u"наименов")
        c_gr = _dist_pick_col(d.columns, u"групп")
        nms = d[c_nm].astype(str) if c_nm else pd.Series(u"", index=d.index)
        grs = d[c_gr].astype(str) if c_gr else pd.Series(u"", index=d.index)
        for bc, q, nm, gr in zip(d[u"Штрих-код"], qs, nms, grs):
            b = fmt_barcode(bc)
            nm = u"" if nm in (u"nan", u"None") else nm
            why = _dist_skip_row(b, nm, gr)
            if why:
                k = (b, nm, why)
                cut[k] = cut.get(k, 0.0) + float(q)
                continue
            if nm:
                DIST_NAMES[b] = nm
            tot[b] = tot.get(b, 0.0) + float(q)
    if cut:
        R.check("INFO", u"Из вывоза исключено до раздачи",
                u"%d ШК, %g ед. - розлив/кеги/кофе/тех.товар"
                % (len(cut), round(sum(cut.values()), 3)))
        for (b, nm, why), q in sorted(cut.items(), key=lambda t: -t[1])[:40]:
            R.problem(u"Не перемещаем (из файла вывоза)", "", b, nm,
                      u"%s, %g ед." % (why, round(q, 3)))
    return tot, sheets


def load_dist_list(path, quiet=False):
    """Ассортимент для раздачи. Принимается:
       * TXT/CSV - строки 'ШК' или 'ШК;количество';
       * папка ВЫВОЗ_<дата> (или её TXT) - сумма файлов всех складов;
       * Вывоз_<дата>.xlsx - строки с «Вывоз = ДА».
       -> (множество ШК, {ШК: сколько лежит на Полевой})"""
    raw, src = {}, os.path.basename(path.rstrip("\\/"))
    try:
        if os.path.isdir(path):
            sub = os.path.join(path, "TXT")
            folder = sub if os.path.isdir(sub) else path
            raw, n = _dist_from_txt_dir(folder)
            src = u"%s, файлов складов: %d" % (os.path.basename(folder), n)
        elif path.lower().endswith((".xlsx", ".xlsm")):
            raw, n = _dist_from_xlsx(path)
            src = u"%s, листов: %d" % (src, n)
        else:
            raw = _dist_from_lines(_dist_read_txt(path))
    except (IOError, OSError) as e:
        if not quiet:
            R.check("ERROR", u"Список раздачи", u"%s: %s" % (src, e))
        return set(), {}

    bcs, avail = set(), {}
    for bc0, q in raw.items():
        bc, _n = restore_barcode(fmt_barcode(bc0))
        if not bc.isdigit():
            continue
        bcs.add(bc)
        if q:
            avail[bc] = round(avail.get(bc, 0.0) + float(q), 6)
    if not quiet:
        R.check("OK" if bcs else "ERROR", u"Список раздачи",
                u"%s -> %d ШК, из них с количеством %d"
                % (src, len(bcs), len(avail)))
    return bcs, avail


def dist_count(path):
    """Сколько позиций в выбранном источнике - для подписи в окне."""
    try:
        return len(load_dist_list(path, quiet=True)[0])
    except Exception:
        return 0


def distribute_new(orders, w, weight_map, bcs, avail):
    """Разнарядка по позициям без истории на точках."""
    if DIST_NEW_MODE == "net" or w is None or not avail:
        return orders
    src = route_key(dist_source())
    shop_rows = w[w["shop"].map(lambda s: route_key(s) != src)]
    if shop_rows.empty:
        return orders
    have = set(orders["bc"]) if len(orders) else set()
# "Нет и не было на остатках" = ШК вообще отсутствует в ведомости либо
# за период по нему нулевое движение. Товар, который на точках есть и
# продаётся, разнарядкой не досылаем: у него просто закрыт спрос.
    _mv = w[["sale", "out", "start", "end"]].abs().sum(axis=1)
    seen_bc = set(w.loc[_mv > 1e-9, "bc"])
    skip_seen = sorted((bcs & seen_bc) - have)
    if skip_seen:
        R.check("INFO", u"Разнарядка: пропущены товары с историей",
                u"%d ШК есть на точках, но потребности нет - остаются на Полевой"
                % len(skip_seen))
        for b in skip_seen[:40]:
            R.problem(u"История есть, потребности нет", "", b, "",
                      u"разнарядкой не раздаётся")
    new_bc = [b for b in sorted(bcs)
              if b not in have and b not in seen_bc
              and float(avail.get(b, 0)) > 0]
    if not new_bc:
        return orders

    tw = shop_rows.groupby("shop")["sale"].sum()
    if DIST_NEW_MODE == "ravno" or float(tw.sum()) <= 0:
        tw = pd.Series(1.0, index=sorted(set(shop_rows["shop"])))
    tw = tw[tw > 0].sort_values(ascending=False)
    if DIST_NEW_TOP:
        tw = tw.head(DIST_NEW_TOP)
    if tw.empty:
        return orders
    share = tw / float(tw.sum())
    names = w.groupby("bc")["name"].first().to_dict()

    rows, spread = [], 0
    for bc in new_bc:
        total = float(avail[bc])
        is_w = bool((weight_map or {}).get(bc))
        step = WEIGHT_STEP if is_w else 1.0
        nmax = int(total // max(DIST_NEW_MIN, step))
        sh = share if nmax >= len(share) else share.head(max(nmax, 1))
        sh = sh / float(sh.sum())
        left = total
        for shop in sh.index:
            q = math.floor(total * float(sh[shop]) / step + 1e-9) * step
            q = round(min(q, left), 6)
            if q < DIST_NEW_MIN - 1e-9:
                continue
            left = round(left - q, 6)
            rows.append({"shop": shop, "bc": bc, "name": names.get(bc) or DIST_NAMES.get(bc, u""),
                         "qty": q, "raw": q, "step": step})
        if left > 1e-9:
            R.problem(u"Разнарядка: остаток на Полевой", "", bc,
                      names.get(bc, ""), u"не разошлось %g из %g" % (left, total))
        spread += 1

    if not rows:
        R.check("INFO", u"Разнарядка", u"делить нечего: доли меньше минимума")
        return orders
    add = pd.DataFrame(rows)
    for c in orders.columns:
        if c not in add.columns:
            add[c] = 0.0 if orders[c].dtype.kind in "fiu" else ""
    add = add[orders.columns]
    R.check("OK", u"Разнарядка по новым позициям",
            u"%d ШК без истории разложены на %d точек (%s), строк %d"
            % (spread, len(share), DIST_NEW_MODE, len(add)))
    return (pd.concat([orders, add], ignore_index=True)
              .sort_values(["shop", "name"], kind="stable"))


def drop_excluded(orders):
    """Убирает из заказа то, что перевозке не подлежит."""
    if orders.empty or (DIST_SKIP_ONLY_IN_DIST and not dist_on()):
        return orders
    bad = orders["bc"].isin(set(DIST_SKIP_BC))
    low = orders["name"].astype(str).str.lower()
    for pat in DIST_SKIP_NAME:
        bad = bad | low.str.contains(pat, regex=False, na=False)
    if bad.any():
        _skip = set(orders.loc[bad, "bc"])
        R.anomalies = [a for a in R.anomalies if a[1] not in _skip]
    if not bad.any():
        R.check("OK", u"Исключения из перемещения", u"под правило ничего не попало")
        return orders
    cut = orders[bad]
    R.check("INFO", u"Исключено из перемещения",
            u"%d строк, %g ед. - розлив/овощи/кеги/кофе не возим"
            % (len(cut), round(float(cut["qty"].sum()), 3)))
    for bc, g in cut.groupby("bc"):
        R.problem(u"Исключено: не перемещаем", "", bc, g.iloc[0]["name"],
                  u"точек %d, %g ед." % (len(g), round(float(g["qty"].sum()), 3)))
    return orders[~bad].copy()



# ============== ВЫВОЗ ПОДЧИСТУЮ (раздача с Полевой) ==============
DIST_SHIP_ALL = True


def _dist_banned(bc, name):
    """Кеги, розлив, овощи, кофе не раздаём даже разнарядкой."""
    if bc in DIST_SKIP_BC:
        return True
    n = norm(name or "")
    return any(k in n for k in DIST_SKIP_NAME) if n else False


def _dist_shares(w):
    """Доли точек по обороту. Источник исключён."""
    if w is None or not len(w):
        return None
    src = route_key(dist_source())
    rows = w[w["shop"].map(lambda s: route_key(s) != src)]
    if rows.empty:
        return None
    tw = rows.groupby("shop")["sale"].sum()
    if DIST_NEW_MODE == "ravno" or float(tw.sum()) <= 0:
        tw = pd.Series(1.0, index=sorted(set(rows["shop"])))
    tw = tw[tw > 0].sort_values(ascending=False)
    if DIST_NEW_TOP:
        tw = tw.head(DIST_NEW_TOP)
    return None if tw.empty else tw / float(tw.sum())


def _spread_qty(total, share, step, min_q):
    """Разложить total по точкам, сумма выдачи строго равна total."""
    total = round(float(total), 6)
    if total <= 0 or share is None or len(share) == 0:
        return {}
    units = int(math.floor(total / step + 1e-9))
    tail = round(total - units * step, 6)
    if units <= 0:
        return {share.index[0]: total}
    kmin = max(1, int(round(float(min_q) / step)))
    n = max(1, min(len(share), units // kmin))
    sh = share.head(n)
    sh = sh / float(sh.sum())
    raw = [units * float(sh[s]) for s in sh.index]
    base = [int(math.floor(x)) for x in raw]
    order = sorted(range(len(raw)), key=lambda i: raw[i] - base[i], reverse=True)
    for i in order[:units - sum(base)]:
        base[i] += 1
    give = {s: k for s, k in zip(sh.index, base) if k > 0}
    if not give:
        give = {sh.index[0]: units}
    if sum(give.values()) < units:
        top = max(give, key=lambda s: give[s])
        give[top] += units - sum(give.values())
    out = {s: round(k * step, 6) for s, k in give.items()}
    if tail > 1e-6:
        top = max(out, key=lambda s: out[s])
        out[top] = round(out[top] + tail, 6)
    return out


_BC_CACHE = {}


def _bc_shares(w, bc, share, st=None):
    """Доли точек для конкретного ШК: продажи этого ШК + общий оборот,
    без точек, где ШК лежит мёртвым. Нет продаж нигде -> общий оборот."""
    if share is None or w is None or not len(w):
        return share
    if not (DIST_SHARE_BY_BC or DIST_SKIP_DEAD):
        return share
    key = (id(w), len(w))
    if _BC_CACHE.get("key") != key:
        src = route_key(dist_source())
        rows = w[w["shop"].map(lambda s: route_key(s) != src)]
        _BC_CACHE.clear()
        _BC_CACHE["key"] = key
        _BC_CACHE["g"] = rows.groupby(["bc", "shop"])[["sale", "end"]].sum()
    g = _BC_CACHE["g"]
    try:
        sub = g.xs(bc, level="bc")
    except KeyError:
        if st is not None:
            st["shop"] += 1
        return share
    base = share.copy()
    if DIST_SKIP_DEAD:
        dead = [s for s in sub.index[(sub["sale"] <= 0) & (sub["end"] > 0)]
                if s in base.index]
        if dead and len(dead) < len(base):
            base = base.drop(dead)
            if st is not None:
                st["dead"] += len(dead)
    base = base / float(base.sum())
    if DIST_SHARE_BY_BC:
        s = sub["sale"].reindex(base.index).fillna(0.0).clip(lower=0.0)
        if float(s.sum()) > 0:
            k = float(DIST_BC_WEIGHT)
            mix = k * s / float(s.sum()) + (1.0 - k) * base
            mix = mix[mix > 0]
            if st is not None:
                st["bc"] += 1
            return (mix / float(mix.sum())).sort_values(ascending=False)
    if st is not None:
        st["shop"] += 1
    return base.sort_values(ascending=False)


def apply_dist_list(orders, w=None, weight_map=None):
    """Вывоз списка с Полевой подчистую."""
    if not dist_on():
        return orders
    if not DIST_SHIP_ALL:
        return apply_dist_list_legacy(orders, w, weight_map)
    bcs, avail = load_dist_list(DIST_FILE)
    if not bcs:
        return orders

    src = route_key(dist_source())
    was = len(orders) if orders is not None else 0
    cols = (list(orders.columns) if orders is not None and len(orders.columns)
            else ["shop", "bc", "name", "qty", "raw", "step", "is_weight"])
    if orders is None or orders.empty:
        o = pd.DataFrame(columns=cols)
    else:
        o = orders[orders["bc"].isin(bcs)].copy()
        o = o[o["shop"].map(lambda s: route_key(s) != src)]
    R.check("OK", u"Раздача по списку",
            u"потребность: из %d строк осталось %d, точек %d"
            % (was, len(o), o["shop"].nunique() if len(o) else 0))

    share = _dist_shares(w)
    names = {}
    if w is not None and len(w):
        names = {k: v for k, v in w.groupby("bc")["name"].first().to_dict().items() if v}
    try:
        for b, nm0 in DIST_NAMES.items():
            if nm0:
                names.setdefault(b, nm0)
    except NameError:
        pass

    keep, add, rest, ban = [], [], 0.0, 0
    st = {"need": 0, "extra": 0, "full": 0, "noqty": 0, "cut": 0}
    _bs = {"bc": 0, "shop": 0, "dead": 0}
    for bc in sorted(bcs):
        g = o[o["bc"] == bc]
        nm = names.get(bc, u"")
        if _dist_banned(bc, nm):
            ban += 1
            continue
        have = avail.get(bc)
        have = float(have) if have is not None else None
        is_w = bool((weight_map or {}).get(bc))
        step = WEIGHT_STEP if is_w else 1.0
        if have is None or have <= 0:
            if len(g):
                keep.append(g)
            else:
                st["noqty"] += 1
            continue

        left = round(have, 6)
        if len(g):
            g = g.sort_values("raw", ascending=False, kind="stable")
            rows = []
            for i, r0 in g.iterrows():
                if left <= 1e-9:
                    break
                q = float(r0["qty"])
                if q > left + 1e-9:
                    q = math.floor(left / step + 1e-9) * step
                    st["cut"] += 1
                q = round(min(q, left), 6)
                if q < step - 1e-9:
                    continue
                rows.append((i, q))
                left = round(left - q, 6)
            if rows:
                gg = g.loc[[i for i, _ in rows]].copy()
                gg["qty"] = [q for _, q in rows]
                keep.append(gg)
                st["need"] += len(rows)

        if left > 1e-9 and share is not None:
            for shop, q in _spread_qty(left, _bc_shares(w, bc, share, _bs), step, DIST_NEW_MIN).items():
                add.append({"shop": shop, "bc": bc, "name": nm, "qty": q,
                            "raw": q, "step": step, "is_weight": is_w})
                left = round(left - q, 6)
            st["extra" if len(g) else "full"] += 1
        if left > 1e-9:
            rest += left
            R.problem(u"Осталось на Полевой", "", bc, nm, u"не разошлось %g" % left)

    res = pd.concat(keep) if keep else o.iloc[0:0]
    if add:
        a = pd.DataFrame(add)
        for c in cols:
            if c not in a.columns:
                a[c] = 0.0 if (len(res) and res[c].dtype.kind in "fiu") else ""
        res = pd.concat([res, a[cols]], ignore_index=True)
    if len(res):
        res = res.reset_index(drop=True)
        if res.duplicated(["shop", "bc"]).any():
            agg = dict((c, ("sum" if c in ("qty", "raw") else "first"))
                       for c in res.columns if c not in ("shop", "bc"))
            res = res.groupby(["shop", "bc"], as_index=False, sort=False).agg(agg)[cols]
        res = res.sort_values(["shop", "name"], kind="stable")

    ship = res.groupby("bc")["qty"].sum() if len(res) else pd.Series(dtype=float)
    over = [b for b in ship.index
            if b in avail and float(ship[b]) > float(avail[b]) + 1e-6]
    R.check("ERROR" if over else "OK", u"ПЕРЕБОР по наличию на Полевой",
            (u"%d ШК больше остатка: %s" % (len(over), ", ".join(over[:5])))
            if over else u"перебора нет")
    R.check("INFO", u"Доли раздачи остатка",
            u"по продажам ШК (вес %g): %d ШК; по обороту точки: %d ШК; "
            u"пропущено точек, где ШК лежит без продаж: %d"
            % (DIST_BC_WEIGHT, _bs["bc"], _bs["shop"], _bs["dead"]))
    R.check("OK" if rest <= 1e-9 else "WARN", u"Вывоз подчистую",
            u"ШК %d, строк %d, единиц %g; потребностью %d строк, урезано %d, "
            u"доложено разнарядкой %d ШК, роздано целиком %d ШК; "
            u"не перемещаем %d ШК; без количества %d ШК; осталось %g ед."
            % (res["bc"].nunique() if len(res) else 0, len(res),
               round(float(res["qty"].sum()), 3) if len(res) else 0,
               st["need"], st["cut"], st["extra"], st["full"],
               ban, st["noqty"], round(rest, 3)))
    return res


def apply_dist_list_legacy(orders, w=None, weight_map=None):
    """Только ШК из списка, источник убираем, режем под наличие на Полевой.
    Дефицит отдаём точкам с наибольшей потребностью целыми строками:
    количества уже кратны упаковке, пересчитывать округление нельзя."""
    if not dist_on() or orders.empty:
        return orders
    bcs, avail = load_dist_list(DIST_FILE)
    if not bcs:
        return orders
    was = len(orders)
    src = route_key(dist_source())
    o = orders[orders["bc"].isin(bcs)].copy()
    o = o[o["shop"].map(lambda s: route_key(s) != src)]
    R.check("OK", u"Раздача по списку",
            u"из %d строк осталось %d, точек-получателей %d"
            % (was, len(o), o["shop"].nunique() if len(o) else 0))
    dead = sorted(bcs - set(o["bc"]))
    if dead:
        R.check("INFO", u"Из списка никому не нужны",
                u"%d ШК - потребности нет, остаются на Полевой" % len(dead))
        for b in dead[:MAX_REPORT_ROWS]:
            R.problem(u"Потребности нет ни на одной точке", "", b, "",
                      u"остаётся на Полевой")
    if not (DIST_LIMIT_STOCK and avail) or o.empty:
        return distribute_new(o, w, weight_map, bcs, avail)

    keep, cut = [], 0
    for bc, g in o.groupby("bc", sort=False):
        have = avail.get(bc)
        if have is None:
            keep.append(g)
            continue
        g = g.sort_values("raw", ascending=False, kind="stable")
        left, rows = float(have), []
        for i, r0 in g.iterrows():
            if float(r0["qty"]) <= left + ROUND_EPS:
                left -= float(r0["qty"])
                rows.append(i)
        if len(rows) < len(g):
            cut += len(g) - len(rows)
            R.problem(u"На Полевой не хватает", "", bc, g.iloc[0]["name"],
                      u"нужно %g, есть %g - без отгрузки %d точек"
                      % (round(float(g["qty"].sum()), 3), have, len(g) - len(rows)))
        keep.append(g.loc[rows])
    res = pd.concat(keep) if keep else o.iloc[0:0]
    R.check("OK" if not cut else "WARN", u"Ограничение по наличию на Полевой",
            u"отгружаем %d строк, не хватило на %d" % (len(res), cut))
    res = res.sort_values(["shop", "name"], kind="stable")
    return distribute_new(res, w, weight_map, bcs, avail)


def build_supplier_order(orders):
    """Сводит заказы магазинов в закупку у поставщика с округлением до коробки."""
    if orders.empty:
        return orders
    g = (orders.groupby("bc", as_index=False)
               .agg(name=("name", "first"), art=("art", "first"),
                    shops=("shop", "nunique"), need=("qty", "sum"),
                    step_tr=("step", "first"), ut=("unit_note", "first")))
    sup_step, sup_note, sup_qty = [], [], []
    for bc, need, st in zip(g["bc"], g["need"], g["step_tr"]):
        s, note = resolve_supplier_step(bc, st)
        q = apply_rounding_thr(need, s, SUPPLIER_MIN_FRAC)
        sup_step.append(s); sup_note.append(note); sup_qty.append(q)
    g["sup_step"], g["sup_note"], g["order"] = sup_step, sup_note, sup_qty
    g["izlishek"] = [round(o - n, 3) for o, n in zip(g["order"], g["need"])]

    split = int(sum(1 for a, b in zip(g["sup_step"], g["step_tr"]) if a > b + 1e-9))
    R.check("INFO", "Заказ поставщику",
            "%d позиций; из них %d вскрываются при перемещении (коробка > шага)"
            % (len(g), split))
    big = g[g["izlishek"] > 2 * g["sup_step"]]
    if len(big):
        R.check("WARN", "Заказ поставщику: крупный излишек округления",
                "%d позиций, см. лист заказа" % len(big))
    return g[g["order"] > 0].sort_values("name", kind="stable")


SUP_HDRS = ["№", "Название товара", "Артикул", "Штрих-код", "Магазинов",
            "Нужно магазинам", "Шаг перемещения", "Коробка", "ЗАКАЗ У ПОСТАВЩИКА",
            "Излишек", "Комментарий"]
SUP_W = [5, 46, 11, 15, 11, 16, 16, 10, 20, 11, 34]


def write_supplier_xlsx(path, day, g):
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказ поставщику"
    ws["A1"] = "Сводный заказ поставщику на %s" % day.strftime("%d.%m.%Y")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Потребность магазинов сведена и округлена вверх до целой коробки. "
                "Перемещение на точки допускает дробление.")
    ws["A2"].font = Font(italic=True, size=9)
    for j, h in enumerate(SUP_HDRS, start=1):
        c0 = ws.cell(row=4, column=j, value=h)
        c0.font = Font(bold=True); c0.fill = HEAD_FILL; c0.border = BORDER
        c0.alignment = Alignment(horizontal="center", wrap_text=True)
    r = 4
    for n, (_, row) in enumerate(g.iterrows(), start=1):
        r += 1
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value=row["name"])
        ws.cell(row=r, column=3, value=row["art"])
        bc = ws.cell(row=r, column=4, value=row["bc"]); bc.number_format = "@"
        ws.cell(row=r, column=5, value=int(row["shops"]))
        ws.cell(row=r, column=6, value=round(float(row["need"]), 3))
        ws.cell(row=r, column=7, value=float(row["step_tr"]))
        ws.cell(row=r, column=8, value=float(row["sup_step"]))
        q = ws.cell(row=r, column=9, value=round(float(row["order"]), 3))
        q.font = Font(bold=True)
        iz = ws.cell(row=r, column=10, value=float(row["izlishek"]))
        if float(row["izlishek"]) > 2 * float(row["sup_step"]):
            iz.fill = WARN_FILL
        ws.cell(row=r, column=11, value=row["sup_note"])
        for j in range(1, len(SUP_HDRS) + 1):
            ws.cell(row=r, column=j).border = BORDER
            if 5 <= j <= 10:
                ws.cell(row=r, column=j).number_format = "0.###"
    r += 1
    ws.cell(row=r, column=8, value="ИТОГО:").font = Font(bold=True)
    t = ws.cell(row=r, column=9, value=round(float(g["order"].sum()), 3))
    t.font = Font(bold=True); t.fill = WARN_FILL
    for j, wd in enumerate(SUP_W, start=1):
        ws.column_dimensions[get_column_letter(j)].width = wd
    ws.freeze_panes = "A5"
    wb.save(path)


def apply_rounding_thr(x, step, thr):
    if x is None or x <= 0:
        return 0.0
    u = round(float(x) / float(step), 9)
    fl = math.floor(u)
    n = fl if (u - fl) <= thr else fl + 1
    return round(n * float(step), 6)


def build_orders(w, weight_map):
    import numpy as np

    WEIGHT_CONFLICTS.clear()
    del STEP_FROM_NAME[:]

    h, i, l = w["sale"].to_numpy(float), w["out"].to_numpy(float), w["end"].to_numpy(float)
    _neg = -np.minimum(l, 0.0) if NEG_AS_NEED else np.zeros_like(l)
    vec = np.where(l <= 0, (h + i) * KOEF + _neg,
                   np.where(h > 0, np.maximum((h + i) * KOEF - l, 0.0), 0.0))
    if NEG_AS_NEED:
        _nm = l < 0
        if int(_nm.sum()):
            R.check("INFO", "Минусовой остаток включён в потребность",
                    "%d строк, добавлено %.3f ед. к расчёту"
                    % (int(_nm.sum()), float(-l[_nm].sum())))
    loop = np.array([calc_order(a, b, cc) for a, b, cc in zip(h, i, l)])
    ev = np.array([calc_order_eval(a, b, cc) for a, b, cc in zip(h, i, l)])

    d1 = float(np.max(np.abs(vec - loop))) if len(vec) else 0.0
    d2 = float(np.max(np.abs(vec - ev))) if len(vec) else 0.0
    if max(d1, d2) > 1e-9:
        R.check("ERROR", "Сверка трёх реализаций формулы",
                "макс. расхождение %.12f" % max(d1, d2))
    else:
        R.check("OK", "Сверка трёх реализаций формулы",
                "векторная = построчная = разбор текста, %d строк" % len(vec))

    w = w.copy()
    w["raw"] = vec

    grp = (w.groupby(["shop", "bc"], as_index=False)
            .agg(name=("name", "first"), art=("art", "first"),
                 sale=("sale", "sum"), out=("out", "sum"),
                 end=("end", "sum"), raw=("raw", "sum"), dup=("bc", "size")))

    dups = int((grp["dup"] > 1).sum())
    R.check("INFO", "Дубли товара внутри магазина",
            "сведено позиций: %d" % dups if dups else "дублей нет")

    grp["is_weight"] = grp["bc"].map(lambda b: bool(weight_map.get(b, False)))
    if DROP_BELOW > 0:
        grp.loc[grp["raw"] < DROP_BELOW, "raw"] = 0.0

    _s, _t, _m, _nt = [], [], [], []
    for _b, _hw, _nm, _ar in zip(grp["bc"], grp["is_weight"], grp["name"], grp["art"]):
        a, b2, m2, c2 = resolve_step(_b, _hw, _nm, _ar)
        _s.append(a); _t.append(b2); _m.append(m2); _nt.append(c2)
    grp["step"], grp["thr"], grp["min_q"], grp["unit_note"] = _s, _t, _m, _nt

    # весовой ящик: вес из названия важнее шага из истории
    boxes, steps2, n_nm, n_ref = [], [], 0, 0
    for _b, _hw, _nm2, _st in zip(grp["bc"], grp["is_weight"], grp["name"], grp["step"]):
        if not _hw:
            boxes.append(0.0)
            steps2.append(float(_st))
            continue
        _n2 = _norm_name(_nm2).strip()
        _por = next((v for k, v in VES_PORTION_BY_NAME.items() if _n2.startswith(k)), 0.0)
        if _por > 0:
            boxes.append(0.0)          # ящик не навязываем: есть подфасовка
            steps2.append(_por)
            continue
        nb = box_from_name(_nm2)
        if nb > 0:
            steps2.append(nb)          # вес из названия - всегда шаг, даже мелкий
        else:
            steps2.append(float(_st))
        bx = 0.0
        if nb >= VES_BOX_MIN:
            bx = nb
            n_nm += 1
        elif nb <= 0 and str(REF.get(_b, {}).get("unit_type", "")) == "ves_fix" \
                and float(_st) >= VES_BOX_MIN:
            bx = float(_st)
            n_ref += 1
        boxes.append(bx)
    grp["boxkg"] = boxes
    grp["step"] = steps2
    # ПАТЧ 1: порог обязан соответствовать шагу. Если шаг подменён весом ящика,
    # thr от весового режима (почти ноль) поднимает до двух ящиков, VES_BOX_RULE
    # срезает обратно до одного -> ложный ERROR "Округление не вверх".
    _thr2, _n_thr = [], 0
    for _bk2, _th2 in zip(grp["boxkg"], grp["thr"]):
        if float(_bk2 or 0.0) >= VES_BOX_MIN and float(_th2) <= 0.01:
            _thr2.append(FIX_MIN_FRAC)
            _n_thr += 1
        else:
            _thr2.append(float(_th2))
    grp["thr"] = _thr2
    if _n_thr:
        R.check("INFO", "Порог согласован с шагом ящика",
                "%d позиций: шаг = целый ящик, порог %s" % (_n_thr, FIX_MIN_FRAC))
    if n_nm or n_ref:
        R.check("INFO", "Весовой товар в ящиках",
                "ящик из названия: %d позиций, из справочника: %d" % (n_nm, n_ref))

    if WEIGHT_CONFLICTS:
        uniq = {c[0]: c for c in WEIGHT_CONFLICTS}
        R.check("WARN", "Справочник помечает весовой товар как штучный",
                "%d позиций переведены на шаг %g по данным ведомости"
                % (len(uniq), WEIGHT_STEP))
        for bc_, nm_, st_ in list(uniq.values())[:20]:
            R.problem("Тип в справочнике: 'шт', фактически вес", "", bc_, nm_,
                      "было %g -> стало %g; поправить unit_type в BigQuery" % (st_, WEIGHT_STEP))

    in_ref = int(grp["bc"].map(lambda b: b in REF).sum())
    R.check("INFO", "Покрытие справочником",
            "%d из %d позиций найдены в transfer_pack" % (in_ref, len(grp)))

    if STEP_FROM_NAME:
        _u = {}
        for _bx, _nx, _sx, _boxx, _ntx in STEP_FROM_NAME:
            _u[_bx] = (_nx, _sx, _ntx)
        R.check("WARN", "Шаг взят из названия (нет в справочнике)",
                "%d артикулов - проверь лист 'Проблемы данных'" % len(_u))
        for _bx, (_nx, _sx, _ntx) in list(_u.items())[:80]:
            R.problem("Шаг из названия (проверь)", "", _bx, _nx,
                      "шаг %g; %s" % (_sx, _ntx))

    _ex = grp["bc"].map(lambda b: REF.get(b, {}).get("unit_type") == "exclude")
    if int(_ex.sum()):
        R.check("INFO", "Исключено справочником", "%d позиций" % int(_ex.sum()))
        grp = grp[~_ex].copy()

    # 1) базовое округление по шагу и порогу
    grp["base"] = [apply_rounding_thr(x, s, t)
                   for x, s, t in zip(grp["raw"], grp["step"], grp["thr"])]

    # 2) нижняя граница: страховка при нулевом остатке + min_q из справочника
    floors, cuts, n_ins, n_mq, n_cut = [], [], 0, 0, 0
    for base, raw, end, step, thr, mq in zip(grp["base"], grp["raw"], grp["end"],
                                             grp["step"], grp["thr"], grp["min_q"]):
        fl, cut = 0.0, False
        if base <= 0 and raw > 0 and end <= 0 and thr > 0.01:
            fl = float(step)
            n_ins += 1
        if mq and mq > 0 and max(base, fl) > 0:
            mq_c = apply_rounding_thr(mq, step, ROUND_EPS)
            if max(base, fl) + 1e-9 < mq_c:
                if step <= WEIGHT_STEP + 1e-9 and raw < VES_MIN_FILL * mq_c:
                    cut = True          # мельче доли порции - в эту волну не везём
                    n_cut += 1
                else:
                    fl = mq_c
                    n_mq += 1
        floors.append(0.0 if cut else fl)
        cuts.append(cut)
    grp["floor"] = floors
    grp["qty"] = [0.0 if c else max(b, f)
                  for b, f, c in zip(grp["base"], grp["floor"], cuts)]

    # 2а) ящик не вскрывают: либо целый ящик, либо ноль.
    #     Остаток не больше VES_BOX_STOCK_MAX - везём ящик, иначе ждём.
    # ПАТЧ 2: помечаем позиции, количество которых задала политика ящика,
    # и разделяем подъём и срез - раньше срез считался как "поднято".
    grp["ves_box"] = False
    if VES_BOX_RULE:
        qs, box_hit, n_up, n_down, n_zero, n_wait = [], [], 0, 0, 0, 0
        wait_rows = []
        for q, bx, end_, raw_, shop_, bc_, nm_ in zip(
                grp["qty"], grp["boxkg"], grp["end"], grp["raw"],
                grp["shop"], grp["bc"], grp["name"]):
            bx = float(bx or 0.0)
            if bx < VES_BOX_MIN:
                qs.append(q)
                box_hit.append(False)
                continue
            if float(end_) <= VES_BOX_STOCK_MAX:
                if (VES_BOX_NEED_DEMAND and float(raw_) <= 1e-9
                        and float(end_) > VES_BOX_CRUMB):
                    n_wait += 1
                    wait_rows.append((shop_, bc_, nm_, float(end_)))
                    qs.append(0.0)
                    box_hit.append(True)
                    continue
                nq = bx * VES_BOX_QTY
                if nq > float(q) + 1e-9:
                    n_up += 1
                elif nq < float(q) - 1e-9:
                    n_down += 1
                qs.append(nq)
                box_hit.append(True)
            else:
                if float(q) > 0:
                    n_zero += 1
                qs.append(0.0)
                box_hit.append(True)
        grp["ves_box"] = box_hit
        if n_wait:
            R.check("INFO", "Весовой ящик: ждём спроса",
                    "%d позиций (продаж нет, хвост больше %g кг)"
                    % (n_wait, VES_BOX_CRUMB))
            for _s, _b, _n, _e in wait_rows[:40]:
                R.problem("Ящик не повезли: нет продаж", _s, _b, _n,
                          "остаток %.3f кг, продаж 0 - нужен ли товар в точке?" % _e)
        if n_down:
            R.check("WARN", "Весовой ящик: заказ срезан до целого ящика",
                    "%d позиций (больше %d ящика за волну не возим)"
                    % (n_down, VES_BOX_QTY))
        grp["qty"] = qs
        if n_up:
            R.check("INFO", "Весовой ящик: поднято до целого ящика",
                    "%d позиций (остаток не больше %g кг)" % (n_up, VES_BOX_STOCK_MAX))
        if n_zero:
            R.check("INFO", "Весовой ящик: потребность обнулена",
                    "%d позиций (остаток больше %g кг, ящик не вскрываем)"
                    % (n_zero, VES_BOX_STOCK_MAX))
    if n_cut:
        R.check("INFO", "Весовой: мелкая потребность отброшена",
                "%d позиций (меньше %g порции по истории перемещений)"
                % (n_cut, VES_MIN_FILL))

    if n_ins:
        R.check("INFO", "Страховка при нулевом остатке",
                "поднято до 1 упаковки: %d позиций (порог %s срезал бы в ноль)"
                % (n_ins, FIX_MIN_FRAC))
    if n_mq:
        R.check("INFO", "Поднято до min_q справочника", "%d позиций" % n_mq)

    _chk = grp[grp["step"] > 1.0]
    _bad = _chk[((_chk["qty"] / _chk["step"]) -
                 (_chk["qty"] / _chk["step"]).round()).abs() > 1e-6]
    if len(_bad):
        R.check("ERROR", "Кратность нарушена", "%d позиций" % len(_bad))
    else:
        R.check("OK", "Кратность соблюдена",
                "все заказы кратны шагу (%d позиций с шагом > 1)" % len(_chk))

    # контроль направления округления
    # для фасовки/шоубоксов недобор допустим, но не больше порога FIX_MIN_FRAC от упаковки
    # ПАТЧ 3: позиции, заданные правилом ящика, из проверки округления убраны:
    # их количество определяет политика отгрузки, а не формула потребности.
    free = grp[~grp["ves_box"].astype(bool)]
    plain = free[(free["qty"] > 0) & (free["thr"] <= 0.01)]
    wrong = plain[plain["qty"] + 1e-9 < plain["raw"]]

    packs = free[(free["qty"] > 0) & (free["thr"] > 0.01)]
    over = packs[packs["raw"] - packs["qty"] > packs["thr"] * packs["step"] + 1e-9]
    cut = int((packs["raw"] - packs["qty"] > 1e-9).sum())

    if len(wrong) or len(over):
        R.check("ERROR", "Округление не вверх",
                "штучных/весовых занижено %d, фасовок ниже допустимого %d"
                % (len(wrong), len(over)))
        for _, b0 in list(wrong.iterrows())[:10] + list(over.iterrows())[:10]:
            R.problem("Занижено против расчёта", b0["shop"], b0["bc"], b0["name"],
                      "расчёт %.3f -> заказ %.3f, шаг %g" % (b0["raw"], b0["qty"], b0["step"]))
    else:
        R.check("OK", "Округление вверх",
                "занижений нет; неполная упаковка отброшена по порогу %s: %d позиций"
                % (FIX_MIN_FRAC, cut))

    if BLOCK_MIN_FILL > 0:
        _blk = [float(transfer_block(nm, bc)) for nm, bc in zip(grp["name"], grp["bc"])]
        _bag = [bool(is_bag(nm, bc)) for nm, bc in zip(grp["name"], grp["bc"])]
        _blk = pd.Series(_blk, index=grp.index)
        _bag = pd.Series(_bag, index=grp.index)
        # пакеты не отбрасываем: 1 шт потребности -> всё равно везём 50
        _thin = (_blk > 0) & (~_bag if BAG_ALWAYS_UP else True) & (grp["raw"] < _blk * BLOCK_MIN_FILL)
        if int(_thin.sum()):
            R.check("INFO", "Блочный товар: мелкая потребность отброшена",
                    "%d позиций (меньше %g блока)" % (int(_thin.sum()), BLOCK_MIN_FILL))
            grp.loc[_thin, "qty"] = 0.0

    res = grp[grp["qty"] > 0].copy()
    zero = len(grp) - len(res)
    R.check("INFO", "Отброшено нулевых позиций", zero)

    for _, r0 in res.iterrows():
        why = []
        if r0["sale"] <= 0 and r0["qty"] > 0:
            why.append("заказ без реализации (остаток <= 0)")
        if r0["end"] < 0:
            _ut = str(REF.get(r0["bc"], {}).get("unit_type", ""))
            _piece = (float(r0["step"]).is_integer() and float(r0["step"]) >= 1
                      and not _ut.startswith("ves") and not bool(r0.get("is_weight")))
            if _piece or r0["end"] <= -NEG_END_TOL:
                why.append("отрицательный остаток %.3f%s"
                           % (r0["end"], " (штучный товар!)" if _piece else ""))
        if r0["qty"] > ANOMALY_ABS:
            why.append("очень крупный заказ")
        if r0["sale"] > 0 and r0["qty"] > ANOMALY_FACTOR * r0["sale"] and r0["qty"] >= 10:
            why.append("заказ > %dx реализации" % ANOMALY_FACTOR)
        if why:
            R.anomaly(r0["shop"], r0["bc"], r0["name"], r0["sale"], r0["out"],
                      r0["end"], r0["raw"], r0["qty"], "; ".join(why))
    if R.anomalies:
        R.check("WARN", "Аномальные позиции", "%d шт. - см. лист 'Аномалии'" % len(R.anomalies))
    else:
        R.check("OK", "Аномальные позиции", "не обнаружены")

    return res.sort_values(["shop", "name"], kind="stable")

def replace_cups_with_packs(res, w):
    """Стаканы в магазин везут упаковками. Поштучную позицию убираем,
    вместо неё ставим упаковку: количество = потребность в штуках / вместимость,
    округление вверх, кратность 1 упаковка."""
    if res is None or res.empty:
        return res
    import math as _m
    import pandas as _pd
    hit = res["bc"].isin(PACK_SWAP)
    if not bool(hit.any()):
        return res
    src = w.groupby(["shop", "bc"], as_index=False).agg(
        name=("name", "first"), art=("art", "first"),
        sale=("sale", "sum"), out=("out", "sum"), end=("end", "sum"))
    base = {c: 0.0 for c in ("sale", "out", "end", "raw", "base", "floor")}
    new_rows, moved = [], 0
    for _, r0 in res[hit].iterrows():
        pack_bc, per_pack = PACK_SWAP[r0["bc"]]
        need_pcs = float(r0["qty"])
        n_pack = max(1, int(_m.ceil(need_pcs / float(per_pack) - 1e-9)))
        moved += 1
        sel = (res["shop"] == r0["shop"]) & (res["bc"] == pack_bc)
        if bool(sel.any()):
            res.loc[sel, "qty"] = res.loc[sel, "qty"].astype(float) + n_pack
            continue
        info = REF.get(pack_bc) or {}
        row = dict(base)
        row.update(shop=r0["shop"], bc=pack_bc, qty=float(n_pack), step=1.0,
                   thr=ROUND_EPS, min_q=0.0, dup=1, is_weight=False,
                   unit_note="упаковка замість %g шт" % need_pcs)
        f = src[(src["shop"] == r0["shop"]) & (src["bc"] == pack_bc)]
        row["name"] = (f["name"].iloc[0] if len(f) else str(info.get("pname") or pack_bc))
        row["art"] = (f["art"].iloc[0] if len(f) else "")
        if len(f):
            row["sale"], row["out"], row["end"] = (float(f["sale"].iloc[0]),
                                                   float(f["out"].iloc[0]),
                                                   float(f["end"].iloc[0]))
        new_rows.append(row)
    res = res[~hit].copy()
    if new_rows:
        res = _pd.concat([res, _pd.DataFrame(new_rows)], ignore_index=True)
    R.check("INFO", "Поштучные карточки заменены упаковками",
            "%d позиций (стаканы и пакеты: закупщице нужны упаковки, не штуки)"
            % moved)
    return res.sort_values(["shop", "name"], kind="stable")


def add_bottle_accessories(res, w):
    """К пивным бутылкам добавляет крышки (1:1) и ручки (1:1 к бутылкам 1-2 л).
    Проверено по истории мая-августа: крышек на бутылку 1,02; ручек на бутылку 1-2 л 0,90."""
    if res is None or res.empty:
        return res
    import pandas as _pd
    base = {c: 0.0 for c in ("sale", "out", "end", "raw", "base", "floor")}
    src = w.groupby(["shop", "bc"], as_index=False).agg(
        name=("name", "first"), art=("art", "first"),
        sale=("sale", "sum"), out=("out", "sum"), end=("end", "sum"))
    added, raised = 0, 0
    new_rows = []
    for shop, part in res.groupby("shop", sort=False):
        got = dict(zip(part["bc"], part["qty"]))
        need = {
            BC_KRISHKA: sum(v for b, v in got.items() if b in BC_PLYASHKI_ALL),
            BC_RUCHKA:  sum(v for b, v in got.items() if b in BC_PLYASHKI_RUCHKA),
        }
        for bc, n in need.items():
            if n <= 0:
                continue
            info = REF.get(bc) or {}
            step = float(info.get("step") or 1.0) or 1.0
            n_up = apply_rounding_thr(n, step, ROUND_EPS)
            if bc in got:
                if got[bc] + 1e-9 < n_up:
                    res.loc[(res["shop"] == shop) & (res["bc"] == bc), "qty"] = n_up
                    raised += 1
                continue
            row = dict(base)
            row.update(shop=shop, bc=bc, qty=n_up, step=step, thr=ROUND_EPS,
                       min_q=0.0, dup=1, is_weight=False,
                       unit_note="комплект к пивній пляшці")
            f = src[(src["shop"] == shop) & (src["bc"] == bc)]
            row["name"] = (f["name"].iloc[0] if len(f)
                           else str(info.get("pname") or bc))
            row["art"] = (f["art"].iloc[0] if len(f) else "")
            if len(f):
                row["sale"], row["out"], row["end"] = (float(f["sale"].iloc[0]),
                                                       float(f["out"].iloc[0]),
                                                       float(f["end"].iloc[0]))
            new_rows.append(row); added += 1
    if new_rows:
        res = _pd.concat([res, _pd.DataFrame(new_rows)], ignore_index=True)
    if added or raised:
        R.check("INFO", "Комплект к пивным бутылкам",
                "добавлено позиций: %d, поднято до нормы: %d (крышка 1:1, ручка к 1-2 л)"
                % (added, raised))
    return res.sort_values(["shop", "name"], kind="stable")


# ========================== ЗАПИСЬ ===========================

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
ERR_FILL = PatternFill("solid", fgColor="FCE4E4")

HDRS = ["№", "Название товара", "Артикул", "Штрих-код", "[-] реализация",
        "[-] перемещение", "на конец", "Расчёт (скрипт)", "Расчёт (формула Excel)",
        "ЗАКАЗ", "Контроль", "Упаковка"]
WIDTHS = [5, 44, 11, 15, 13, 14, 10, 15, 17, 10, 14, 24]
LAST_COL = len(HDRS)          # 12 -> L


def excel_qty_formula(row_i, step, thr, floor=0.0):
    """Excel-формула округления, идентичная логике скрипта (с порогом строки и минимумом)."""
    s = repr(float(step))
    t = repr(float(thr))
    u = "I{r}/{s}".format(r=row_i, s=s)
    f = "IF({u}-INT({u})<={t},INT({u}),INT({u})+1)*{s}".format(u=u, t=t, s=s)
    if floor and float(floor) > 0:
        f = "MAX({f},{fl})".format(f=f, fl=repr(float(floor)))
    return f


def write_shop_xlsx(path, shop, day, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказ"
    ws["A1"] = "Заказ: %s" % shop
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Дата ведомости: %s     Позиций: %d     Единиц: %s     Формула: %s     "
                "Округление: %s, порог фасовки %s"
                % (day.strftime("%d.%m.%Y"), len(rows),
                   round(float(rows["qty"].sum()), 3),
                   FORMULA_TEXT, ROUND_MODE, FIX_MIN_FRAC))
    ws["A2"].font = Font(italic=True, size=9)

    hr = 4
    for j, h in enumerate(HDRS, start=1):
        c0 = ws.cell(row=hr, column=j, value=h)
        c0.font = Font(bold=True)
        c0.fill = HEAD_FILL
        c0.border = BORDER
        c0.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    r = hr
    for n, (_, row) in enumerate(rows.iterrows(), start=1):
        r += 1
        step = float(row["step"])
        thr = float(row["thr"])
        floor = float(row.get("floor", 0.0) or 0.0)

        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value=row["name"])
        ws.cell(row=r, column=3, value=row["art"])
        bc = ws.cell(row=r, column=4, value=row["bc"])
        bc.number_format = "@"
        ws.cell(row=r, column=5, value=float(row["sale"]))
        ws.cell(row=r, column=6, value=float(row["out"]))
        ws.cell(row=r, column=7, value=float(row["end"]))
        ws.cell(row=r, column=8, value=round(float(row["raw"]), 4))
        ws.cell(row=r, column=9, value=(
            ("=IF(G{r}<=0,(E{r}+F{r})*{k}-G{r},IF(E{r}>0,MAX((E{r}+F{r})*{k}-G{r},0),0))" if NEG_AS_NEED else "=IF(G{r}<=0,(E{r}+F{r})*{k},IF(E{r}>0,MAX((E{r}+F{r})*{k}-G{r},0),0))")
        ).format(r=r, k=repr(KOEF)))
        q = ws.cell(row=r, column=10, value=round(float(row["qty"]), 3))
        q.font = Font(bold=True)
        ws.cell(row=r, column=11, value=(
            '=IF(AND(ABS(H{r}-I{r})<0.0005,ABS({f}-J{r})<0.0005),"OK","РАСХОЖДЕНИЕ")'
        ).format(r=r, f=excel_qty_formula(r, step, thr, floor)))
        ws.cell(row=r, column=12, value=row.get("unit_note", ""))

        for j in range(1, LAST_COL + 1):
            cell = ws.cell(row=r, column=j)
            cell.border = BORDER
            if 5 <= j <= 10:
                cell.number_format = "0.###"
                cell.alignment = Alignment(horizontal="center")

    last = r
    r += 1
    ws.cell(row=r, column=9, value="ИТОГО:").font = Font(bold=True)
    tot = ws.cell(row=r, column=10, value=round(float(rows["qty"].sum()), 3))
    tot.font = Font(bold=True)
    tot.fill = WARN_FILL

    ws["A3"] = ('=IF(COUNTIF(K5:K{n},"РАСХОЖДЕНИЕ")=0,"Excel-контроль: OK, расчёт скрипта '
                'совпал с формулой во всех строках","Excel-контроль: РАСХОЖДЕНИЙ "'
                '&COUNTIF(K5:K{n},"РАСХОЖДЕНИЕ"))').format(n=last)
    ws["A3"].font = Font(bold=True, color="1F6F3F")

    for j, wd in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = wd
    ws.freeze_panes = "A%d" % (hr + 1)
    ws.auto_filter.ref = "A%d:%s%d" % (hr, get_column_letter(LAST_COL), last)
    wb.save(path)


def write_shop_txt(path, rows):
    # PATCH-TXT-FRAC v3: формат числа задаёт САМО количество, а не step.
    # Дробное qty при целом step уходило в int(round(q)): 0.175 -> 0, 1.1 -> 1.
    lines = []
    for _, row in rows.iterrows():
        q = float(row["qty"])
        try:
            step = float(row["step"])
        except Exception:
            step = 1.0
        as_w = (abs(q - round(q)) > 1e-9
                or step <= 0 or not float(step).is_integer()
                or bool(row.get("is_weight", False)))
        lines.append("%s;%s" % (row["bc"], qty_to_text(q, as_w)))
    with open(path, "w", encoding=TXT_ENCODING, errors="replace", newline="") as f:
        f.write(TXT_NEWLINE.join(lines) + TXT_NEWLINE)
    return lines


# =================== ПРОВЕРКА ЗАПИСАННОГО =====================

TXT_LINE_RE = re.compile(r"^\d{6,14};\d+(?:[.,]\d+)?$")


def verify_written(shop_dir, fname, rows):
    """Обратное чтение xlsx и txt со сверкой с расчётом."""
    problems = []
    xlsx_path = os.path.join(shop_dir, fname + ".xlsx")
    txt_path = os.path.join(shop_dir, fname + ".txt")

    expected = {}
    for _, row in rows.iterrows():
        expected[row["bc"]] = round(float(row["qty"]), 3)

    wb = load_workbook(xlsx_path, data_only=False)
    ws = wb.active
    got = {}
    r = 5
    while ws.cell(row=r, column=4).value not in (None, ""):
        b = fmt_barcode(ws.cell(row=r, column=4).value)
        v = ws.cell(row=r, column=10).value
        got[b] = round(float(v), 3) if isinstance(v, (int, float)) else None
        r += 1
    wb.close()
    if got != expected:
        _diff = []
        for _b in sorted(set(expected) | set(got)):
            _e, _g = expected.get(_b), got.get(_b)
            if _e != _g:
                _diff.append("%s: расчёт %s / xlsx %s" % (_b, _e, _g))
        problems.append("xlsx: позиций %d из %d, расхождений %d -> %s"
                        % (len(got), len(expected), len(_diff), "; ".join(_diff[:8])))

    with open(txt_path, "r", encoding=TXT_ENCODING) as f:
        raw_lines = [ln.strip() for ln in f.read().splitlines() if ln.strip()]
    if len(raw_lines) != len(expected):
        problems.append("txt: строк %d вместо %d" % (len(raw_lines), len(expected)))
    tgot = {}
    for ln in raw_lines:
        if not TXT_LINE_RE.match(ln):
            problems.append("txt: неверный формат строки '%s'" % ln)
            continue
        b, q = ln.split(";")
        tgot[b] = round(float(q.replace(",", ".")), 3)
    if tgot != expected:
        _diff = []
        for _b in sorted(set(expected) | set(tgot)):
            _e, _g = expected.get(_b), tgot.get(_b)
            if _e != _g:
                _diff.append("%s: расчёт %s / txt %s" % (_b, _e, _g))
        problems.append("txt: расхождений %d -> %s"
                        % (len(_diff), "; ".join(_diff[:8])))
    if os.path.getsize(txt_path) == 0:
        problems.append("txt: файл пустой")
    return problems

# ======================= ОТЧЁТЫ ==============================

def write_summary(path, day, summary, src):
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = "Сводка заказов на %s" % day.strftime("%d.%m.%Y")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Источник: %s" % src
    ws["A2"].font = Font(italic=True, size=9)
    for j, h in enumerate(["№", "Адрес (Центр учета)", "Позиций", "Единиц",
                           "Проверка файлов", "Папка"], start=1):
        c0 = ws.cell(row=4, column=j, value=h)
        c0.font = Font(bold=True)
        c0.fill = HEAD_FILL
    for n, it in enumerate(summary, start=1):
        ws.cell(row=4 + n, column=1, value=n)
        ws.cell(row=4 + n, column=2, value=it["shop"])
        ws.cell(row=4 + n, column=3, value=it["positions"])
        ws.cell(row=4 + n, column=4, value=it["units"])
        st = ws.cell(row=4 + n, column=5, value=it["verify"])
        st.fill = OK_FILL if it["verify"] == "OK" else ERR_FILL
        ws.cell(row=4 + n, column=6, value=it["folder"])
    for j, wd in enumerate([5, 32, 10, 10, 34, 62], start=1):
        ws.column_dimensions[get_column_letter(j)].width = wd
    ws.freeze_panes = "A5"
    wb.save(path)


def write_validation(path, day, src, shops_all, shops_with_order):
    wb = Workbook()
    ws = wb.active
    ws.title = "Проверки"
    ws["A1"] = "Отчёт валидации заказа на %s" % day.strftime("%d.%m.%Y")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Источник: %s   |   запуск %s" % (src, datetime.now().strftime("%d.%m.%Y %H:%M"))
    ws["A2"].font = Font(italic=True, size=9)
    verdict = "ЕСТЬ ОШИБКИ" if R.errors else ("ЕСТЬ ПРЕДУПРЕЖДЕНИЯ" if R.warns else "ВСЁ ЧИСТО")
    v = ws["A3"]
    v.value = "ИТОГ: %s (ошибок %d, предупреждений %d)" % (verdict, len(R.errors), len(R.warns))
    v.font = Font(bold=True, size=12)
    v.fill = ERR_FILL if R.errors else (WARN_FILL if R.warns else OK_FILL)

    for j, h in enumerate(["Статус", "Проверка", "Детали"], start=1):
        c0 = ws.cell(row=5, column=j, value=h)
        c0.font = Font(bold=True)
        c0.fill = HEAD_FILL
    for n, (st, nm, det) in enumerate(R.checks, start=1):
        ws.cell(row=5 + n, column=1, value=st).fill = {
            "OK": OK_FILL, "WARN": WARN_FILL, "ERROR": ERR_FILL,
            "INFO": PatternFill("solid", fgColor="EDEDED")}[st]
        ws.cell(row=5 + n, column=2, value=nm)
        ws.cell(row=5 + n, column=3, value=det)
    for j, wd in enumerate([10, 52, 90], start=1):
        ws.column_dimensions[get_column_letter(j)].width = wd

    ws2 = wb.create_sheet("Проблемы данных")
    for j, h in enumerate(["Тип", "Адрес", "Штрих-код", "Товар", "Детали"], start=1):
        c0 = ws2.cell(row=1, column=j, value=h)
        c0.font = Font(bold=True)
        c0.fill = HEAD_FILL
    for n, row in enumerate(R.problems, start=1):
        for j, v2 in enumerate(row, start=1):
            ws2.cell(row=1 + n, column=j, value=v2)
    for j, wd in enumerate([38, 26, 16, 42, 40], start=1):
        ws2.column_dimensions[get_column_letter(j)].width = wd

    ws3 = wb.create_sheet("Аномалии")
    for j, h in enumerate(["Адрес", "Штрих-код", "Товар", "реализация", "перемещение",
                           "на конец", "Расчёт", "ЗАКАЗ", "Причина"], start=1):
        c0 = ws3.cell(row=1, column=j, value=h)
        c0.font = Font(bold=True)
        c0.fill = HEAD_FILL
    for n, row in enumerate(R.anomalies, start=1):
        for j, v2 in enumerate(row, start=1):
            ws3.cell(row=1 + n, column=j, value=v2)
    for j, wd in enumerate([26, 16, 42, 12, 13, 11, 11, 10, 44], start=1):
        ws3.column_dimensions[get_column_letter(j)].width = wd

    ws4 = wb.create_sheet("Адреса")
    ws4.cell(row=1, column=1, value="Адрес (Центр учета)").font = Font(bold=True)
    ws4.cell(row=1, column=2, value="Есть заказ").font = Font(bold=True)
    for n, s in enumerate(shops_all, start=1):
        ws4.cell(row=1 + n, column=1, value=s)
        ws4.cell(row=1 + n, column=2, value="да" if s in shops_with_order else "нет")
    ws4.column_dimensions["A"].width = 34
    wb.save(path)

# =========================== MAIN ============================

# ---- итог раздачи для окна (PATCH-DIST-SUM v1) ----
_PRICE_KEYS = ((u"себест", u"собівар"),      # себестоимость - приоритет
               (u"закуп",),
               (u"цена", u"ціна"),
               (u"сумма", u"сума"))          # сумма строки / количество


def _num_col(s):
    return pd.to_numeric(s.astype(str).str.replace(u"\u00a0", u"")
                          .str.replace(u" ", u"").str.replace(u",", u"."),
                         errors="coerce")


def _dist_price_map():
    """ШК -> цена за единицу из файла вывоза. -> (dict, откуда взято)"""
    res, used = {}, u""
    try:
        book = pd.read_excel(DIST_FILE, sheet_name=None, dtype=str)
    except Exception as e:
        return res, u"файл вывоза не прочитан (%s)" % type(e).__name__
    for sh, d in book.items():
        cols = {norm(c): c for c in d.columns}
        c_bc = next((cols[k] for k in cols if u"штрих" in k), None)
        if c_bc is None:
            continue
        c_pr, kind = None, None
        for n, keys in enumerate(_PRICE_KEYS):
            c_pr = next((cols[k] for k in cols if any(x in k for x in keys)), None)
            if c_pr is not None:
                kind = n
                break
        if c_pr is None:
            continue
        pr = _num_col(d[c_pr])
        if kind == len(_PRICE_KEYS) - 1:
            c_q = next((cols[k] for k in cols if u"колич" in k or u"кільк" in k), None)
            if c_q is None:
                continue
            qq = _num_col(d[c_q])
            pr = pr / qq.where(qq > 0)
        for b, p in zip(d[c_bc], pr):
            if pd.isna(p) or p <= 0:
                continue
            b = restore_barcode(fmt_barcode(b))[0]
            if b and b not in res:
                res[b] = float(p)
        if not used:
            used = u"лист '%s', колонка '%s'" % (sh, c_pr)
    return res, used


def _money(x):
    return (u"{:,.2f}".format(float(x))).replace(u",", u" ")


def dist_summary_text(orders):
    """Строка итога раздачи для окна + разбивка по точкам в лог."""
    try:
        if not dist_on() or orders is None or not len(orders):
            return u""
        o = orders[orders["qty"] > 0]
        units = round(float(o["qty"].sum()), 3)
        txt = (u"\n\nРаздача с %s: %s ед. · точек %d · строк %d · ШК %d"
               % (dist_source(), units, o["shop"].nunique(), len(o),
                  o["bc"].nunique()))
        prices, used = {}, u""   # цен нет - сумму не считаем
        R.say(u"-" * 70)
        R.say(u"Раздача с %s      : %s ед." % (dist_source(), units))
        if prices:
            pv = o["bc"].map(prices)
            o = o.assign(_m=o["qty"] * pv)
            money = float(o["_m"].sum())
            miss = int(pv.isna().sum())
            txt += u"\nСумма раздачи: %s грн" % _money(money)
            if miss:
                txt += u" (без цены %d строк)" % miss
            R.say(u"Сумма раздачи       : %s грн, цены: %s, без цены строк: %d"
                  % (_money(money), used, miss))
            per = o.groupby("shop")[["qty", "_m"]].sum().sort_values("_m", ascending=False)
            for s, r in per.iterrows():
                R.say(u"  %-26s %10s ед. %14s грн"
                      % (s, round(float(r["qty"]), 3), _money(r["_m"])))
        else:
            per = o.groupby("shop")["qty"].sum().sort_values(ascending=False)
            for s, q in per.items():
                R.say(u"  %-26s %10s ед." % (s, round(float(q), 3)))
        return txt
    except Exception as e:
        R.say(u"Итог раздачи не посчитан: %s: %s" % (type(e).__name__, e))
        return u""


def plan_jobs():
    """[(юрлицо или None, выбор для pick_source или None)] - по прогону на базу.
    (None, None) = прежний путь: один файл, pick_source выбирает сам."""
    if GUI_SOURCES:
        rest = tuple(GUI_CHOICE[1:])
        return [(e, (p,) + rest) for e, p in GUI_SOURCES]
    args = [a for a in sys.argv[1:] if a.strip()]
    auto = any(a.lower().lstrip("-/") in ("auto", "a") for a in args)
    if any(not a.startswith(("-", "/")) for a in args) or (USE_GUI and not auto):
        return [(None, None)]
    if dist_on():
        return [(None, None)]          # раздача - по одной ведомости, как раньше
    found = find_sources_by_entity()
    if not found:
        return [(None, None)]          # pick_source сам скажет, что файла нет
    day, origin = guess_day(found.get("FM") or found["UA"])
    return [(e, (found[e], day, True, False, origin))
            for e in ("FM", "UA") if e in found]


def main():
    """Ведомость каждой базы считается отдельно, Family первой."""
    global JOB_CHOICE, OTHER_SHOPS
    ensure_dirs()
    jobs = plan_jobs()
    done = []
    for n, (ent, choice) in enumerate(jobs):
        if n:
            R.checks, R.problems, R.anomalies = [], [], []
            R.say("")
        RUN_INFO.clear()
        JOB_CHOICE = choice
        OTHER_SHOPS = []
        if len(jobs) > 1:
            if done:
                OTHER_SHOPS = list(done[-1].get("shops_list", []))
            else:
                OTHER_SHOPS = read_shops(jobs[1][1][0])
        try:
            code = main_one(ent)
        finally:
            JOB_CHOICE, OTHER_SHOPS = None, []
        info = dict(RUN_INFO)
        info.update({"entity": ENTITY, "code": code, "dist": dist_on(),
                     "errors": len(R.errors), "warns": len(R.warns),
                     "err_text": [u"%s: %s" % (n, d) if d else n
                                  for _s, n, d in R.errors[:3]]})
        done.append(info)
    RUN_INFO.clear()
    RUN_INFO["jobs"] = done
    return max(i["code"] for i in done)


def main_one(entity=None):
    """Один прогон по одной ведомости. entity - юрлицо, если его задали
    (окно, автопоиск); None - определяется по магазинам файла."""
    global ENTITY
    R.say("=" * 70)
    R.say("РАСЧЁТ ЗАКАЗА С ВАЛИДАЦИЕЙ  |  %s" % datetime.now().strftime("%d.%m.%Y %H:%M:%S"))
    R.say("=" * 70)

    if not self_test():
        R.say("Самотестирование не пройдено, работа прервана.")
        return 2

    picked = pick_source()
    if not picked:
        R.check("ERROR", "Файл-источник",
                "не выбран (отмена) либо не найден - положите ведомость в %s или %s"
                % (IN_FM_DIR, IN_UA_DIR))
        return 2
    src, day, do_clean, do_archive, date_origin = picked
    ENTITY = entity or file_entity(src)
    ENTITY = ENTITY if ENTITY in ENTITY_LABEL else "FM"
    R.check("OK", "Файл-источник", src)
    R.check("INFO", "Юрлицо (база Торгсофта)", ENTITY_LABEL[ENTITY])
    R.check("INFO", "Дата заказа", "%s (%s)" % (day.strftime("%d.%m.%Y"), date_origin))

    stamp = read_export_stamp(src)
    if stamp:
        age = (datetime.now() - stamp).total_seconds() / 3600.0
        R.check("WARN" if age > STALE_EXPORT_HOURS else "OK",
                "Выгрузка из Торгсофта",
                "%s, это %.0f ч назад" % (stamp.strftime("%d.%m.%Y %H:%M"), age))
    else:
        R.check("INFO", "Выгрузка из Торгсофта", "отметки времени в файле нет")

    fp = source_fingerprint(src)
    prev = history_same_file(fp)
    if prev:
        R.check("WARN", "Повторный расчёт",
                "тот же файл уже считался %s в папку %s" % (prev["run"], prev["day"]))

    off = (day - date.today()).days
    if abs(off) > DATE_SANITY_DAYS:
        R.check("WARN", "Дата заказа",
                "отстоит от сегодняшней на %d дн., проверьте имя файла" % off)

    busy = locked_files(os.path.join(OUT_DIR, day.strftime("%Y-%m-%d") + entity_suffix()))
    if busy:
        R.check("ERROR", "Файлы прошлого расчёта заняты",
                "закройте в Excel и запустите снова: %s"
                % "; ".join(os.path.basename(b) for b in busy[:5]))
        return 2

    hr = detect_header_row(src)
    R.check("OK" if hr == HEADER_ROW_HINT else "WARN", "Строка заголовков", hr)

    df = pd.read_excel(src, sheet_name=0, header=hr - 1).dropna(how="all")
    R.check("INFO", "Прочитано строк", len(df))

    c = resolve_columns(df)
    missing = [k for k in ("bc", "sale", "out", "end", "shop") if not c.get(k)]
    if missing:
        R.check("ERROR", "Обязательные столбцы",
                "не найдены: %s; заголовки файла: %s" % (missing, list(df.columns)))
        return 2
    R.check("OK", "Сопоставление столбцов",
            "H='%s', I='%s', L='%s', адрес='%s'" % (c["sale"], c["out"], c["end"], c["shop"]))

    balance_check(df, c)
    w = drop_no_order_shops(prepare(df, c))
    if not barcode_check(w):
        return 2
    weight_map = detect_weight(w)

    shops_all = sorted(set(w["shop"]))
    R.check("INFO", "Уникальных адресов", "%d: %s" % (len(shops_all), "; ".join(shops_all)))
    RUN_INFO["shops_list"] = shops_all
    got = entity_of_shops(shops_all)
    if got == "MIX":
        R.check("ERROR", "Юрлицо",
                u"в одной ведомости магазины двух баз (с «%s» в конце: %s). "
                u"Выгрузите ведомость каждой базы отдельно"
                % (UA_SHOP_SUFFIX, "; ".join(x for x in shops_all if is_ua_shop(x))))
    elif got and got != ENTITY:
        R.check("ERROR", "Юрлицо",
                u"ведомость выбрана как %s, а магазины в ней %s - проверьте, тот ли файл"
                % (ENTITY_LABEL[ENTITY], ENTITY_LABEL[got]))
    else:
        R.check("OK", "Юрлицо", u"все %d магазинов из базы %s"
                % (len(shops_all), ENTITY_LABEL[ENTITY]))
    check_route(day, shops_all, OTHER_SHOPS)

    names, collide = {}, []
    for s in shops_all:
        names.setdefault(safe_name(s).lower(), []).append(s)
    for k, v in names.items():
        if len(v) > 1:
            collide.append(" / ".join(v))
    if collide:
        R.check("WARN", "Совпадение имён папок после очистки",
                "; ".join(collide) + " - будут суффиксы (2), (3)")
    else:
        R.check("OK", "Имена файлов/папок", "коллизий нет")

    REF.clear()
    REF.update(load_reference())

    orders = build_orders(w, weight_map)
    orders = replace_cups_with_packs(orders, w)
    orders = add_bottle_accessories(orders, w)
    orders = drop_excluded(orders)
    if R.errors:
        R.say("Обнаружены критические ошибки, файлы не записаны.")
        write_validation(os.path.join(LOG_DIR, "ВАЛИДАЦИЯ_СБОЙ_%s.xlsx"
                                      % datetime.now().strftime("%Y-%m-%d_%H-%M-%S")),
                         day, src, shops_all, set())
        return 2
    if orders.empty:
        R.check("WARN", "Результат", "по формуле ни одна позиция не дала количество > 0")
        return 0

    orders = apply_dist_list(orders, w, weight_map)
    if dist_on():
        R.check("OK", u"Режим", u"раздача с %s по списку %s"
                % (dist_source(), os.path.basename(DIST_FILE)))
        if orders.empty:
            R.check("WARN", u"Раздача", u"ни одной точке ничего не нужно")
    day_dir = os.path.join(OUT_DIR, day.strftime("%Y-%m-%d") + dist_suffix()
                           + entity_suffix())
    if do_clean:
        moved = clear_day_dir(day_dir)
        if moved:
            R.check("INFO", "Прошлый результат этой даты",
                    "убран в %s" % moved)
    os.makedirs(day_dir, exist_ok=True)

    R.say("-" * 70)
    summary, used, verify_fail = [], {}, 0
    for shop, rows in orders.groupby("shop", sort=True):
        base_nm = safe_name(shop)
        nm, k = base_nm, 2
        while nm.lower() in used:
            nm = "%s (%d)" % (base_nm, k)
            k += 1
        used[nm.lower()] = shop

        shop_dir = os.path.join(day_dir, nm)
        os.makedirs(shop_dir, exist_ok=True)
        write_shop_xlsx(os.path.join(shop_dir, nm + ".xlsx"), shop, day, rows)
        write_shop_txt(os.path.join(shop_dir, nm + ".txt"), rows)

        probs = verify_written(shop_dir, nm, rows)
        status = "OK" if not probs else "СБОЙ: " + "; ".join(probs)
        if probs:
            verify_fail += 1
            for p in probs:
                R.problem("Проверка записанных файлов", shop, "", "", p)
        summary.append({"shop": shop, "positions": len(rows),
                        "units": round(float(rows["qty"].sum()), 3),
                        "folder": shop_dir, "verify": status})
        R.say("%-24s позиций: %4d   единиц: %9s   %s"
              % (shop, len(rows), round(float(rows["qty"].sum()), 3),
                 "OK" if not probs else "СБОЙ"))

    if verify_fail:
        R.check("ERROR", "Обратная сверка записанных файлов",
                "расхождения в %d магазинах" % verify_fail)
    else:
        R.check("OK", "Обратная сверка записанных файлов",
                "xlsx и txt перечитаны, совпали с расчётом (%d магазинов)" % len(summary))

    tot_units = round(sum(s["units"] for s in summary), 3)
    if abs(tot_units - round(float(orders["qty"].sum()), 3)) > 0.001:
        R.check("ERROR", "Сверка итогов", "сумма по файлам не равна сумме расчёта")
    else:
        R.check("OK", "Сверка итогов", "всего единиц %s" % tot_units)

    if MAKE_SUPPLIER_ORDER and not dist_on():
        sup = build_supplier_order(orders)
        if not sup.empty:
            sp = os.path.join(day_dir, "_ЗАКАЗ_ПОСТАВЩИКУ.xlsx")
            write_supplier_xlsx(sp, day, sup)
            R.say("Заказ поставщику    : %d позиций -> %s" % (len(sup), sp))

    write_summary(os.path.join(day_dir, "_СВОДКА.xlsx"), day, summary, src)
    write_validation(os.path.join(day_dir, "_ВАЛИДАЦИЯ.xlsx"), day, src,
                     shops_all, set(s["shop"] for s in summary))

    R.say("-" * 70)
    R.say("Магазинов с заказом : %d из %d" % (len(summary), len(shops_all)))
    R.say("Всего позиций       : %d" % sum(s["positions"] for s in summary))
    R.say("Всего единиц        : %s" % tot_units)
    R.say("Ошибок / предупр.   : %d / %d" % (len(R.errors), len(R.warns)))
    R.say("Результат           : %s" % day_dir)
    R.say("Отчёт валидации     : %s" % os.path.join(day_dir, "_ВАЛИДАЦИЯ.xlsx"))
    RUN_INFO.update({"day_dir": day_dir, "day": day,
                     "validation": os.path.join(day_dir, "_ВАЛИДАЦИЯ.xlsx"),
                     "shops": len(summary), "shops_all": len(shops_all),
                     "positions": sum(s["positions"] for s in summary),
                     "units": tot_units,
                     "dist_txt": dist_summary_text(orders)})

    history_write(day, src, fp, len(summary),
                  ENTITY_LABEL[ENTITY] + (u", раздача" if dist_on() else u""))
    if do_archive and not R.errors:
        moved = archive_source(src)
        if moved:
            R.say("Ведомость в АРХИВ   : %s" % os.path.basename(moved))

    lp = os.path.join(LOG_DIR, "log_%s.txt" % datetime.now().strftime("%Y-%m-%d_%H-%M-%S"))
    with open(lp, "w", encoding="utf-8") as f:
        f.write("\n".join(R.log))
    print("Лог: %s" % lp)
    return 2 if R.errors else 0


# ======================= ОКНО ПРОГРАММЫ ======================

MONTHS_RU = [u"января", u"февраля", u"марта", u"апреля", u"мая", u"июня",
             u"июля", u"августа", u"сентября", u"октября", u"ноября", u"декабря"]


def _res(name):
    """Файл, вшитый в exe (иконка), либо лежащий рядом со скриптом."""
    return os.path.join(getattr(sys, "_MEIPASS", _app_dir()), name)


def _win11_chrome(win, dark=False):
    """Скруглённые углы и цвет заголовка как у обычных окон Windows 11."""
    try:
        import ctypes
        win.update_idletasks()
        h = ctypes.windll.user32.GetParent(win.winfo_id())
        ctypes.windll.dwmapi.DwmSetWindowAttribute(
            h, 33, ctypes.byref(ctypes.c_int(2)), 4)          # скругление
        ctypes.windll.dwmapi.DwmSetWindowAttribute(
            h, 20, ctypes.byref(ctypes.c_int(1 if dark else 0)), 4)
    except Exception:
        pass


def run_app():
    """Окно программы: боковое меню (Ведомость, Раздача со склада, Результаты,
    Настройки) и папки по типам документов; ход расчёта и итог - внутри окна."""
    import threading, queue, json
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
    from tkinter import font as tkfont

    app = tk.Tk()
    app.withdraw()
    app.title(u"Расчёт заказа перемещений")
    try:
        app.iconbitmap(_res("raschet.ico"))
    except Exception:
        pass

    fams = set(tkfont.families(app))
    F = next((f for f in ("Segoe UI", "DejaVu Sans") if f in fams), "TkDefaultFont")
    MONO = next((f for f in ("Cascadia Mono", "Consolas", "DejaVu Sans Mono")
                 if f in fams), "TkFixedFont")
    SIDE, SIDE_ACT, SIDE_FG, SIDE_MUT, SIDE_LINE = (
        "#1f2430", "#2c3342", "#c9cfdb", "#7c8597", "#353c4b")
    BG, CARD, BORD, FG, MUT = "#f4f6f9", "#ffffff", "#e3e7ee", "#1d2433", "#667085"
    ACC, ACC_D, ACC_L, ACC_OFF = "#2f6fed", "#1f5bd0", "#e8effd", "#a9c1f5"
    GRN, GRN_L, AMB, AMB_L, RED, RED_L = ("#15803d", "#e8f6ee", "#b45309",
                                          "#fdf3e4", "#dc2626", "#fdeaea")
    app.configure(bg=BG)
    S = ttk.Style(app)
    try:
        S.theme_use("clam")
    except Exception:
        pass
    S.configure("Run.Horizontal.TProgressbar", troughcolor=BORD, background=ACC,
                bordercolor=BORD, lightcolor=ACC, darkcolor=ACC, thickness=6)

    def fnt(size, bold=False):
        return (F, size, "bold") if bold else (F, size)

    def lab(parent, text=u"", fg=FG, size=10, bold=False, **kw):
        return tk.Label(parent, text=text, fg=fg, bg=parent["bg"], font=fnt(size, bold),
                        anchor="w", justify="left", **kw)

    def card(parent):
        return tk.Frame(parent, bg=CARD, highlightbackground=BORD,
                        highlightcolor=BORD, highlightthickness=1)

    def button(parent, text, cmd, primary=False):
        """Кнопка-надпись: свой цвет на любой теме Windows, есть «неактивна»."""
        b = tk.Label(parent, text=text, font=fnt(10, primary),
                     padx=18 if primary else 12, pady=8 if primary else 5,
                     highlightthickness=0 if primary else 1,
                     highlightbackground=BORD)
        b.on = True

        def paint(hover=False):
            if primary:
                b.config(fg="#ffffff",
                         bg=(ACC_D if hover else ACC) if b.on else ACC_OFF)
            else:
                b.config(fg=FG if b.on else MUT,
                         bg=BG if (hover and b.on) else CARD)

        def enable(flag):
            b.on = bool(flag)
            b.config(cursor="hand2" if b.on else "arrow")
            paint()
        b.bind("<Enter>", lambda e: paint(True))
        b.bind("<Leave>", lambda e: paint())
        b.bind("<Button-1>", lambda e: cmd() if b.on else None)
        b.enable = enable
        enable(True)
        return b

    def link(parent, text, cmd):
        lb = tk.Label(parent, text=text, bg=parent["bg"], fg=MUT, font=fnt(9),
                      cursor="hand2")
        lb.bind("<Button-1>", lambda e: cmd())
        lb.bind("<Enter>", lambda e: lb.config(fg=ACC))
        lb.bind("<Leave>", lambda e: lb.config(fg=MUT))
        return lb

    CHIP = {"ok": (GRN, GRN_L), "none": (MUT, BG), "bad": (RED, RED_L),
            "warn": (AMB, AMB_L), "run": (ACC, ACC_L)}

    def chip(parent):
        c = tk.Label(parent, font=fnt(9, True), padx=8, pady=2)
        c.set = lambda text, kind: c.config(text=text, fg=CHIP[kind][0],
                                            bg=CHIP[kind][1])
        return c

    def opt(lbl, text, fg=None):
        """Строка, которая видна только с текстом: пустые не занимают места."""
        lbl.config(text=text, **({"fg": fg} if fg else {}))
        if text and not lbl.winfo_manager():
            lbl.pack(**lbl.pk)
        elif not text and lbl.winfo_manager():
            lbl.pack_forget()

    def wrap(lbl, box, minus=34):
        """Перенос строк по ширине карточки, а не по жёсткому числу."""
        box.bind("<Configure>", lambda e: lbl.config(
            wraplength=max(200, e.width - minus)), add="+")
        return lbl

    def open_path(pth):
        if not os.path.splitext(pth)[1]:
            try:
                os.makedirs(pth, exist_ok=True)
            except Exception:
                pass
        try:
            os.startfile(pth)
        except Exception:
            messagebox.showinfo(u"Путь", pth, parent=app)

    def rel(pth):
        try:
            r_ = os.path.relpath(pth, BASE_DIR) if pth else u""
        except ValueError:            # другой диск
            return pth
        return pth if r_.startswith("..") else r_

    def num(x):
        s_ = (u"%.3f" % float(x)).rstrip("0").rstrip(".")
        whole, _sep, frac = s_.partition(".")
        whole = u"{:,}".format(int(whole)).replace(",", u" ")
        return whole + (u"," + frac if frac else u"")

    def human(d):
        return u"%d %s" % (d.day, MONTHS_RU[d.month - 1])

    # ---------------- настройки окна: запоминаются ----------------
    prefs_path = os.path.join(LOG_DIR, u"настройки_окна.json")
    prefs = {"clean": 1, "arch": 1}
    try:
        with io.open(prefs_path, encoding="utf-8") as f_:
            prefs.update(json.load(f_))
    except Exception:
        pass
    clean_v = tk.IntVar(value=1 if prefs.get("clean") else 0)
    arch_v = tk.IntVar(value=1 if prefs.get("arch") else 0)

    def save_prefs():
        try:
            with io.open(prefs_path, "w", encoding="utf-8") as f_:
                json.dump({"clean": clean_v.get(), "arch": arch_v.get()}, f_)
        except Exception:
            pass

    # ---------------- боковое меню ----------------
    side = tk.Frame(app, bg=SIDE, width=228)
    side.pack(side="left", fill="y")
    side.pack_propagate(False)
    lab(side, u"Расчёт заказа", "#ffffff", 14, True).pack(fill="x", padx=20,
                                                          pady=(22, 0))
    lab(side, u"РЦ → магазины, файлы ТСД", SIDE_MUT, 9).pack(
        fill="x", padx=20, pady=(2, 18))
    body = tk.Frame(app, bg=BG)
    body.pack(side="left", fill="both", expand=True)
    pages, nav, cur = {}, {}, ["ved"]

    def nav_item(key, text, ic):
        fr = tk.Frame(side, bg=SIDE, cursor="hand2")
        fr.pack(fill="x", padx=10, pady=1)
        mark = tk.Frame(fr, bg=SIDE, width=3)
        mark.pack(side="left", fill="y")
        i = tk.Label(fr, text=ic, bg=SIDE, fg=SIDE_MUT, font=fnt(10, True), width=2)
        i.pack(side="left", padx=(10, 4), pady=9)
        t = tk.Label(fr, text=text, bg=SIDE, fg=SIDE_FG, font=fnt(10), anchor="w")
        t.pack(side="left", fill="x", expand=True)
        badge = tk.Label(fr, text=u"", bg=SIDE, font=fnt(9, True))
        badge.pack(side="right", padx=10)
        for w_ in (fr, mark, i, t, badge):
            w_.bind("<Button-1>", lambda e: show(key))
        nav[key] = (fr, mark, i, t, badge)

    def show(key):
        cur[0] = key
        for k, (fr, mark, i, t, badge) in nav.items():
            on = k == key
            bgc = SIDE_ACT if on else SIDE
            for w_ in (fr, i, t, badge):
                w_.config(bg=bgc)
            mark.config(bg=ACC if on else bgc)
            i.config(fg="#ffffff" if on else SIDE_MUT)
            t.config(fg="#ffffff" if on else SIDE_FG, font=fnt(10, on))
        for k, p_ in pages.items():
            if k == key:
                p_.pack(fill="both", expand=True, padx=28, pady=(22, 18))
            else:
                p_.pack_forget()
        if key == "res" and not running[0]:
            render_history()

    def badge(key, text, fg):
        nav[key][4].config(text=text, fg=fg)

    for k_, t_, i_ in (("ved", u"Ведомость", u"1"), ("dist", u"Раздача со склада", u"2"),
                       ("res", u"Результаты", u"3"), ("set", u"Настройки", u"⚙")):
        nav_item(k_, t_, i_)

    tk.Frame(side, bg=SIDE_LINE, height=1).pack(fill="x", padx=20, pady=(18, 10))
    lab(side, u"ПАПКИ", SIDE_MUT, 8, True).pack(fill="x", padx=20, pady=(0, 4))
    for t_, pth_ in ((u"Ведомости Family", IN_FM_DIR),
                     (u"Ведомости ЮА Маркет", IN_UA_DIR),
                     (u"Списки раздачи", DIST_IN_DIR),
                     (u"Готовые заказы", OUT_DIR),
                     (u"Архив ведомостей", ARC_DIR),
                     (u"Журнал и сбои", LOG_DIR)):
        fl = lab(side, u"▸  " + t_, SIDE_FG, 9, cursor="hand2")
        fl.pack(fill="x", padx=22, pady=2)
        fl.bind("<Button-1>", lambda e, p_=pth_: open_path(p_))
        fl.bind("<Enter>", lambda e, w_=fl: w_.config(fg="#ffffff"))
        fl.bind("<Leave>", lambda e, w_=fl: w_.config(fg=SIDE_FG))
    side_ref = lab(side, u"", SIDE_MUT, 8)
    side_ref.pack(side="bottom", fill="x", padx=20, pady=14)

    def page(key, title, sub):
        p_ = tk.Frame(body, bg=BG)
        pages[key] = p_
        p_.title = lab(p_, title, FG, 18, True)
        p_.title.pack(fill="x")
        p_.sub = wrap(lab(p_, sub, MUT, 10), p_, 0)
        p_.sub.pack(fill="x", pady=(2, 14))
        p_.foot = tk.Frame(p_, bg=BG)          # раньше содержимого: не обрежется
        p_.foot.pack(side="bottom", fill="x", pady=(10, 0))
        return p_

    # ================= Ведомость =================
    pv = page("ved", u"Ведомость",
              u"Оборотные ведомости из Торгсофта. Каждая база считается отдельно, "
              u"в свою папку заказов. В Торгсофт программа ничего не пишет.")
    row = tk.Frame(pv, bg=BG)
    row.pack(fill="x")
    row.grid_columnconfigure(0, weight=1, uniform="c")
    row.grid_columnconfigure(1, weight=1, uniform="c")
    src = {"FM": tk.StringVar(), "UA": tk.StringVar()}
    bad = {"FM": u"", "UA": u""}     # почему файл не подходит карточке
    fc = {}
    OWN = {"FM": IN_FM_DIR, "UA": IN_UA_DIR}

    def file_card(ent, col):
        c = card(row)
        c.grid(row=0, column=col, sticky="nsew", padx=(0, 8) if col == 0 else (8, 0))
        top = tk.Frame(c, bg=CARD)
        top.pack(fill="x", padx=16, pady=(14, 6))
        lab(top, ENTITY_LABEL[ent], FG, 12, True).pack(side="left")
        ch = chip(top)
        ch.pack(side="right")
        nm = wrap(lab(c, u"", FG, 10), c)
        nm.pack(fill="x", padx=16)
        info = wrap(lab(c, u"", MUT, 9), c)
        info.pack(fill="x", padx=16, pady=(2, 0))
        warn = wrap(lab(c, u"", RED, 9), c)
        warn.pk = dict(fill="x", padx=16, pady=(6, 0), after=info)
        bt = tk.Frame(c, bg=CARD)
        bt.pack(fill="x", side="bottom", padx=16, pady=(12, 14))
        pick = button(bt, u"Выбрать файл…", lambda: choose(ent))
        pick.pack(side="left")
        rm = link(bt, u"Убрать", lambda: load(ent, u""))
        fc[ent] = dict(chip=ch, nm=nm, info=info, warn=warn, pick=pick, rm=rm)

    file_card("FM", 0)
    file_card("UA", 1)

    dc = card(pv)
    dc.pack(fill="x", pady=(16, 0))
    dr = tk.Frame(dc, bg=CARD)
    dr.pack(fill="x", padx=16, pady=(14, 6))
    lab(dr, u"Дата отгрузки", FG, 11, True).pack(side="left")
    dv = tk.StringVar()
    tk.Entry(dr, textvariable=dv, width=11, justify="center", font=fnt(12, True),
             relief="flat", bg=CARD, fg=FG, insertbackground=FG,
             highlightthickness=1, highlightbackground=BORD, highlightcolor=ACC
             ).pack(side="left", padx=(16, 10), ipady=4)
    dorig = lab(dr, u"", MUT, 9)
    dorig.pack(side="left")
    route_l = wrap(lab(dc, u"", FG, 10), dc)
    route_l.pack(fill="x", padx=16)
    lab(dc, u"Задаёт папку результата и сверку куста, на количество не влияет.",
        MUT, 9).pack(fill="x", padx=16, pady=(2, 0))
    dspace = tk.Frame(dc, bg=CARD, height=14)
    dspace.pack(fill="x")
    dwarn = wrap(lab(dc, u"", AMB, 9), dc)
    dwarn.pk = dict(fill="x", padx=16, pady=(6, 0), before=dspace)

    go_b = button(pv.foot, u"Рассчитать заказ", lambda: start("order"), primary=True)
    go_b.pack(side="right")
    why_l = wrap(lab(pv.foot, u"", MUT, 9), pv.foot, 220)
    why_l.pack(side="right", padx=14)

    # ================= Раздача со склада =================
    pdi = page("dist", u"Раздача со склада",
               u"Раздать со склада по магазинам конкретный товар по списку. "
               u"Обычный заказ этот режим не затрагивает.")
    dcard = card(pdi)
    dcard.pack(fill="x")
    dent, distv = tk.StringVar(value="FM"), tk.StringVar()
    dcount = {}                       # список -> позиций: файл читается один раз
    nums = []

    def step(n, title):
        r_ = tk.Frame(dcard, bg=CARD)
        r_.pack(fill="x", padx=16, pady=(14 if n == 1 else 10, 0))
        nb = tk.Label(r_, text=str(n), font=fnt(10, True), width=2)
        nb.pack(side="left", anchor="n")
        nums.append(nb)
        lab(r_, title, FG, 10, True, width=14).pack(side="left", anchor="n", padx=10)
        box = tk.Frame(r_, bg=CARD)
        box.pack(side="left", fill="x", expand=True)
        return box

    b1 = step(1, u"Ведомость")
    drb = {}
    for e_ in ("FM", "UA"):
        drb[e_] = tk.Radiobutton(b1, variable=dent, value=e_, bg=CARD, fg=FG,
                                 activebackground=CARD, selectcolor=CARD,
                                 font=fnt(10), anchor="w", highlightthickness=0,
                                 bd=0, command=lambda: refresh())
        drb[e_].pack(fill="x")
    d_day = wrap(lab(b1, u"", MUT, 9), b1, 0)
    d_day.pack(fill="x", pady=(2, 0))
    b2 = step(2, u"Склад-источник")
    d_src = lab(b2, u"", FG, 10)
    d_src.pack(fill="x")
    b3 = step(3, u"Список товара")
    d_list = wrap(lab(b3, u"", MUT, 10), b3, 0)
    d_list.pack(fill="x")
    lb_ = tk.Frame(b3, bg=CARD)
    lb_.pack(fill="x", pady=(8, 0))
    button(lb_, u"Файл списка…", lambda: dist_pick()).pack(side="left")
    button(lb_, u"Папка ВЫВОЗ…", lambda: dist_pick_dir()).pack(
        side="left", padx=8)
    link(lb_, u"Убрать", lambda: (distv.set(u""), refresh())).pack(side="left", padx=6)
    tk.Frame(dcard, bg=CARD, height=16).pack(fill="x")
    wrap(lab(pdi, u"Список: файл Вывоз_ДД.ММ.ГГГГ.xlsx, папка ВЫВОЗ_<дата> или "
                  u"txt/csv «штрих-код;количество». Склад в получатели не входит. "
                  u"Раздача считается по одной ведомости.", MUT, 9),
         pdi, 0).pack(fill="x", pady=(12, 0))
    dist_b = button(pdi.foot, u"Раздать со склада", lambda: start("dist"),
                    primary=True)
    dist_b.pack(side="right")
    dwhy = wrap(lab(pdi.foot, u"", MUT, 9), pdi.foot, 220)
    dwhy.pack(side="right", padx=14)

    def dist_pick():
        f = filedialog.askopenfilename(
            parent=app, title=u"Список товара для раздачи со склада",
            initialdir=(os.path.dirname(distv.get()) if distv.get()
                        else DIST_IN_DIR if os.path.isdir(DIST_IN_DIR) else BASE_DIR),
            filetypes=[(u"Вывоз и списки", "*.xlsx *.txt *.csv"),
                       (u"Все файлы", "*.*")])
        if f:
            distv.set(os.path.abspath(f))
        refresh()

    def dist_pick_dir():
        d = filedialog.askdirectory(
            parent=app, title=u"Папка ВЫВОЗ_<дата> со скрипта вывоза",
            initialdir=(distv.get() if os.path.isdir(distv.get() or "")
                        else DIST_IN_DIR if os.path.isdir(DIST_IN_DIR) else BASE_DIR))
        if d:
            distv.set(os.path.abspath(d))
        refresh()

    # ================= Результаты =================
    pr = page("res", u"Результаты", u"В этом окне расчёта ещё не было.")
    runbox = tk.Frame(pr, bg=BG)
    bar = ttk.Progressbar(runbox, mode="indeterminate",
                          style="Run.Horizontal.TProgressbar")
    bar.pack(fill="x")
    cards_row = tk.Frame(pr, bg=BG)
    next_l = wrap(lab(pr, u"", MUT, 9), pr, 0)
    seg = tk.Frame(pr, bg=BG)
    tabs, view = {}, ["hist"]
    for k_, t_ in (("hist", u"Прошлые расчёты"), ("log", u"Ход расчёта")):
        tl = tk.Label(seg, text=t_, bg=BG, font=fnt(10, True), cursor="hand2",
                      padx=0, pady=4)
        tl.pack(side="left", padx=(0, 18))
        tl.bind("<Button-1>", lambda e, k_=k_: set_view(k_))
        tabs[k_] = tl
    hist = card(pr)
    logf = tk.Frame(pr, bg=BG)
    sb = ttk.Scrollbar(logf)
    sb.pack(side="right", fill="y")
    txt = tk.Text(logf, bg="#1b1f27", fg="#d6d6d6", font=(MONO, 9), bd=0,
                  padx=12, pady=10, wrap="none", yscrollcommand=sb.set,
                  height=12, highlightthickness=0)
    txt.pack(fill="both", expand=True)
    sb.config(command=txt.yview)
    txt.configure(state="disabled")
    for tag, col in (("ERROR", "#ff6b6b"), ("WARN", "#ffc266"),
                     ("OK", "#7fd18c"), ("INFO", "#8fb6e8")):
        txt.tag_configure(tag, foreground=col)
    new_b = button(pr.foot, u"Новый расчёт", lambda: show("ved"), primary=True)
    new_b.pack(side="right")

    def set_view(k):
        view[0] = k
        for k_, tl in tabs.items():
            tl.config(fg=ACC if k_ == k else MUT)
        hist.pack_forget()
        logf.pack_forget()
        if k == "hist":
            hist.pack(fill="x")
        else:
            logf.pack(fill="both", expand=True)

    def res_layout(state):
        """state: idle / run / done - что показывать в «Результатах»."""
        for w_ in (runbox, cards_row, next_l, seg, hist, logf):
            w_.pack_forget()
        if state == "run":
            runbox.pack(fill="x", pady=(0, 12))
        if state == "done":
            cards_row.pack(fill="x")
            next_l.pack(fill="x", pady=(10, 0))
        seg.pack(fill="x", pady=(16 if state == "done" else 0, 6))
        set_view("log" if state == "run" else view[0])

    def render_history():
        for w_ in hist.winfo_children():
            w_.destroy()
        rows_ = history_read()[-8:][::-1]
        if not rows_:
            lab(hist, u"Журнал запусков пуст: здесь появятся прошлые расчёты.",
                MUT, 9).pack(fill="x", padx=16, pady=14)
            return
        dows = (u"пн", u"вт", u"ср", u"чт", u"пт", u"сб", u"вс")

        def when(s_):
            try:
                return datetime.strptime(s_, "%Y-%m-%d %H:%M:%S").strftime("%d.%m %H:%M")
            except Exception:
                return s_

        def ship(s_):
            try:
                d_ = datetime.strptime(s_, "%Y-%m-%d")
                return u"%s %s" % (dows[d_.weekday()], d_.strftime("%d.%m"))
            except Exception:
                return s_
        for j, h_ in enumerate((u"Запуск", u"Отгрузка", u"База", u"Маг.", u"Файл")):
            tk.Label(hist, text=h_, bg=BG, fg=MUT, font=fnt(9, True), anchor="w",
                     padx=10, pady=5).grid(row=0, column=j, sticky="we")
        for k, h_ in enumerate(rows_, 1):
            for j, v in enumerate((when(h_.get("run", u"")), ship(h_.get("day", u"")),
                                   h_.get("base") or u"—", h_.get("shops", u""),
                                   (lambda f_: f_ if len(f_) <= 48 else
                                    f_[:46] + u"…")(h_.get("file", u"")))):
                tk.Label(hist, text=v, bg=CARD, fg=FG, font=fnt(9), anchor="w",
                         padx=10, pady=4).grid(row=k, column=j, sticky="we")
        hist.grid_columnconfigure(4, weight=1)

    # ================= Настройки =================
    ps = page("set", u"Настройки", u"Меняются редко. Запоминаются между запусками.")
    sc = card(ps)
    sc.pack(fill="x")

    def opt_row(var, title, sub):
        r_ = tk.Frame(sc, bg=CARD)
        r_.pack(fill="x", padx=16, pady=(12, 0))
        tk.Checkbutton(r_, variable=var, bg=CARD, activebackground=CARD,
                       highlightthickness=0, bd=0, command=save_prefs
                       ).pack(side="left", anchor="n", pady=(1, 0))
        x = tk.Frame(r_, bg=CARD)
        x.pack(side="left", fill="x", padx=6)
        t = lab(x, title, FG, 10, cursor="hand2")
        t.pack(fill="x")
        t.bind("<Button-1>", lambda e: (var.set(0 if var.get() else 1), save_prefs()))
        lab(x, sub, MUT, 9).pack(fill="x")

    opt_row(clean_v, u"Прежний расчёт дня убирать в ЗАКАЗЫ\\_замененные",
            u"чтобы в папке дня не остался старый txt магазина; ничего не удаляется")
    opt_row(arch_v, u"Ведомость после расчёта переносить в «Архив ведомостей»",
            u"чтобы завтра не подставился вчерашний файл")
    tk.Frame(sc, bg=CARD, height=14).pack(fill="x")
    lab(ps, u"Справочник упаковок", FG, 11, True).pack(fill="x", pady=(18, 6))
    rc = card(ps)
    rc.pack(fill="x")
    kv = {}
    for k_ in (u"Источник", u"Кэш", u"Рабочая папка"):
        r_ = tk.Frame(rc, bg=CARD)
        r_.pack(fill="x", padx=16, pady=(12 if k_ == u"Источник" else 4, 0))
        lab(r_, k_, MUT, 9, width=14).pack(side="left", anchor="n")
        kv[k_] = wrap(lab(r_, u"", FG, 9), r_, 140)
        kv[k_].pack(side="left", fill="x", expand=True)
    kv[u"Источник"].config(text=u"BigQuery %s + ручные правки (%s)"
                           % (BQ_TABLE.split(".")[-1], BQ_MANUAL.split(".")[-1]))
    kv[u"Рабочая папка"].config(text=BASE_DIR)
    rb_ = tk.Frame(rc, bg=CARD)
    rb_.pack(fill="x", padx=16, pady=(10, 14))
    ref_b = button(rb_, u"Обновить справочник сейчас", lambda: refresh_ref())
    ref_b.pack(side="left")
    ref_msg = wrap(lab(rb_, u"", MUT, 9), rb_, 260)
    ref_msg.pack(side="left", padx=12)

    def ref_sync():
        if not os.path.exists(REF_FILE):
            kv[u"Кэш"].config(text=u"%s · файла нет: при расчёте возьмётся из BigQuery, "
                                   u"без него - по названиям товаров" % rel(REF_FILE))
            side_ref.config(text=u"справочник: кэша нет")
            return
        ts = datetime.fromtimestamp(os.path.getmtime(REF_FILE))
        dd = (date.today() - ts.date()).days
        when = (u"сегодня в %s" % ts.strftime("%H:%M") if dd == 0 else
                u"вчера в %s" % ts.strftime("%H:%M") if dd == 1 else
                ts.strftime("%d.%m.%Y %H:%M"))
        try:
            with io.open(REF_FILE, encoding="utf-8-sig", errors="replace") as f_:
                n = max(0, sum(1 for _ in f_) - 1)
        except Exception:
            n = 0
        fresh = (time.time() - os.path.getmtime(REF_FILE)) < REF_TTL_H * 3600
        kv[u"Кэш"].config(text=u"%s · обновлён %s · %s позиций · %s"
                              % (rel(REF_FILE), when, num(n),
                                 u"свежий" if fresh else
                                 u"при расчёте обновится из BigQuery"))
        side_ref.config(text=u"справочник: %s" % when)

    ref_busy = [False]

    def refresh_ref():
        if REF_OFFLINE:
            ref_msg.config(text=u"Режим офлайн: BigQuery выключен в настройках "
                                u"скрипта (REF_OFFLINE)", fg=AMB)
            return
        if running[0] or ref_busy[0]:
            return
        ref_busy[0] = True
        ref_b.enable(False)
        refresh()
        ref_msg.config(text=u"Обновляю из BigQuery…", fg=MUT)
        res = {}

        def work():
            try:
                res["r"] = refresh_reference_from_bq()
            except Exception as e:
                res["r"] = (False, u"%s: %s" % (type(e).__name__, e))
        th = threading.Thread(target=work, daemon=True)
        th.start()

        def poll():
            if th.is_alive():
                app.after(200, poll)
                return
            ok, msg = res["r"]
            ref_msg.config(text=(u"Обновлено: " if ok else u"Не обновлено: ") + msg,
                           fg=GRN if ok else RED)
            ref_busy[0] = False
            ref_b.enable(True)
            ref_sync()
            refresh()
        app.after(200, poll)

    # ================= проверка выбора =================
    origin = [u""]
    running = [False]

    def picked():
        return [(e, src[e].get()) for e in ("FM", "UA") if src[e].get()]

    def parse_dv():
        t = dv.get().strip().replace("/", ".").replace("-", ".")
        m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})$", t)
        if not m:
            return None
        d, mo, y = (int(g) for g in m.groups())
        y = y + 2000 if y < 100 else y
        try:
            return date(y, mo, d)
        except ValueError:
            return None

    def out_name(d, ent, dist=False):
        return u"ЗАКАЗЫ\\%s%s%s" % (d.strftime("%Y-%m-%d"),
                                    DIST_DIR_SUFFIX if dist else u"",
                                    ENTITY_DIR_SUFFIX[ent])

    def refresh(*_a):
        d = parse_dv()
        busy = (u"Идёт расчёт - итог в разделе «Результаты»" if running[0] else
                u"Обновляется справочник, подождите" if ref_busy[0] else u"")
        # ---- Ведомость ----
        wrong = [e for e in ("FM", "UA") if src[e].get() and bad[e]]
        why = (busy or
               (u"Выберите ведомость хотя бы одной базы" if not picked() else
                u"Ведомость %s не подходит - замените или уберите"
                % ENTITY_LABEL[wrong[0]] if wrong else
                u"Укажите дату отгрузки в виде 15.10.2026" if d is None else u""))
        go_b.enable(not why)
        if why:
            why_l.config(text=u"⚠ " + why, fg=AMB)
        else:
            why_l.config(fg=MUT, text=u"Будет записано: " + u";  ".join(
                u"%s → %s" % (ENTITY_LABEL[e], out_name(d, e)) for e, _p in picked()))
        if not dv.get().strip():
            route_l.config(text=u"")
            opt(dwarn, u"")
        elif d is None:
            route_l.config(text=u"Дата не понята. Нужен вид 15.10.2026", fg=RED)
            opt(dwarn, u"")
        else:
            rk = ROUTE_BY_DOW.get(d.weekday())
            route_l.config(fg=FG, text=(
                u"%s, %s · группа %s · по графику %d магазинов"
                % (DOW_RU[d.weekday()].capitalize(), human(d), rk, len(ROUTES[rk]))
                if rk else u"%s — по графику отгрузки нет" % human(d)))
            msgs = []
            for e, _p in picked():
                _dd, n, ts = day_dir_state(d, ENTITY_DIR_SUFFIX[e])
                if n:
                    msgs.append(u"%s: на %s расчёт уже делали %s — %d магазинов. "
                                u"Прежний уберём в папку замен."
                                % (ENTITY_LABEL[e], human(d),
                                   ts.strftime("%d.%m в %H:%M") if ts else "?", n))
            off = (d - date.today()).days
            if abs(off) > DATE_SANITY_DAYS:
                msgs.append(u"Это %s от сегодняшнего дня — проверьте, тот ли файл."
                            % (u"на %d дн. вперёд" % off if off > 0
                               else u"на %d дн. назад" % -off))
            opt(dwarn, u"\n".join(msgs), AMB)
        # ---- Раздача со склада ----
        avail = [e for e in ("FM", "UA") if src[e].get() and not bad[e]]
        for e, rb in drb.items():
            rb.config(state="normal" if e in avail else "disabled",
                      text=u"%s · %s" % (ENTITY_LABEL[e],
                                         os.path.basename(src[e].get())
                                         if e in avail else u"ведомость не выбрана"))
        if avail and dent.get() not in avail:
            dent.set(avail[0])
        de = dent.get()
        d_day.config(text=(u"отгрузка %s · меняется в разделе «Ведомость»"
                           % d.strftime("%d.%m.%Y")) if d else
                     u"сначала выберите ведомость и дату в разделе «Ведомость»")
        d_src.config(text=u"%s — по базе ведомости"
                     % (DIST_SOURCE_SHOP_UA if de == "UA" else DIST_SOURCE_SHOP))
        p_ = distv.get()
        if p_ and p_ not in dcount:
            dcount[p_] = dist_count(p_)
        d_list.config(text=(u"%s · %d позиций" % (os.path.basename(p_.rstrip("\\/")),
                                                  dcount.get(p_, 0))
                            if p_ else u"не выбран"), fg=FG if p_ else MUT)
        for nb, done in zip(nums, (bool(avail and d), bool(avail), bool(p_))):
            nb.config(bg=GRN_L if done else ACC_L, fg=GRN if done else ACC)
        dwhy_t = (busy or
                  (u"Выберите ведомость в разделе «Ведомость»" if not avail else
                   u"Укажите дату отгрузки в разделе «Ведомость»" if d is None else
                   u"Выберите список товара" if not p_ else u""))
        dist_b.enable(not dwhy_t)
        dwhy.config(text=(u"⚠ " + dwhy_t) if dwhy_t else
                    u"Будет записано: " + out_name(d, de, True),
                    fg=AMB if dwhy_t else MUT)
        ref_b.enable(not (running[0] or ref_busy[0]))

    def set_date():
        """Дата из имени ведомости Family, если её нет - из ведомости ЮА."""
        p_ = src["FM"].get() or src["UA"].get()
        if not p_:
            dv.set(u"")
            dorig.config(text=u"")
            return
        g, o = guess_day(p_)
        origin[0] = o
        dv.set(g.strftime("%d.%m.%Y"))
        dorig.config(text={u"из имени файла": u"взята из названия файла"}.get(o, o))

    def load(ent, path):
        src[ent].set(path or u"")
        c = fc[ent]
        bad[ent] = u""
        opt(c["warn"], u"")
        c["pick"].config(text=u"Заменить…" if path else u"Выбрать файл…")
        if not path:
            c["rm"].pack_forget()
            c["chip"].set(u"не выбрана", "none")
            c["nm"].config(fg=MUT, text=(
                u"В папке «Ведомости Family» файла нет" if ent == "FM" else
                u"Нужна, если в этот день везём в магазины «… ЮА»"))
            c["info"].config(fg=MUT, text=u"Положите выгрузку в папку «%s» или "
                                          u"нажмите «Выбрать файл…»"
                                          % os.path.basename(OWN[ent]))
            set_date()
            refresh()
            return
        c["rm"].pack(side="left", padx=12)
        c["nm"].config(fg=FG, text=os.path.basename(path))
        where = (u"" if folder_entity(path) == ent
                 else u"\nпапка: " + rel(os.path.dirname(path)))
        stamp = read_export_stamp(path)
        stale = False
        if stamp:
            age = (datetime.now() - stamp).total_seconds() / 3600.0
            stale = age > STALE_EXPORT_HOURS
            c["info"].config(fg=AMB if stale else GRN, text=(
                u"выгружена %s, это %s%s"
                % (stamp.strftime("%d.%m.%Y в %H:%M"),
                   u"%.0f ч назад" % age if age < 48 else u"%.0f дн. назад" % (age / 24.0),
                   where)))
        else:
            c["info"].config(fg=MUT, text=u"в файле не записано, когда его выгрузили"
                                          + where)
        got = file_entity(path)
        if got == "MIX":
            bad[ent] = (u"В файле магазины обеих баз. Выгрузите ведомость "
                        u"каждой базы отдельно")
        elif got and got != ent:
            bad[ent] = (u"Это ведомость базы %s (у магазинов %s «%s» в конце). "
                        u"Переложите её в папку «%s»"
                        % (ENTITY_LABEL[got], u"есть" if got == "UA" else u"нет",
                           UA_SHOP_SUFFIX, os.path.basename(OWN[got])))
        elif ent == "UA" and path == src["FM"].get():
            bad[ent] = u"Этот же файл выбран в карточке Family"
        pv_ = history_same_file(source_fingerprint(path))
        if bad[ent]:
            c["chip"].set(u"✕ не подходит", "bad")
        elif stale or pv_:
            c["chip"].set(u"проверьте", "warn")
        else:
            c["chip"].set(u"✓ готово", "ok")
        opt(c["warn"], u"\n".join(x for x in (
            bad[ent],
            (u"Этот же файл уже считали %s — заказ лёг в папку %s"
             % (pv_["run"], pv_["day"])) if pv_ else u"") if x),
            RED if bad[ent] else AMB)
        set_date()
        refresh()

    def choose(ent):
        cur_ = src[ent].get()
        f = filedialog.askopenfilename(
            parent=app, title=u"Оборотная ведомость из базы %s" % ENTITY_LABEL[ent],
            initialdir=(os.path.dirname(cur_) if cur_ else
                        OWN[ent] if os.path.isdir(OWN[ent]) else BASE_DIR),
            filetypes=[(u"Книги Excel", "*.xlsx *.xlsm"), (u"Все файлы", "*.*")])
        if f:
            load(ent, os.path.abspath(f))

    def reload_sources():
        found = find_sources_by_entity(quiet=True)
        load("FM", found.get("FM", u""))
        load("UA", found.get("UA", u""))

    dv.trace_add("write", refresh)

    # ================= расчёт =================
    def put(line):
        tag = ("ERROR" if line.startswith("[ОШИБ") else
               "WARN" if line.startswith("[ВНИМ") else
               "OK" if line.startswith("[ OK ") else
               "INFO" if line.startswith("[ИНФО") else "")
        txt.configure(state="normal")
        txt.insert("end", line + "\n", tag)
        txt.see("end")
        txt.configure(state="disabled")

    def start(kind):
        global GUI_CHOICE, GUI_SOURCES, DIST_FILE
        b = go_b if kind == "order" else dist_b
        if running[0] or ref_busy[0] or not b.on:
            return
        d = parse_dv()
        running[0] = True
        if kind == "dist":
            DIST_FILE = distv.get()
            GUI_SOURCES = [(dent.get(), src[dent.get()].get())]
        else:
            DIST_FILE = u""
            GUI_SOURCES = picked()
        GUI_CHOICE = (GUI_SOURCES[0][1], d, bool(clean_v.get()),
                      bool(arch_v.get()), origin[0])
        refresh()
        pr.title.config(text=u"Считаю…", fg=FG)
        pr.sub.config(text=u"%s%s · отгрузка %s"
                      % (u"Раздача со склада · " if kind == "dist" else u"",
                         u" + ".join(ENTITY_LABEL[e] for e, _p in GUI_SOURCES),
                         human(d)))
        badge("res", u"…", ACC_OFF)
        new_b.enable(False)
        txt.configure(state="normal")
        txt.delete("1.0", "end")
        txt.configure(state="disabled")
        res_layout("run")
        show("res")
        bar.start(12)

        R.checks, R.problems, R.anomalies, R.log = [], [], [], []
        RUN_INFO.clear()
        q = queue.Queue()

        def say(m):
            R.log.append(str(m))
            q.put(str(m))
        R.say = say

        def worker():
            try:
                code = main()
            except Exception as e:
                import traceback
                q.put(u"СБОЙ: %s: %s" % (type(e).__name__, e))
                for ln in traceback.format_exc().splitlines():
                    q.put(ln)
                code = 2
            q.put(("done", code))

        threading.Thread(target=worker, daemon=True).start()

        def poll():
            done = None
            while True:
                try:
                    it = q.get_nowait()
                except queue.Empty:
                    break
                if isinstance(it, tuple):
                    done = it[1]
                else:
                    put(it)
            if done is None:
                app.after(70, poll)
            else:
                finish(done)
        app.after(70, poll)

    def res_card(j, col, n):
        c = card(cards_row)
        c.grid(row=0, column=col, sticky="nsew",
               padx=(0, 0) if n == 1 else (0, 8) if col == 0 else (8, 0))
        top = tk.Frame(c, bg=CARD)
        top.pack(fill="x", padx=16, pady=(14, 8))
        lab(top, ENTITY_LABEL.get(j["entity"], u"") +
            (u" · раздача" if j.get("dist") else u""), FG, 12, True).pack(side="left")
        ch = chip(top)
        ch.pack(side="right")
        ok = j["code"] == 0 and not j["errors"]
        if ok and j.get("day_dir"):
            ch.set(u"✓ готово", "ok")
        elif ok:
            ch.set(u"заказывать нечего", "warn")
        elif j.get("day_dir"):
            ch.set(u"✕ в ТСД не заливать", "bad")
        else:
            ch.set(u"✕ файлы не записаны", "bad")
        if j.get("day_dir"):
            g = tk.Frame(c, bg=CARD)
            g.pack(fill="x", padx=16)
            for k, (lb, v) in enumerate(((u"магазинов", u"%d / %d" % (j["shops"],
                                                                      j["shops_all"])),
                                         (u"позиций", num(j["positions"])),
                                         (u"единиц", num(j["units"])),
                                         (u"предупр.", str(j["warns"])))):
                cell = tk.Frame(g, bg=CARD)
                cell.pack(side="left", expand=True, fill="x")
                lab(cell, v, AMB if k == 3 and j["warns"] else FG, 15, True).pack(fill="x")
                lab(cell, lb, MUT, 8).pack(fill="x")
        msg = []
        if not ok:
            msg.append(u"Ошибок %d, предупреждений %d. %s"
                       % (j["errors"], j["warns"],
                          u"Файлы записаны, но проверку не прошли."
                          if j.get("day_dir") else u""))
            msg += [u"• " + t_[:160] for t_ in j.get("err_text", [])]
        elif not j.get("day_dir"):
            msg.append(u"По формуле ни одна позиция не дала количество больше нуля. "
                       u"Предупреждений %d." % j["warns"])
        if j.get("dist_txt"):
            msg.append(j["dist_txt"].strip())
        if j.get("day_dir"):
            msg.append(rel(j["day_dir"]))
        wrap(lab(c, u"\n".join(msg), RED if not ok else MUT, 9), c).pack(
            fill="x", padx=16, pady=(10, 0))
        bt = tk.Frame(c, bg=CARD)
        bt.pack(fill="x", side="bottom", padx=16, pady=(10, 14))
        if j.get("day_dir"):
            button(bt, u"Открыть заказы", lambda: open_path(j["day_dir"])).pack(side="left")
            button(bt, u"Отчёт проверки", lambda: open_path(j["validation"])).pack(
                side="left", padx=8)
        elif not ok:
            button(bt, u"Журнал и сбои", lambda: open_path(LOG_DIR)).pack(side="left")

    def finish(code):
        running[0] = False
        bar.stop()
        jobs = RUN_INFO.get("jobs", [])
        ok = [j for j in jobs if j["code"] == 0 and not j["errors"] and j.get("day_dir")]
        fail = [j for j in jobs if j["code"] or j["errors"]]
        for w_ in cards_row.winfo_children():
            w_.destroy()
        for k in (0, 1):
            cards_row.grid_columnconfigure(k, weight=1 if k < max(1, len(jobs)) else 0,
                                           uniform="r")
        for k, j in enumerate(jobs[:2]):
            res_card(j, k, min(2, len(jobs)))
        if not jobs:
            wrap(lab(cards_row, u"Расчёт прервался, итогов нет - причина в «Ходе "
                                u"расчёта».", RED, 10), cards_row, 0).grid(
                                    row=0, column=0, sticky="we")
        title, col = ((u"Готово", GRN) if ok and not fail else
                      (u"Готово не всё", AMB) if ok else
                      (u"Есть ошибки", RED) if fail or not jobs else
                      (u"Заказывать нечего", AMB))
        pr.title.config(text=title, fg=col)
        d = jobs[0].get("day") if jobs else None
        pr.sub.config(text=u"Расчёт %s%s" % (
            datetime.now().strftime("в %H:%M"),
            u", отгрузка %s" % human(d) if d else u""))
        nxt = []
        if ok:
            nxt.append(u"Откройте «Отчёт проверки» и просмотрите жёлтые строки. "
                       u"В папке каждого магазина xlsx для печати и txt для ТСД; "
                       u"txt базы ЮА заливаются из базы ЮА.")
        if fail or not jobs:
            nxt.append(u"Где ошибки: красные строки в «Ходе расчёта», отчёт - "
                       u"_ВАЛИДАЦИЯ.xlsx в папке заказов или ВАЛИДАЦИЯ_СБОЙ_…xlsx в "
                       u"«Журнал и сбои». Исправьте файл и запустите ещё раз.")
        next_l.config(text=u"Что дальше: " + u" ".join(nxt) if nxt else u"")
        view[0] = "log" if (fail or not jobs) else "hist"
        badge("res", u"✓" if ok and not fail else u"!", GRN if ok and not fail else
              AMB if ok else RED)
        new_b.enable(True)
        render_history()
        res_layout("done")
        txt.see("end")
        reload_sources()     # обработанные ведомости могли уехать в архив
        ref_sync()

    app.bind("<Return>", lambda e: start("order") if cur[0] == "ved" else
             start("dist") if cur[0] == "dist" else None)
    ensure_dirs()
    ref_sync()
    reload_sources()
    res_layout("idle")
    show("ved")
    sw, sh = app.winfo_screenwidth(), app.winfo_screenheight()
    w_, h_ = min(1000, sw - 40), min(660, sh - 80)
    app.geometry("%dx%d+%d+%d" % (w_, h_, max(0, (sw - w_) // 2),
                                  max(0, (sh - h_) // 3)))
    app.minsize(min(900, w_), min(600, h_))
    app.deiconify()
    _win11_chrome(app)
    app.lift()
    app.focus_force()
    app.mainloop()
    return 0


if __name__ == "__main__":
    _args = [a.lower().lstrip("-/") for a in sys.argv[1:]]
    if USE_GUI and "auto" not in _args and "a" not in _args:
        try:
            sys.exit(run_app())
        except Exception as _e:
            try:
                import tkinter.messagebox as _mb
                _mb.showerror("Расчёт заказа",
                              "Окно не открылось:\n%s: %s\n\n"
                              "Запустите с ключом --auto для работы без окна."
                              % (type(_e).__name__, _e))
            except Exception:
                print("СБОЙ: %s: %s" % (type(_e).__name__, _e))
            sys.exit(2)
    try:
        code = main()
    except Exception as e:
        print("СБОЙ: %s: %s" % (type(e).__name__, e))
        import traceback
        traceback.print_exc()
        code = 2
    try:
        input("\nГотово. Нажмите Enter для выхода...")
    except (EOFError, RuntimeError):
        pass
    sys.exit(code)
